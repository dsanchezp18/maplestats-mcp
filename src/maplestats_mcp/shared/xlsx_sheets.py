"""Read `.xlsx` workbooks sheet by sheet as published text, for agency tables.

Agency workbooks (BC Stats, provincial statistics bureaus) are laid out for
people: title rows above the table, two- and three-row column headers, years
across columns. Nothing here reshapes them; a sheet comes back as rows of
text with a guessed header row, so the caller sees what the agency published.
nl_stats keeps its own copy of the same cell and header logic (it also reads
legacy `.xls`) and is left as is.
"""

from __future__ import annotations

import io
import warnings
from datetime import date, datetime
from typing import Any

from maplestats_mcp.shared.errors import InvalidInput
from maplestats_mcp.shared.executor import check_deadline
from maplestats_mcp.shared.fr_typography import call_error

# A header row is the first of the first rows with this many filled cells.
HEADER_MIN_CELLS = 3
HEADER_SEARCH_ROWS = 30
# Declared sheet sizes above this are formatting, not data (see fix_dimensions).
SUSPICIOUS_SHEET_CELLS = 10_000_000


def cell_text(value: Any) -> str:
    """One cell as text: whole floats lose the `.0`, dates are ISO, blanks are ''."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() and abs(value) < 1e15 else repr(value)
    if isinstance(value, datetime):
        return value.isoformat(sep=" ") if value.time() else value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return " ".join(str(value).split())


def trim_rows(rows: list[list[str]]) -> list[list[str]]:
    """Drop trailing empty rows and columns, and pad every row to one width."""
    while rows and not any(rows[-1]):
        rows.pop()
    width = max((max((i for i, c in enumerate(r) if c), default=-1) + 1 for r in rows), default=0)
    return [r[:width] + [""] * (width - len(r[:width])) for r in rows]


def open_workbook(body: bytes) -> Any:
    """A read-only, values-only openpyxl workbook over the file's bytes."""
    # openpyxl takes ~0.7 s to import, so it loads on first use.
    from openpyxl import load_workbook

    with warnings.catch_warnings():
        # Agency workbooks carry print-area names openpyxl cannot set; harmless for reading.
        warnings.simplefilter("ignore", UserWarning)
        return load_workbook(io.BytesIO(body), read_only=True, data_only=True)


def fix_dimensions(sheet: Any) -> None:
    """Re-read a sheet whose declared <dimension> cannot be trusted.

    Two writer bugs, both confirmed live:

    - a lone A1 on a full sheet, so openpyxl yields one row;
    - an inflated range from formatted blank cells: the NWT traffic workbook
      declares 65,536 rows by 16,217 columns for about 9,800 rows of 109, and
      openpyxl pads every row to the declared width, so one pass took 250 s
      (7.5 s once the dimensions were reset).

    Resetting is not free: a sheet with millions of styled empty rows (a 3.8 MB
    workbook of 30 real rows) takes 10 s per pass, so ordinary dimensions are
    kept.
    """
    rows, columns = sheet.max_row or 0, sheet.max_column or 0
    if (rows <= 1 and columns <= 1) or rows * columns > SUSPICIOUS_SHEET_CELLS:
        sheet.reset_dimensions()


def sheet_dimensions(body: bytes) -> list[tuple[str, int | None, int | None]]:
    """Sheet names with the row and column counts the file declares (None if absent)."""
    workbook = open_workbook(body)
    try:
        return [(ws.title, ws.max_row, ws.max_column) for ws in workbook.worksheets]
    finally:
        workbook.close()


def read_sheet(body: bytes, sheet: str, max_rows: int) -> tuple[list[list[str]], bool]:
    """Rows of one sheet as text, capped at `max_rows`; the flag says rows were cut."""
    workbook = open_workbook(body)
    try:
        if sheet not in workbook.sheetnames:
            raise call_error(
                InvalidInput,
                f"No sheet {sheet!r}; sheets are {workbook.sheetnames}.",
                f"Aucune feuille {sheet!r} ; les feuilles sont {workbook.sheetnames}.",
            )
        rows: list[list[str]] = []
        capped = False
        worksheet = workbook[sheet]
        fix_dimensions(worksheet)
        for raw in worksheet.iter_rows(values_only=True):
            if len(rows) >= max_rows:
                capped = True
                break
            if len(rows) % 100 == 0:
                check_deadline()
            rows.append([cell_text(c) for c in raw])
        return trim_rows(rows), capped
    finally:
        workbook.close()


def guess_header(rows: list[list[str]]) -> int | None:
    """0-based index of the first row with at least three filled cells."""
    for index, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        if sum(1 for cell in row if cell) >= HEADER_MIN_CELLS:
            return index
    return None
