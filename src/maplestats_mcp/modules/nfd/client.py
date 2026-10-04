"""Client for the National Forestry Database's published files.

Checked live 2026-10-02 against all 25 tables (CSV), their data
dictionaries and comments workbooks:

1. The Download page (EN and FR) lists 25 tables, not the 37 the brief
   mentioned: one `<table>` per dataset with a caption "3.2.1. Area burned
   by cause class" and links to the CSV, XLSX, dictionary (.xlsx) and
   comments (.xls). The numbering is the NFD's own (2, 3.1.1 ... 8.2.3).
2. Every CSV is one bilingual file: UTF-8 with BOM, CRLF, with an English
   and a French column for each label and for the year and the value
   ("Year,Année,ISO,Jurisdiction,Juridiction,<dimension pairs>,<value>,
   Data qualifier,<valeur>,Qualificatifs"). The two value columns were
   identical in the files compared. Property losses (3.3) differs: six
   columns with a combined "Year / Année" and no label pairs.
3. Values are numbers or blank; a blank usually comes with a qualifier code
   u, U or n, but not always: the area-burned-by-cause file has a blank
   with qualifier "a" (1999 PE, live 2026-10-03). Every blank is read as
   null whatever its code. Codes are case-sensitive (E and e differ). One
   harvest row has a value and no qualifier.
4. Labels carry footnote markers ("Prescribed burning*b", "Fuelwood*b and
   firewood*c"); the text is in the comments workbook.
5. Yukon is "YK" in the scarification table and "YT" elsewhere; wood supply
   has jurisdiction GC ("Canada") and property losses has NP (national
   parks). French month names in "Area burned by month" appear in two
   capitalizations ("Août" and "août").
6. Some tables repeat a key (same year, jurisdiction and categories): 110
   groups in wood supply, 10 in roundwood harvest. They are kept as
   published and counted in the table description.
7. The data dictionary (.xlsx, sheets "English" and "Francais") starts with
   title, file name, last update and source, then one block per field.
8. The comments workbook has a sheet per language ("21_EN", "21_FR") with
   ISO, jurisdiction, year, comment and footnote text; six fire tables have
   none (HTTP 404).
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from collections.abc import Awaitable, Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any, Literal
from urllib.parse import quote, urljoin

import httpx
from bs4 import BeautifulSoup

from maplestats_mcp.modules.nfd import constants
from maplestats_mcp.modules.nfd.schemas import (
    NfdComment,
    NfdCommentsResult,
    NfdDimension,
    NfdJurisdiction,
    NfdQueryResult,
    NfdRow,
    NfdTableDescription,
    NfdTableList,
    NfdTableSummary,
)
from maplestats_mcp.shared.cache import cached_fetch
from maplestats_mcp.shared.csv_files import decode
from maplestats_mcp.shared.envelope import make_provenance
from maplestats_mcp.shared.errors import InvalidInput, NotFound, UpstreamError, UpstreamUnavailable
from maplestats_mcp.shared.fr_typography import (
    fr_or_en,
    french_spacing,
    lang_error,
    truncation_note_lang,
)
from maplestats_mcp.shared.http import get_raw
from maplestats_mcp.shared.licences import LICENCES_FR, OGL_CANADA
from maplestats_mcp.shared.models import Provenance
from maplestats_mcp.shared.rate_limiter import get_limiter

_LICENCE = f"{OGL_CANADA} NFD terms of use: {constants.TERMS_URL}"
_LICENCE_FR = french_spacing(
    f"{LICENCES_FR[OGL_CANADA]} Conditions d'utilisation de la BNDF : {constants.TERMS_URL}"
)
_SOURCE_ERRORS = (NotFound, UpstreamError, UpstreamUnavailable)

Lang = Literal["en", "fr"]
Filter = str | Sequence[str] | None

_LIMITER = get_limiter(
    constants.RATE_LIMIT_SOURCE,
    rate=constants.RATE_LIMIT_PER_SECOND,
    capacity=constants.RATE_LIMIT_CAPACITY,
)

_FOOTNOTE = re.compile(r"\*([a-z])\b")
_TABLE_ID = re.compile(r"^\s*(\d+(?:\.\d+)*)\.?\s*(.*)$", re.DOTALL)
# The NFD's own spelling of a column, mapped to the name tools expose.
_KEY_OVERRIDES = {"renenues": "revenue_type"}
_UNIT_BY_HEADER = {
    "number": ("number of fires", "nombre d'incendies"),
    "number of seedlings": ("seedlings", "semis"),
    "dollars": ("dollars", "dollars"),
    "value": (None, None),
}
_UNIT_FR = {"hectares": "hectares", "cubic metres": "mètres cubes", "dollars": "dollars"}
_SHEET_LANG = {"en": "EN", "fr": "FR"}


# ------------------------------------------------------------------ helpers


def fold(text: str) -> str:
    """Accent-free, case-free, single-spaced form used to compare labels."""
    decomposed = unicodedata.normalize("NFKD", text.replace("’", "'"))
    stripped = "".join(c for c in decomposed if not unicodedata.combining(c))
    return " ".join(stripped.casefold().split())


def _key_from_label(label: str) -> str:
    base = re.sub(r"\s*\((en|fr)\)\s*$", "", label.strip(), flags=re.IGNORECASE)
    key = re.sub(r"[^a-z0-9]+", "_", fold(base)).strip("_")
    return _KEY_OVERRIDES.get(key, key)


def clean_label(raw: str) -> tuple[str, list[str]]:
    """The label without footnote markers ('Prescribed burning*b'), and the letters."""
    marks = _FOOTNOTE.findall(raw)
    return " ".join(_FOOTNOTE.sub("", raw).split()), marks


def normalize_table_id(table_id: str) -> str:
    return table_id.strip().rstrip(".").replace(" ", "")


def _as_list(value: Filter) -> list[str]:
    if value is None:
        return []
    return [value] if isinstance(value, str) else [v for v in value if v is not None]


def _suggest(value: str, labels: Iterable[str], lang: str = "en") -> str:
    wanted = fold(value)
    close = sorted({lb for lb in labels if lb and wanted and wanted in fold(lb)})
    if not close:
        return ""
    names = ", ".join(repr(v) for v in close[:10])
    return fr_or_en(lang, f"; did you mean {names}?", f" ; vouliez-vous dire {names} ?")


async def _in_lang[T](awaitable: Awaitable[T], lang: str) -> T:
    """Await a download or parse; for lang="fr", frame its English error in French."""
    try:
        return await awaitable
    except _SOURCE_ERRORS as exc:
        if lang != "fr":
            raise
        raise lang_error(
            type(exc),
            "fr",
            "",
            f"nfd : le fichier de la BNDF n'a pas pu être lu (détail technique en anglais) : {exc}",
        ) from exc


# ----------------------------------------------------------------- download


async def _download(url: str, context: str) -> httpx.Response:
    """GET one file, following up to two redirects (the site may move to https)."""
    current = url
    for _ in range(3):
        await _LIMITER.acquire()
        try:
            response = await get_raw(current, timeout=constants.DOWNLOAD_TIMEOUT_SECONDS)
        except httpx.HTTPStatusError as exc:
            status = exc.response.status_code
            location = exc.response.headers.get("location")
            if status in (301, 302, 303, 307, 308) and location:
                current = urljoin(current, location)
                continue
            if status == 404:
                raise NotFound(f"{context}: no file at {current}.") from exc
            raise UpstreamError(f"{context}: {current} returned HTTP {status}.") from exc
        except httpx.HTTPError as exc:
            raise UpstreamUnavailable(
                f"{context}: nfdp.ccfm.org did not answer for {current} "
                f"({type(exc).__name__}); try again shortly."
            ) from exc
        # get_raw returns 406/409 unraised (StatCan uses them); without this
        # check a 409 error page was parsed as data and reported as NotFound
        # "returned a web page" (test_client_edges, 2026-10-03).
        if response.is_error:
            raise UpstreamError(f"{context}: {current} returned HTTP {response.status_code}.")
        if len(response.content) > constants.MAX_FILE_BYTES:
            raise UpstreamError(f"{context}: {current} is larger than this tool reads.")
        return response
    raise UpstreamError(f"{context}: {url} redirected too many times.")


def _file_url(path: str) -> str:
    """Absolute URL for a path from the Download page (names contain spaces and commas)."""
    return constants.BASE_URL + quote(path, safe="/")


# ---------------------------------------------------------------- catalogue


@dataclass(frozen=True)
class CatalogueEntry:
    table_id: str
    section_en: str
    section_fr: str
    title_en: str
    title_fr: str
    csv_path: str
    xlsx_path: str | None
    dictionary_path_en: str | None
    dictionary_path_fr: str | None
    comments_path: str | None


def _parse_page(html: str) -> dict[str, dict[str, Any]]:
    """Table id -> section, title and file paths, from one Download page."""
    soup = BeautifulSoup(html, "html.parser")
    found: dict[str, dict[str, Any]] = {}
    for table in soup.find_all("table"):
        caption = table.find("caption")
        if caption is None:
            continue
        match = _TABLE_ID.match(" ".join(caption.get_text().split()))
        if match is None:
            continue
        heading = table.find_previous("h2")
        entry: dict[str, Any] = {
            "section": " ".join(heading.get_text().split()) if heading else "",
            "title": match.group(2).strip(),
            "csv": None,
            "xlsx": None,
            "dictionary": None,
            "comments": None,
        }
        for link in table.find_all("a", href=True):
            href = str(link["href"])
            lowered = href.lower()
            if lowered.endswith(".csv"):
                entry["csv"] = href
            elif "data_dictionary" in lowered:
                entry["dictionary"] = href
            elif "/comments/" in lowered:
                entry["comments"] = href
            elif lowered.endswith(".xlsx"):
                entry["xlsx"] = href
        found[match.group(1)] = entry
    return found


async def _load_catalogue() -> tuple[dict[str, CatalogueEntry], bool]:
    async def fetch() -> dict[str, CatalogueEntry]:
        pages = {}
        for lang, url in (("en", constants.PAGE_EN), ("fr", constants.PAGE_FR)):
            response = await _download(url, "nfd download page")
            pages[lang] = _parse_page(decode(response.content))
        english, french = pages["en"], pages["fr"]
        entries: dict[str, CatalogueEntry] = {}
        for table_id, en in english.items():
            if not en["csv"]:
                continue
            fr = french.get(table_id, {})
            entries[table_id] = CatalogueEntry(
                table_id=table_id,
                section_en=en["section"],
                section_fr=fr.get("section") or en["section"],
                title_en=en["title"],
                title_fr=fr.get("title") or en["title"],
                csv_path=en["csv"],
                xlsx_path=en["xlsx"],
                dictionary_path_en=en["dictionary"],
                dictionary_path_fr=fr.get("dictionary") or en["dictionary"],
                comments_path=en["comments"],
            )
        if not entries:
            raise UpstreamError("nfd download page: no tables found; the page layout changed.")
        return entries

    return await cached_fetch("nfd:catalogue", constants.CATALOGUE_TTL_SECONDS, fetch)


async def _entry(table_id: str, lang: str = "en") -> tuple[CatalogueEntry, bool]:
    entries, cached = await _in_lang(_load_catalogue(), lang)
    wanted = normalize_table_id(table_id)
    entry = entries.get(wanted)
    if entry is None:
        listing = ", ".join(sorted(entries, key=_id_sort))
        raise lang_error(
            InvalidInput,
            lang,
            f"Unknown NFD table {table_id!r}; table ids are {listing}.",
            f"tableau de la BNDF {table_id!r} inconnu ; les tableaux sont {listing}.",
        )
    return entry, cached


def _id_sort(table_id: str) -> tuple[int, ...]:
    return tuple(int(part) for part in table_id.split("."))


def _section(entry: CatalogueEntry, lang: Lang) -> str:
    return entry.section_fr if lang == "fr" else entry.section_en


def _title(entry: CatalogueEntry, lang: Lang) -> str:
    return entry.title_fr if lang == "fr" else entry.title_en


def _url(path: str | None) -> str | None:
    return _file_url(path) if path else None


def _summary(entry: CatalogueEntry, lang: Lang) -> NfdTableSummary:
    dictionary = entry.dictionary_path_fr if lang == "fr" else entry.dictionary_path_en
    return NfdTableSummary(
        table_id=entry.table_id,
        section=_section(entry, lang),
        title=_title(entry, lang),
        csv_url=_file_url(entry.csv_path),
        xlsx_url=_url(entry.xlsx_path),
        dictionary_url=_url(dictionary),
        comments_url=_url(entry.comments_path),
    )


def _provenance(
    url: str,
    cached: bool,
    schema: str,
    *,
    freshness: str | None = None,
    as_of: datetime | None = None,
    coverage: str | None = None,
    limits: str | None = None,
    lang: str = "en",
) -> Provenance:
    return make_provenance(
        source=constants.SOURCE_NAME,
        url=url,
        cached=cached,
        schema_name=schema,
        freshness=freshness
        or fr_or_en(
            lang,
            "Updated a few times a year by the NFD (data dictionaries dated April to June 2026).",
            "Mise à jour quelques fois par année par la BNDF (dictionnaires de données datés "
            "d'avril à juin 2026).",
        ),
        as_of=as_of,
        coverage=coverage,
        limits=limits,
        licence=_LICENCE_FR if lang == "fr" else _LICENCE,
        lang=lang,
    )


async def list_tables(section: str | None = None, lang: Lang = "en") -> NfdTableList:
    entries, cached = await _in_lang(_load_catalogue(), lang)
    chosen = sorted(entries.values(), key=lambda e: _id_sort(e.table_id))
    if section:
        wanted = fold(section)
        chosen = [
            e
            for e in chosen
            if wanted in fold(e.section_en)
            or wanted in fold(e.section_fr)
            or wanted in fold(e.title_en)
            or wanted in fold(e.title_fr)
        ]
        if not chosen:
            names = sorted({_section(e, lang) for e in entries.values()})
            raise lang_error(
                InvalidInput,
                lang,
                f"No NFD table matches {section!r}; sections are {names}.",
                f"aucun tableau de la BNDF ne correspond à {section!r} ; les sections sont "
                f"{names}.",
            )
    sections: list[str] = []
    for e in sorted(entries.values(), key=lambda e: _id_sort(e.table_id)):
        name = _section(e, lang)
        if name not in sections:
            sections.append(name)
    return NfdTableList(
        tables=[_summary(e, lang) for e in chosen],
        count=len(chosen),
        sections=sections,
        provenance=_provenance(
            constants.PAGE_FR if lang == "fr" else constants.PAGE_EN,
            cached,
            "nfd.NfdTableList",
            coverage=fr_or_en(
                lang,
                f"{len(entries)} tables on the NFD Download page",
                f"{len(entries)} tableaux sur la page Téléchargement de la BNDF",
            ),
            lang=lang,
        ),
    )


# ---------------------------------------------------------------- the table


@dataclass(frozen=True)
class DimensionSpec:
    key: str
    label_en: str
    label_fr: str


@dataclass(frozen=True)
class Record:
    year: int
    iso: str
    jurisdiction: tuple[str, str]
    labels: tuple[tuple[str, str], ...]
    value: float | None
    qualifier: str
    footnotes: tuple[str, ...]


@dataclass
class Dataset:
    url: str
    last_modified: str | None
    dimensions: list[DimensionSpec]
    value_label_en: str
    value_label_fr: str
    records: list[Record]
    unparsed_values: int = 0
    duplicate_groups: int = 0
    remapped_iso: set[str] = field(default_factory=set)
    mixed_case_labels: list[str] = field(default_factory=list)


def _layout(header: list[str]) -> tuple[int, int, int, int, int, int, int, int | None]:
    """Column positions: year, iso, jurisdiction EN, FR, first dimension, value, qualifier, FR value."""
    names = [h.strip() for h in header]
    if len(names) == 6 and "/" in names[0]:
        # Property losses: Year / Année, ISO, Jurisdiction, Juridiction, Dollars, qualifier.
        return 0, 1, 2, 3, 4, 4, 5, None
    if len(names) >= 9 and (len(names) - 9) % 2 == 0 and names[2].strip().lower() == "iso":
        n = len(names)
        return 0, 2, 3, 4, 5, n - 4, n - 3, n - 2
    raise UpstreamError(
        f"nfd table file: unexpected columns {names}; expected Year, Année, ISO, Jurisdiction, "
        "Juridiction, label pairs, then value and data-qualifier columns."
    )


def parse_table(text: str, url: str, last_modified: str | None) -> Dataset:
    reader = csv.reader(io.StringIO(text))
    try:
        header = next(reader)
    except StopIteration:
        raise UpstreamError(f"nfd table file {url} is empty.") from None
    year_i, iso_i, jen_i, jfr_i, dim_i, val_i, qual_i, valfr_i = _layout(header)
    labelled = valfr_i is not None
    dim_end = val_i if labelled else dim_i
    specs = [
        DimensionSpec(_key_from_label(header[i]), header[i].strip(), header[i + 1].strip())
        for i in range(dim_i, dim_end, 2)
    ]
    dataset = Dataset(
        url=url,
        last_modified=last_modified,
        dimensions=specs,
        value_label_en=header[val_i].strip(),
        value_label_fr=header[valfr_i].strip() if valfr_i is not None else header[val_i].strip(),
        records=[],
    )
    seen: set[tuple[Any, ...]] = set()
    repeated: set[tuple[Any, ...]] = set()
    case_variants: dict[str, set[str]] = {}
    for line_no, row in enumerate(reader, start=2):
        if not any(cell.strip() for cell in row):
            continue
        if len(row) != len(header):
            raise UpstreamError(f"nfd table file {url} line {line_no}: {len(row)} columns.")
        try:
            year = int(row[year_i].strip())
        except ValueError:
            raise UpstreamError(
                f"nfd table file {url} line {line_no}: year {row[year_i]!r}."
            ) from None
        iso = row[iso_i].strip().upper()
        if iso in constants.ISO_ALIASES:
            dataset.remapped_iso.add(iso)
            iso = constants.ISO_ALIASES[iso]
        labels: list[tuple[str, str]] = []
        marks: list[str] = []
        for i in range(dim_i, dim_end, 2):
            en, en_marks = clean_label(row[i])
            fr, fr_marks = clean_label(row[i + 1])
            labels.append((en, fr))
            marks.extend(m for m in en_marks + fr_marks if m not in marks)
        raw_value = row[val_i].strip().replace(",", "") if row[val_i].strip() else ""
        value: float | None = None
        if raw_value:
            try:
                value = float(raw_value)
            except ValueError:
                dataset.unparsed_values += 1
        key = (year, iso, tuple(en for en, _ in labels))
        if key in seen:
            repeated.add(key)
        seen.add(key)
        for spec, (en, _) in zip(specs, labels, strict=True):
            case_variants.setdefault(f"{spec.key}\0{fold(en)}", set()).add(en)
        dataset.records.append(
            Record(
                year=year,
                iso=iso,
                jurisdiction=(row[jen_i].strip(), row[jfr_i].strip()),
                labels=tuple(labels),
                value=value,
                qualifier=row[qual_i].strip(),
                footnotes=tuple(sorted(set(marks))),
            )
        )
    if not dataset.records:
        raise UpstreamError(f"nfd table file {url} has a header and no rows.")
    dataset.duplicate_groups = len(repeated)
    dataset.mixed_case_labels = sorted(
        {" / ".join(sorted(v)) for v in case_variants.values() if len(v) > 1}
    )
    return dataset


async def _load_dataset(entry: CatalogueEntry) -> tuple[Dataset, bool]:
    url = _file_url(entry.csv_path)

    async def fetch() -> Dataset:
        response = await _download(url, f"nfd table {entry.table_id}")
        text = decode(response.content)
        if text.lstrip().lower().startswith(("<!doctype", "<html")):
            raise NotFound(f"nfd table {entry.table_id}: {url} returned a web page, not a CSV.")
        return parse_table(text, url, response.headers.get("last-modified"))

    return await cached_fetch(f"nfd:table:{url}", constants.TABLE_TTL_SECONDS, fetch)


def _unit(dataset: Dataset, lang: Lang) -> str | None:
    header = dataset.value_label_fr if lang == "fr" else dataset.value_label_en
    folded = fold(dataset.value_label_en)
    if folded in _UNIT_BY_HEADER:
        return _UNIT_BY_HEADER[folded][1 if lang == "fr" else 0]
    match = re.search(r"\(([^)]*)\)", dataset.value_label_en)
    if match is None:
        return None
    english = match.group(1).strip()
    if lang == "en":
        return english
    fr_match = re.search(r"\(([^)]*)\)", header)
    return _UNIT_FR.get(english, fr_match.group(1).strip() if fr_match else english)


def _has_unit_dimension(dataset: Dataset) -> bool:
    return any(spec.key == "unit_of_measure" for spec in dataset.dimensions)


def _label(pair: tuple[str, str], lang: Lang) -> str:
    return (pair[1] or pair[0]) if lang == "fr" else pair[0]


def _jurisdiction_name(record: Record, lang: Lang) -> str:
    return (
        (record.jurisdiction[1] or record.jurisdiction[0])
        if lang == "fr"
        else record.jurisdiction[0]
    )


# ---------------------------------------------------------------- dictionary


async def _load_dictionary(entry: CatalogueEntry, lang: Lang) -> dict[str, Any] | None:
    """Title, last update and source from the data dictionary; None when it is not readable."""
    path = entry.dictionary_path_fr if lang == "fr" else entry.dictionary_path_en
    if not path:
        return None
    url = _file_url(path)

    async def fetch() -> dict[str, Any] | None:
        try:
            response = await _download(url, f"nfd dictionary {entry.table_id}")
        except (NotFound, UpstreamError, UpstreamUnavailable):
            return None
        # openpyxl takes ~0.7 s to import, so it loads on first use.
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(io.BytesIO(response.content), read_only=True, data_only=True)
        except Exception:  # noqa: BLE001 - a corrupt dictionary must not fail the table
            return None
        try:
            sheet = workbook.worksheets[1 if lang == "fr" and len(workbook.worksheets) > 1 else 0]
            top = [row for row in sheet.iter_rows(min_row=1, max_row=6, values_only=True)]
        finally:
            workbook.close()
        cells = {
            fold(str(row[0])): row[1] for row in top if row and row[0] is not None and len(row) > 1
        }
        title = next((v for k, v in cells.items() if k in ("title", "titre")), None)
        updated = next(
            (v for k, v in cells.items() if k.startswith(("last update", "derniere"))), None
        )
        source = next((v for k, v in cells.items() if k.startswith("source")), None)
        return {
            "title": " ".join(str(title).split()) if title else None,
            "last_updated": updated.date() if isinstance(updated, datetime) else None,
            "source": " ".join(str(source).split()) if source else None,
            "url": url,
        }

    data, _ = await cached_fetch(f"nfd:dictionary:{url}", constants.TABLE_TTL_SECONDS, fetch)
    return data


# -------------------------------------------------------------------- describe


def _dimension_values(dataset: Dataset, index: int, lang: Lang) -> list[str]:
    return sorted({_label(r.labels[index], lang) for r in dataset.records}, key=fold)


async def describe_table(table_id: str, lang: Lang = "en") -> NfdTableDescription:
    entry, cached_catalogue = await _entry(table_id, lang)
    dataset, cached = await _in_lang(_load_dataset(entry), lang)
    dictionary = await _in_lang(_load_dictionary(entry, lang), lang)
    records = dataset.records
    years = [r.year for r in records]
    spans: dict[str, tuple[Record, int, int]] = {}
    for r in records:
        found = spans.get(r.iso)
        spans[r.iso] = (
            (found[0], min(found[1], r.year), max(found[2], r.year))
            if found
            else (r, r.year, r.year)
        )
    jurisdictions = [
        NfdJurisdiction(
            iso=iso,
            name=_jurisdiction_name(first, lang),
            first_year=low,
            last_year=high,
        )
        for iso, (first, low, high) in sorted(spans.items())
    ]
    dimensions = []
    for index, spec in enumerate(dataset.dimensions):
        values = _dimension_values(dataset, index, lang)
        dimensions.append(
            NfdDimension(
                key=spec.key,
                label=spec.label_fr if lang == "fr" else spec.label_en,
                values=values[: constants.VALUES_LISTED_MAX],
                n_values=len(values),
                values_truncated=len(values) > constants.VALUES_LISTED_MAX,
            )
        )
    legend = constants.QUALIFIERS_FR if lang == "fr" else constants.QUALIFIERS_EN
    present = {r.qualifier for r in records if r.qualifier}
    quirks = _quirks(dataset, lang)
    return NfdTableDescription(
        table_id=entry.table_id,
        section=_section(entry, lang),
        title=_title(entry, lang),
        dictionary_title=dictionary["title"] if dictionary else None,
        source=dictionary["source"] if dictionary else None,
        last_updated=dictionary["last_updated"] if dictionary else None,
        value_label=dataset.value_label_fr if lang == "fr" else dataset.value_label_en,
        unit=_unit(dataset, lang),
        first_year=min(years),
        last_year=max(years),
        n_rows=len(records),
        jurisdictions=jurisdictions,
        dimensions=dimensions,
        qualifiers={code: legend[code] for code in sorted(present) if code in legend},
        quirks=quirks,
        csv_url=dataset.url,
        dictionary_url=dictionary["url"] if dictionary else None,
        comments_url=_url(entry.comments_path),
        provenance=_provenance(
            dataset.url,
            cached and cached_catalogue,
            "nfd.NfdTableDescription",
            as_of=(
                datetime.combine(dictionary["last_updated"], datetime.min.time())
                if dictionary and dictionary["last_updated"]
                else None
            ),
            coverage=fr_or_en(
                lang,
                f"{len(records)} rows, {min(years)} to {max(years)}",
                f"{len(records)} lignes, de {min(years)} à {max(years)}",
            ),
            lang=lang,
        ),
    )


def _quirks(dataset: Dataset, lang: str = "en") -> list[str]:
    quirks: list[str] = []
    if dataset.duplicate_groups:
        quirks.append(
            fr_or_en(
                lang,
                f"{dataset.duplicate_groups} combinations of year, jurisdiction and categories "
                "occur on more than one row; they are kept as published, so sum with care.",
                f"{dataset.duplicate_groups} combinaisons d'année, d'administration et de "
                "catégories figurent sur plus d'une ligne ; elles sont gardées telles que "
                "publiées : additionnez avec prudence.",
            )
        )
    if dataset.remapped_iso:
        quirks.append(
            fr_or_en(
                lang,
                "The file codes Yukon as 'YK' on some rows; shown as 'YT'.",
                "Le fichier code le Yukon « YK » sur certaines lignes ; affiché « YT ».",
            )
        )
    isos = {r.iso for r in dataset.records}
    if "GC" in isos:
        quirks.append(
            fr_or_en(
                lang,
                "Jurisdiction 'GC' is the Government of Canada (federal), not a total.",
                "L'administration « GC » est le gouvernement du Canada (fédéral), et non un total.",
            )
        )
    if "NP" in isos:
        quirks.append(
            fr_or_en(
                lang,
                "Jurisdiction 'NP' is national parks, not a province.",
                "L'administration « NP » désigne les parcs nationaux, et non une province.",
            )
        )
    if any(r.footnotes for r in dataset.records):
        quirks.append(
            fr_or_en(
                lang,
                "Some labels carry footnote markers (shown in `footnotes`); the text is in "
                "nfd_table_comments.",
                "Certains libellés portent des appels de note (dans `footnotes`) ; le texte est "
                "dans nfd_table_comments.",
            )
        )
    if dataset.mixed_case_labels:
        spellings = "; ".join(dataset.mixed_case_labels[:6])
        quirks.append(
            fr_or_en(
                lang,
                f"Labels spelled in more than one way in the file: {spellings}.",
                f"Libellés écrits de plus d'une façon dans le fichier : {spellings}.",
            )
        )
    if any(r.value is None for r in dataset.records):
        quirks.append(
            fr_or_en(
                lang,
                "Blank values have no figure (usually qualifier u, U or n, but a few carry "
                "another code such as a); they are returned as null, not 0.",
                "Les valeurs vides n'ont pas de chiffre (le plus souvent qualificatif u, U ou n, "
                "parfois un autre code comme a) ; elles sont renvoyées comme null, et non 0.",
            )
        )
    if _has_unit_dimension(dataset):
        quirks.append(
            fr_or_en(
                lang,
                "Rows mix rates and totals: filter or group_by `unit_of_measure` before summing.",
                "Les lignes mêlent taux et totaux : filtrez `unit_of_measure` ou mettez-le dans "
                "group_by avant d'additionner.",
            )
        )
    if dataset.unparsed_values:
        quirks.append(
            fr_or_en(
                lang,
                f"{dataset.unparsed_values} values were not numbers and are returned as null.",
                f"{dataset.unparsed_values} valeurs n'étaient pas des nombres et sont renvoyées "
                "comme null.",
            )
        )
    return quirks


# ----------------------------------------------------------------------- query


def _resolve_dimension(dataset: Dataset, name: str, lang: str = "en") -> int:
    wanted = fold(name)
    for index, spec in enumerate(dataset.dimensions):
        if (
            wanted in (fold(spec.key), fold(spec.label_en), fold(spec.label_fr))
            or wanted == spec.key
        ):
            return index
    names = [spec.key for spec in dataset.dimensions]
    raise lang_error(
        InvalidInput,
        lang,
        f"Unknown column {name!r}; this table's columns are {names}.",
        f"colonne {name!r} inconnue ; les colonnes de ce tableau sont {names}.",
    )


def _match_jurisdictions(dataset: Dataset, wanted: list[str], lang: str = "en") -> set[str]:
    by_name: dict[str, str] = {}
    for r in dataset.records:
        for name in (r.iso, r.jurisdiction[0], r.jurisdiction[1]):
            by_name[fold(name)] = r.iso
    chosen: set[str] = set()
    for value in wanted:
        code = value.strip().upper()
        iso = by_name.get(fold(constants.ISO_ALIASES.get(code, code))) or by_name.get(fold(value))
        if iso is None:
            names = sorted({f"{r.iso} ({r.jurisdiction[0]})" for r in dataset.records})
            raise lang_error(
                InvalidInput,
                lang,
                f"No jurisdiction {value!r} in this table; jurisdictions are {names}.",
                f"aucune administration {value!r} dans ce tableau ; les administrations sont "
                f"{names}.",
            )
        chosen.add(iso)
    return chosen


def _matching_labels(
    dataset: Dataset, index: int, wanted: list[str], spec: DimensionSpec, lang: str = "en"
) -> set[str]:
    """English labels whose English or French spelling matches one of `wanted`."""
    variants: dict[str, set[str]] = {}
    for r in dataset.records:
        en, fr = r.labels[index]
        for label in (en, fr):
            if label:
                variants.setdefault(fold(label), set()).add(en)
    chosen: set[str] = set()
    for value in wanted:
        found = variants.get(fold(value))
        if not found:
            close = _suggest(value, {en for pair in variants.values() for en in pair}, lang)
            options = sorted({en for pair in variants.values() for en in pair}, key=fold)
            listed = options[: constants.VALUES_LISTED_MAX]
            raise lang_error(
                InvalidInput,
                lang,
                f"No {spec.key} {value!r} in this table"
                + (close or f"; values are {listed}")
                + ".",
                f"aucune valeur {value!r} pour {spec.key} dans ce tableau"
                + (close or f" ; valeurs : {listed}")
                + ".",
            )
        chosen |= found
    return chosen


def _group_keys(dataset: Dataset, group_by: list[str], lang: str = "en") -> list[str]:
    keys: list[str] = []
    for name in group_by:
        folded = fold(name)
        if folded in ("year", "annee"):
            key = "year"
        elif folded in ("jurisdiction", "juridiction", "province", "iso"):
            key = "jurisdiction"
        else:
            key = dataset.dimensions[_resolve_dimension(dataset, name, lang)].key
        if key not in keys:
            keys.append(key)
    return keys


async def query_table(
    table_id: str,
    province: Filter = None,
    year_from: int | None = None,
    year_to: int | None = None,
    filters: Mapping[str, Filter] | None = None,
    group_by: list[str] | None = None,
    drop_missing: bool = False,
    limit: int = constants.ROWS_DEFAULT,
    lang: Lang = "en",
) -> NfdQueryResult:
    if not 1 <= limit <= constants.ROWS_MAX:
        raise lang_error(
            InvalidInput,
            lang,
            f"limit must be between 1 and {constants.ROWS_MAX}.",
            f"limit doit être compris entre 1 et {constants.ROWS_MAX}.",
        )
    if year_from is not None and year_to is not None and year_from > year_to:
        raise lang_error(
            InvalidInput,
            lang,
            "year_from is after year_to.",
            "year_from est postérieur à year_to.",
        )
    entry, cached_catalogue = await _entry(table_id, lang)
    dataset, cached = await _in_lang(_load_dataset(entry), lang)
    records = dataset.records
    notes: list[str] = []

    wanted_places = _as_list(province)
    if wanted_places:
        isos = _match_jurisdictions(dataset, wanted_places, lang)
        records = [r for r in records if r.iso in isos]
    if year_from is not None:
        records = [r for r in records if r.year >= year_from]
    if year_to is not None:
        records = [r for r in records if r.year <= year_to]
    chosen_by_index: dict[int, set[str]] = {}
    for name, value in (filters or {}).items():
        index = _resolve_dimension(dataset, name, lang)
        spec = dataset.dimensions[index]
        values = _as_list(value)
        if not values:
            continue
        chosen_by_index[index] = _matching_labels(dataset, index, values, spec, lang)
    for index, allowed in chosen_by_index.items():
        records = [r for r in records if r.labels[index][0] in allowed]
    if drop_missing:
        records = [r for r in records if r.value is not None]
    legend = constants.QUALIFIERS_FR if lang == "fr" else constants.QUALIFIERS_EN
    unit = _unit(dataset, lang)

    keys = _group_keys(dataset, group_by, lang) if group_by else []
    if keys and _has_unit_dimension(dataset):
        index = next(i for i, s in enumerate(dataset.dimensions) if s.key == "unit_of_measure")
        single = len({r.labels[index][0] for r in records}) <= 1
        if "unit_of_measure" not in keys and not single:
            raise lang_error(
                InvalidInput,
                lang,
                "This table mixes rates and totals; filter `unit_of_measure` to one value or "
                "include it in group_by before summing.",
                "ce tableau mêle taux et totaux ; limitez `unit_of_measure` à une valeur ou "
                "mettez-le dans group_by avant d'additionner.",
            )

    rows = (
        _aggregate(dataset, records, keys, lang, unit)
        if keys
        else _plain(dataset, records, lang, unit)
    )
    matched = len(rows)
    if keys:
        notes.append(
            fr_or_en(
                lang,
                "Values are summed over source rows; rows with no figure add nothing "
                "(`n_missing` counts them). Sums mix actual and estimated figures (`qualifiers`).",
                "Les valeurs sont additionnées sur les lignes de la source ; les lignes sans "
                "chiffre n'ajoutent rien (`n_missing` les compte). Les sommes mêlent chiffres "
                "réels et estimés (`qualifiers`).",
            )
        )
        if "jurisdiction" not in keys:
            odd = sorted({r.iso for r in records if r.iso in ("GC", "NP")})
            if odd:
                notes.append(
                    fr_or_en(
                        lang,
                        f"The sums include non-provincial jurisdictions {odd}.",
                        f"Les sommes comprennent des administrations non provinciales {odd}.",
                    )
                )
    if dataset.duplicate_groups:
        notes.append(
            fr_or_en(
                lang,
                f"{dataset.duplicate_groups} year/jurisdiction/category combinations repeat in "
                "the file and are kept as published.",
                f"{dataset.duplicate_groups} combinaisons année/administration/catégorie se "
                "répètent dans le fichier et sont gardées telles que publiées.",
            )
        )
    # Rows run oldest year first; when the cap cuts, keep the most recent
    # years (the file starts in 1940 for some tables, so the first rows are
    # rarely the ones wanted), still listed oldest first.
    if matched > limit:
        newest_first = sorted(range(matched), key=lambda i: -(rows[i].year or 0))
        kept = sorted(newest_first[:limit])
        shown = [rows[i] for i in kept]
    else:
        shown = rows
    used = {q for row in shown for q in row.qualifiers}
    limits = truncation_note_lang(
        lang,
        returned=len(shown),
        total=matched,
        unit="rows",
        order="latest",
        how_to_get_more=f"narrow year_from/year_to or filters, or raise limit (max {constants.ROWS_MAX})",
        how_to_get_more_fr="resserrez year_from/year_to ou les filtres, ou augmentez limit "
        f"(max. {constants.ROWS_MAX})",
    )
    return NfdQueryResult(
        table_id=entry.table_id,
        title=_title(entry, lang),
        unit=unit,
        rows=shown,
        returned_count=len(shown),
        matched_count=matched,
        limit=limit,
        group_by=keys,
        qualifiers={c: legend[c] for c in sorted(used) if c in legend},
        notes=notes,
        provenance=_provenance(
            dataset.url,
            cached and cached_catalogue,
            "nfd.NfdQueryResult",
            freshness=fr_or_en(
                lang,
                "Updated a few times a year by the NFD"
                + (
                    f" (file Last-Modified: {dataset.last_modified})"
                    if dataset.last_modified
                    else ""
                )
                + ".",
                "Mise à jour quelques fois par année par la BNDF"
                + (
                    f" (Last-Modified du fichier : {dataset.last_modified})"
                    if dataset.last_modified
                    else ""
                )
                + ".",
            ),
            limits=limits,
            lang=lang,
        ),
    )


def _row_unit(dataset: Dataset, record: Record, lang: Lang, default: str | None) -> str | None:
    if default is not None or not _has_unit_dimension(dataset):
        return default
    index = next(i for i, s in enumerate(dataset.dimensions) if s.key == "unit_of_measure")
    return _label(record.labels[index], lang)


def _plain(dataset: Dataset, records: list[Record], lang: Lang, unit: str | None) -> list[NfdRow]:
    ordered = sorted(records, key=lambda r: (r.year, r.iso, tuple(fold(en) for en, _ in r.labels)))
    return [
        NfdRow(
            year=r.year,
            iso=r.iso,
            jurisdiction=_jurisdiction_name(r, lang),
            dimensions={
                spec.key: _label(pair, lang)
                for spec, pair in zip(dataset.dimensions, r.labels, strict=True)
            },
            value=r.value,
            unit=_row_unit(dataset, r, lang, unit),
            qualifiers=[r.qualifier] if r.qualifier else [],
            n_rows=1,
            n_missing=int(r.value is None),
            footnotes=list(r.footnotes),
        )
        for r in ordered
    ]


def _aggregate(
    dataset: Dataset, records: list[Record], keys: list[str], lang: Lang, unit: str | None
) -> list[NfdRow]:
    positions = {spec.key: i for i, spec in enumerate(dataset.dimensions)}
    groups: dict[tuple[Any, ...], list[Record]] = {}
    for r in records:
        parts: list[Any] = []
        for key in keys:
            if key == "year":
                parts.append(r.year)
            elif key == "jurisdiction":
                parts.append(r.iso)
            else:
                parts.append(r.labels[positions[key]][0])
        groups.setdefault(tuple(parts), []).append(r)
    rows: list[NfdRow] = []
    for group_key, members in groups.items():
        values = [m.value for m in members if m.value is not None]
        first = members[0]
        by_key = dict(zip(keys, group_key, strict=True))
        # A group from one jurisdiction (province="BC" with group_by=["year"])
        # keeps its name; it used to come back null.
        one_place = "jurisdiction" in by_key or len({m.iso for m in members}) == 1
        rows.append(
            NfdRow(
                year=by_key.get("year"),
                iso=first.iso if one_place else None,
                jurisdiction=_jurisdiction_name(first, lang) if one_place else None,
                dimensions={
                    key: _label(first.labels[positions[key]], lang)
                    for key in keys
                    if key not in ("year", "jurisdiction")
                },
                value=sum(values) if values else None,
                unit=_row_unit(dataset, first, lang, unit),
                qualifiers=sorted({m.qualifier for m in members if m.qualifier}),
                n_rows=len(members),
                n_missing=len(members) - len(values),
                footnotes=sorted({f for m in members for f in m.footnotes}),
            )
        )
    rows.sort(
        key=lambda r: (
            r.year if r.year is not None else 0,
            r.iso or "",
            tuple(fold(v) for v in r.dimensions.values()),
        )
    )
    return rows


# -------------------------------------------------------------------- comments


async def table_comments(
    table_id: str,
    province: Filter = None,
    year_from: int | None = None,
    year_to: int | None = None,
    limit: int = 50,
    lang: Lang = "en",
) -> NfdCommentsResult:
    if limit < 1:
        raise lang_error(
            InvalidInput, lang, "limit must be at least 1.", "limit doit être d'au moins 1."
        )
    entry, cached_catalogue = await _entry(table_id, lang)
    title = _title(entry, lang)
    if not entry.comments_path:
        return NfdCommentsResult(
            table_id=entry.table_id,
            title=title,
            comments=[],
            returned_count=0,
            matched_count=0,
            notes=[
                fr_or_en(
                    lang,
                    "The NFD lists no comments workbook for this table.",
                    "La BNDF ne donne aucun classeur de commentaires pour ce tableau.",
                )
            ],
            provenance=_provenance(
                constants.PAGE_EN, cached_catalogue, "nfd.NfdCommentsResult", lang=lang
            ),
        )
    url = _file_url(entry.comments_path)

    async def fetch() -> list[list[list[str]]] | None:
        try:
            response = await _download(url, f"nfd comments {entry.table_id}")
        except NotFound:
            return None
        import xlrd  # imported on first use, like openpyxl

        try:
            book = xlrd.open_workbook(file_contents=response.content)
        except Exception as exc:
            raise UpstreamError(
                f"nfd comments {entry.table_id}: {url} is not a readable .xls."
            ) from exc
        sheets = []
        for sheet in book.sheets():
            rows = [
                [str(v).strip() if v != "" else "" for v in sheet.row_values(i)]
                for i in range(sheet.nrows)
            ]
            sheets.append([[sheet.name]] + rows)
        return sheets

    sheets, cached = await _in_lang(
        cached_fetch(f"nfd:comments:{url}", constants.TABLE_TTL_SECONDS, fetch), lang
    )
    provenance_url = url
    if sheets is None:
        return NfdCommentsResult(
            table_id=entry.table_id,
            title=title,
            comments=[],
            returned_count=0,
            matched_count=0,
            notes=[
                fr_or_en(
                    lang,
                    "The comments file listed for this table answers HTTP 404: no comments "
                    "published.",
                    "Le fichier de commentaires donné pour ce tableau répond HTTP 404 : aucun "
                    "commentaire publié.",
                )
            ],
            provenance=_provenance(provenance_url, cached, "nfd.NfdCommentsResult", lang=lang),
        )
    suffix = "_" + _SHEET_LANG[lang]
    chosen = next((s for s in sheets if s[0][0].upper().endswith(suffix)), sheets[0])
    rows = chosen[1:]
    if not rows:
        raise lang_error(
            UpstreamError,
            lang,
            f"nfd comments {entry.table_id}: sheet {chosen[0][0]!r} is empty.",
            f"nfd, commentaires {entry.table_id} : la feuille {chosen[0][0]!r} est vide.",
        )
    header = [fold(h) for h in rows[0]]

    def find(*names: str) -> int:
        return next((i for i, h in enumerate(header) if h in names), -1)

    iso_i = find("iso")
    name_i = find("jurisdiction", "juridiction")
    year_i = find("year", "annee")
    comment_i = find("comment", "commentaire")
    foot_i = find("footnote", "footnotes", "renvois")
    if min(iso_i, name_i, year_i, comment_i) < 0:
        raise lang_error(
            UpstreamError,
            lang,
            f"nfd comments {entry.table_id}: unexpected columns {rows[0]}.",
            f"nfd, commentaires {entry.table_id} : colonnes inattendues {rows[0]}.",
        )
    known = {iso_i, name_i, year_i, comment_i, foot_i}
    wanted_places = {fold(v) for v in _as_list(province)}
    grouped: dict[tuple[str, str, str, str, str], list[int]] = {}
    for row in rows[1:]:
        if len(row) <= max(known):
            continue
        try:
            year = int(float(row[year_i]))
        except ValueError:
            continue
        code = row[iso_i].upper()
        iso = constants.ISO_ALIASES.get(code) or code
        if wanted_places and not ({fold(iso), fold(row[name_i])} & wanted_places):
            continue
        if year_from is not None and year < year_from:
            continue
        if year_to is not None and year > year_to:
            continue
        context = "; ".join(
            f"{rows[0][i]}: {row[i]}" for i in range(len(row)) if i not in known and row[i]
        )
        foot = row[foot_i] if foot_i >= 0 else ""
        grouped.setdefault((iso, row[name_i], row[comment_i], foot, context), []).append(year)
    comments = [
        NfdComment(
            iso=iso,
            jurisdiction=name,
            years=sorted(set(years)),
            comment=comment,
            footnotes=foot or None,
            context=context or None,
        )
        for (iso, name, comment, foot, context), years in sorted(
            grouped.items(), key=lambda kv: (kv[0][0], min(kv[1]), kv[0][2])
        )
        if comment or foot
    ]
    return NfdCommentsResult(
        table_id=entry.table_id,
        title=title,
        comments=comments[:limit],
        returned_count=min(limit, len(comments)),
        matched_count=len(comments),
        notes=[
            fr_or_en(
                lang,
                "Identical comments are merged and list every year they apply to. Footnote "
                "letters match the markers on labels (see `footnotes` in nfd_query_table rows).",
                "Les commentaires identiques sont fusionnés et donnent toutes les années "
                "visées. Les lettres de renvoi correspondent aux appels de note des libellés "
                "(voir `footnotes` dans les lignes de nfd_query_table).",
            )
        ],
        provenance=_provenance(
            provenance_url,
            cached and cached_catalogue,
            "nfd.NfdCommentsResult",
            limits=fr_or_en(
                lang,
                f"comments capped at {limit}",
                f"commentaires plafonnés à {limit}",
            )
            if len(comments) > limit
            else None,
            lang=lang,
        ),
    )
