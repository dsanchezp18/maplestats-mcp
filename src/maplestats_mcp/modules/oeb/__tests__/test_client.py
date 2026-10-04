"""Mocked tests for the OEB client, shaped on pages and files read live 2026-10-03."""

from __future__ import annotations

import io
import re

import pytest
from openpyxl import Workbook

from maplestats_mcp.modules.oeb import client, constants, pages, records, tools
from maplestats_mcp.shared import cache as cache_module
from maplestats_mcp.shared.errors import InvalidInput, NotFound, UpstreamError

LISTING = re.compile(r"https://www\.oeb\.ca/ontarios-energy-sector/open-data\?page=\d")
LISTING_FR = re.compile(
    r"https://www\.oeb\.ca/fr/secteur-de-lenergie-de-lontario/donnees-ouvertes.*"
)
REL_SLUG = "electricity-reporting-record-keeping-requirements-rrr-section-2142-system-reliability"
REL_URL = "https://www.oeb.ca/documents/opendata/rrr/2024-2/ED%202.1.4.2%20System%20Reliability%20Indicators.xml"
OLD_URL = (
    "https://www.oeb.ca/documents/opendata/rrr/2023/2.1.4.2%20System%20Reliability%20Indicators.xml"
)


@pytest.fixture(autouse=True)
def _clear_cache() -> None:
    cache_module._caches.clear()


def _row(slug: str, title: str, body: str, freq: str, prefix: str = "/open-data/") -> str:
    return (
        '<div class="views-row"><div class="views-field views-field-title"><span class="field-content">'
        f'<a href="{prefix}{slug}" hreflang="en">{title}</a></span></div>'
        f'<div class="views-field views-field-body"><div class="field-content"><p>{body}</p></div></div>'
        '<div class="views-field views-field-field-update-frequency"><span class="views-label">'
        f'Update Frequency: </span><span class="field-content">{freq}</span></div></div>'
    )


def _listing(rows: str, *, next_page: bool = False, yearbook: bool = True) -> str:
    year = (
        '<a href="/sites/default/files/yearbook-System-Reliability-2021.xlsx">System Reliability '
        "(xlsx)</a>"
        if yearbook
        else ""
    )
    pager = (
        '<li class="pager__item pager__item--next"><a href="?page=1" rel="next">Next</a></li>'
        if next_page
        else ""
    )
    return f"<html><body><main>{year}<div class='view-content'>{rows}</div><ul>{pager}</ul></main></body></html>"


# The live reliability page: current file, "Archive:" then <details> releases;
# the 2.1.7 page splits one link in two anchors, kept here.
DATASET_PAGE = """
<html><head><link rel="alternate" hreflang="fr" href="https://www.oeb.ca/fr/donnees-ouvertes/fiabilite" /></head>
<body><main><h1 class="page-title"><span>Electricity RRR: Section 2.1.4.2 System Reliability Indicators</span></h1>
<div class="field field--name-body"><div class="field__label">Description</div><div class="field__item"><p>Provides data on reliability.</p><p>Blanks populated with zeros.</p></div></div>
<div class="field field--name-field-wysiwyg-url"><div class="field__label">URL</div><div class="field__item">
<p>Table 1: the reliability table.<br><a href="https://www.oeb.ca/documents/opendata/rrr/2024-2/ED 2.1.4.2 System Reliability Indicators.xml">https://www.oeb.ca/documents/opendata/rrr/2024-2/ED 2.1.4.2 System Reliability Indic</a><a href="https://www.oeb.ca/documents/opendata/rrr/2024-2/ED 2.1.4.2 System Reliability Indicators.xml">ators.xml</a></p>
<p><strong>Archive:</strong></p>
<details><summary>Data published August 27, 2024</summary><div class="details-wrapper"><p><a href="https://www.oeb.ca/documents/opendata/rrr/2023/2.1.4.2 System Reliability Indicators.xml">x</a></p></div></details>
<p>Date published October 6, 2023 <a href="https://www.oeb.ca/documents/opendata/rrr/2.1.4.2 Old.xml">y</a></p>
</div></div>
<div class="field field--name-field-file-types"><div class="field__label">File Types</div><div class="field__items"><div class="field__item">XML</div></div></div>
<div class="field field--name-field-update-frequency"><div class="field__label">Update Frequency</div><div class="field__item">Annually</div></div>
<div class="field field--name-field-update-frequency-note field__item">Last updated: April 30, 2026</div>
</main></body></html>
"""

# Access export, as served: escaped names, an empty value left out entirely.
ACCESS_XML = b"""<?xml version="1.0" encoding="UTF-8"?>
<dataroot xmlns:od="urn:schemas-microsoft-com:officedata">
<ED_x0020_2142_x0020_System_x0020_Reliability_x0020_Indicators>
<Current_Company_Name>Hydro Ottawa Limited</Current_Company_Name>
<Historical_Company_Name>Hydro Ottawa Limited</Historical_Company_Name>
<Year>2024</Year>
<Cause_of_Interruption>Tree Contacts - Fallen Tree on Right-of-Way</Cause_of_Interruption>
<Total_SAIDI>0.031</Total_SAIDI>
<Salaries_x002C__Wages>1.187263E+08</Salaries_x002C__Wages>
</ED_x0020_2142_x0020_System_x0020_Reliability_x0020_Indicators>
<ED_x0020_2142_x0020_System_x0020_Reliability_x0020_Indicators>
<Current_Company_Name>Alectra Utilities Corporation</Current_Company_Name>
<Historical_Company_Name>PowerStream Inc.</Historical_Company_Name>
<Year>2016</Year>
<Cause_of_Interruption>Tree Contacts</Cause_of_Interruption>
<Total_SAIDI>0.81</Total_SAIDI>
</ED_x0020_2142_x0020_System_x0020_Reliability_x0020_Indicators>
</dataroot>"""

# The licences report: Oracle-style nesting, no Access dataroot.
NESTED_XML = b"""<?xml version="1.0" encoding="UTF-8"?><ALL_LICENCES><LIST_G_LICENCE_NUMBER>
<G_LICENCE_NUMBER><LICENCE_NUMBER>ED-2016-0360</LICENCE_NUMBER><LICENCE_NAME>Alectra Utilities Corporation</LICENCE_NAME><LICENCE_TYPE>Electricity Distributor</LICENCE_TYPE></G_LICENCE_NUMBER>
<G_LICENCE_NUMBER><LICENCE_NUMBER>EB-1</LICENCE_NUMBER><LICENCE_NAME>Some Retailer</LICENCE_NAME><LICENCE_TYPE>Electricity Retailer</LICENCE_TYPE></G_LICENCE_NUMBER>
</LIST_G_LICENCE_NUMBER></ALL_LICENCES>"""


def _xlsx(rows: list[list[object]], title: str = "Sheet 1") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = title
    for row in rows:
        sheet.append(row)
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def _mock_catalogue(httpx_mock) -> None:
    httpx_mock.add_response(
        url=LISTING,
        text=_listing(
            _row(
                REL_SLUG,
                "Electricity RRR: Section 2.1.4.2 System Reliability Indicators",
                "Reliability.",
                "Annually",
            )
        ),
        is_reusable=True,
    )
    httpx_mock.add_response(
        url=constants.DATASET_URL.format(slug=REL_SLUG),
        text=DATASET_PAGE,
        is_reusable=True,
        is_optional=True,
    )


def test_label_decodes_access_escapes_and_namespaces():
    assert (
        records.label("Company_Name__x0028_Merged_or_Current_x0029_")
        == "Company Name (Merged or Current)"
    )
    assert records.label("{urn:schemas-microsoft-com:xml-analysis:rowset}Year") == "Year"
    assert records.label("ED_x0020_212_x0020_Customers_x0020__x0026__x0020_Connections") == (
        "ED 212 Customers & Connections"
    )


def test_to_cell_numbers_and_blanks():
    assert records.to_cell("1.187263E+08") == pytest.approx(118726300.0)
    assert records.to_cell("2024") == 2024
    assert records.to_cell("79.16%") == "79.16%"
    assert records.to_cell("  ") is None
    assert records.to_cell("EB-2015-0245\nDecember 17, 2019") == "EB-2015-0245 December 17, 2019"


def test_xml_records_handles_nested_reports_and_missing_elements():
    rows = [r for _, r in records.xml_records(ACCESS_XML)]
    assert len(rows) == 2
    assert "Salaries, Wages" in rows[0] and "Salaries, Wages" not in rows[1]
    nested = list(records.xml_records(NESTED_XML))
    assert [tag for tag, _ in nested] == ["G LICENCE NUMBER", "G LICENCE NUMBER"]


def test_xml_records_refuses_a_dtd():
    body = b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaa">]><x><r><f>&a;</f></r></x>'
    with pytest.raises(Exception, match="DTD"):
        list(records.xml_records(body))


def test_parse_dataset_page_releases_captions_and_split_links():
    page = pages.parse_dataset_page(DATASET_PAGE, REL_SLUG)
    assert [f.release for f in page.files] == ["current", "2024-08-27", "2023-10-06"]
    assert page.files[0].caption == "Table 1: the reliability table."
    assert page.last_updated is not None and page.last_updated.isoformat() == "2026-04-30"
    assert page.fr_path == "/fr/donnees-ouvertes/fiabilite"
    assert page.file_types == ["XML"]


def test_parse_hop_page_labels_bare_xml_links():
    html = (
        "<main><a href='https://www.oeb.ca/_html/costawards/costawards_case.php?yr=2019'>Intervenor "
        "Cost Award Report: by Case</a><a href='https://www.oeb.ca/_html/costawards/data/"
        "section30_cases_2019.xml'>.xml</a><a href='https://elsewhere.example/x.xml'>other</a></main>"
    )
    files = pages.parse_hop_page(html, "https://www.oeb.ca/x")
    assert [(f.name, f.caption) for f in files] == [
        ("section30_cases_2019.xml", "Intervenor Cost Award Report: by Case")
    ]


async def test_list_datasets_follows_pager_and_adds_yearbook(httpx_mock):
    httpx_mock.add_response(
        url=f"{constants.LISTING_URL}?page=0",
        text=_listing(_row(REL_SLUG, "Reliability", "Reliability.", "Annually"), next_page=True),
        is_reusable=True,
    )
    httpx_mock.add_response(
        url=f"{constants.LISTING_URL}?page=1",
        text=_listing(
            _row("applications-oeb", "Applications before the OEB", "Active applications.", "Daily")
        ),
    )
    httpx_mock.add_response(url=constants.DATASET_URL.format(slug=REL_SLUG), text=DATASET_PAGE)
    httpx_mock.add_response(
        url=constants.DATASET_URL.format(slug="applications-oeb"),
        text=DATASET_PAGE.replace("2.1.4.2", "apps").replace(
            "https://www.oeb.ca/fr/donnees-ouvertes/fiabilite", "https://www.oeb.ca/fr/x"
        ),
    )
    result = await tools.oeb_list_datasets()
    assert [d.slug for d in result.datasets] == [
        REL_SLUG,
        "applications-oeb",
        constants.YEARBOOK_SLUG,
    ]
    rel = result.datasets[0]
    assert (rel.current_files, rel.archived_files, rel.readable) == (1, 2, True)
    assert "Open Government Licence – Ontario" in (result.provenance.licence or "")
    filtered = await tools.oeb_list_datasets("applications")
    assert [d.slug for d in filtered.datasets] == ["applications-oeb"]


async def test_french_titles_come_from_the_french_listing(httpx_mock):
    _mock_catalogue(httpx_mock)
    httpx_mock.add_response(
        url=LISTING_FR,
        text=_listing(
            _row(
                "fiabilite",
                "Indicateurs de fiabilité du système",
                "Fiabilité.",
                "Chaque année",
                "/fr/donnees-ouvertes/",
            ),
            yearbook=False,
        ),
        is_reusable=True,
    )
    result = await tools.oeb_list_datasets("fiabilité", lang="fr")
    assert result.datasets[0].title == "Indicateurs de fiabilité du système"
    assert result.datasets[0].update_frequency == "Chaque année"
    assert "Licence du gouvernement ouvert" in (result.provenance.licence or "")


async def test_query_filters_distributor_year_and_fields(httpx_mock):
    _mock_catalogue(httpx_mock)
    httpx_mock.add_response(url=REL_URL, content=ACCESS_XML)
    result = await tools.oeb_query_dataset(
        "system reliability", distributor="hydro ottawa", year=2024, fields=["SAIDI"]
    )
    assert result.columns == [
        "Current Company Name",
        "Historical Company Name",
        "Year",
        "Total SAIDI",
    ]
    assert result.rows == [
        {
            "Current Company Name": "Hydro Ottawa Limited",
            "Historical Company Name": "Hydro Ottawa Limited",
            "Year": 2024,
            "Total SAIDI": 0.031,
        }
    ]
    assert result.matched_distributors == ["Hydro Ottawa Limited"]


async def test_query_matches_historical_names_and_where(httpx_mock):
    _mock_catalogue(httpx_mock)
    httpx_mock.add_response(url=REL_URL, content=ACCESS_XML)
    result = await tools.oeb_query_dataset(
        REL_SLUG, distributor="PowerStream", where={"cause of interruption": "tree contacts"}
    )
    assert result.total_matched == 1 and result.rows[0]["Year"] == 2016


async def test_query_an_archived_release(httpx_mock):
    _mock_catalogue(httpx_mock)
    httpx_mock.add_response(url=OLD_URL, content=ACCESS_XML)
    result = await tools.oeb_query_dataset(REL_SLUG, release="2024-08-27")
    assert result.file.release == "2024-08-27" and result.total_matched == 2


async def test_unknown_column_lists_the_columns(httpx_mock):
    _mock_catalogue(httpx_mock)
    httpx_mock.add_response(url=REL_URL, content=ACCESS_XML)
    with pytest.raises(InvalidInput, match="Total SAIDI"):
        await tools.oeb_query_dataset(REL_SLUG, where={"colour": "blue"})


async def test_describe_reports_fields_years_and_companies(httpx_mock):
    _mock_catalogue(httpx_mock)
    httpx_mock.add_response(url=REL_URL, content=ACCESS_XML)
    detail = await tools.oeb_describe_dataset(REL_SLUG)
    assert detail.record_count == 2
    assert detail.years == ["2016", "2024"]
    assert "PowerStream Inc." in detail.distributors
    saidi = next(f for f in detail.fields if f.name == "Total SAIDI")
    assert saidi.numeric and saidi.filled == 2
    wages = next(f for f in detail.fields if f.name == "Salaries, Wages")
    assert wages.filled == 1


async def test_an_html_error_page_instead_of_xml_is_an_upstream_error(httpx_mock):
    _mock_catalogue(httpx_mock)
    httpx_mock.add_response(url=REL_URL, content=b"\n\nSorry, maintenance")
    with pytest.raises(UpstreamError, match="not XML"):
        await tools.oeb_query_dataset(REL_SLUG)


async def test_unknown_dataset_is_not_found(httpx_mock):
    _mock_catalogue(httpx_mock)
    with pytest.raises(NotFound):
        await tools.oeb_describe_dataset("pipeline tolls")


async def test_rates_label_fields_from_the_data_keys(httpx_mock):
    spec = constants.RATE_TABLES["electricity_residential"]
    bill = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?><BillDataTable>
    <BillDataRow><Dist>Toronto Hydro-Electric System Limited</Dist><Class>RESIDENTIAL</Class><YEAR>2026</YEAR><SC>42.36</SC><GA_RR_NONRPP>0.0036</GA_RR_NONRPP><VC/></BillDataRow>
    <BillDataRow><Dist>Hydro Ottawa Limited</Dist><Class>RESIDENTIAL</Class><YEAR>2026</YEAR><SC>30.1</SC></BillDataRow>
    </BillDataTable>"""
    keys = _xlsx(
        [
            ["dataname", "description", "value/format"],
            ["<SC>", "monthly fixed charge", "cents/kWh"],
            ["<GA_RR_NONRPP>", "Global adjustment rate rider", "cents/kWh"],
        ],
        "Electricity",
    )
    httpx_mock.add_response(url=spec["url"], content=bill)
    httpx_mock.add_response(url=spec["keys"], content=keys)
    result = await tools.oeb_rates("electricity_residential", distributor="toronto hydro")
    assert result.total_matched == 1 and result.rows[0]["SC"] == 42.36
    fields = {f.name: f for f in result.fields}
    assert fields["GA RR NONRPP"].description == "Global adjustment rate rider"
    assert result.rows[0]["VC"] is None


async def test_rpp_history_reads_the_sheet_with_dates(httpx_mock):
    from datetime import date

    spec = constants.RATE_TABLES["rpp_time_of_use"]
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.title = "Time-of-Use"
    first.append(
        ["Effective date", "Off-Peak price \n(¢ per kWh)", "Mid-Peak price", "On-Peak price"]
    )
    first.append([date(2025, 11, 1), 9.8, 15.7, 20.3])
    workbook.create_sheet("Tiered").append(["Effective date", "Lower", "Threshold", "Higher"])
    out = io.BytesIO()
    workbook.save(out)
    httpx_mock.add_response(url=spec["url"], content=out.getvalue())
    result = await tools.oeb_rates("rpp_time_of_use")
    assert result.rows == [
        {
            "Effective date": "2025-11-01",
            "Off-Peak price (¢ per kWh)": 9.8,
            "Mid-Peak price": 15.7,
            "On-Peak price": 20.3,
        }
    ]
    with pytest.raises(InvalidInput):
        await tools.oeb_rates("rpp_tiered", distributor="Toronto Hydro")


async def test_xlsx_header_below_title_rows(httpx_mock):
    body = _xlsx(
        [
            ["2025 Rates Database"],
            ["This database does not include regulatory charges."],
            [],
            ["Company Name", "Effective_date", "Service Classification", "RateRider_Value"],
            ["Hydro One Networks Inc.", "2025-01-01", "RESIDENTIAL", "0.22"],
        ],
        "Export",
    )
    rows = list(records.xlsx_records(body, "Export"))
    assert rows == [
        {
            "Company Name": "Hydro One Networks Inc.",
            "Effective_date": "2025-01-01",
            "Service Classification": "RESIDENTIAL",
            "RateRider_Value": 0.22,
        }
    ]


def test_select_file_by_year_in_name_not_position():
    files = client._files(
        pages.DatasetPage(
            slug="s",
            title="t",
            description="",
            file_types=[],
            update_frequency=None,
            last_updated=None,
            fr_path=None,
            files=[
                pages.PageFile(
                    url=f"https://www.oeb.ca/{y}-Rates.xlsx",
                    name=f"{y}-Rates.xlsx",
                    caption=None,
                    release="current",
                    format="xlsx",
                )
                for y in (2025, 2024)
            ],
        )
    )
    assert client._select_file(files, "2024", None, "en").name == "2024-Rates.xlsx"
    assert client._select_file(files, "2", None, "en").name == "2024-Rates.xlsx"


def test_docstrings_have_use_for_and_bilingual_keywords():
    for fn in (
        tools.oeb_list_datasets,
        tools.oeb_describe_dataset,
        tools.oeb_query_dataset,
        tools.oeb_rates,
    ):
        doc = fn.__doc__ or ""
        assert "Use for:" in doc and "Keywords:" in doc and "Mots-clés :" in doc
        english = doc.split("Keywords:")[1].split("Mots-clés :")[0]
        french = doc.split("Mots-clés :")[1]
        assert len([k for k in english.split(",") if k.strip()]) >= 8
        assert len([k for k in french.split(",") if k.strip()]) >= 8
