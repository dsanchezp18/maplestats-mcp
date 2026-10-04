"""Tests for modules/cihi/client.py, shaped on live 2026-09-23 pages and files."""

from __future__ import annotations

import io
import re

import pytest
from openpyxl import Workbook

from maplestats_mcp.modules.cihi import client, constants
from maplestats_mcp.shared import cache as cache_module
from maplestats_mcp.shared.errors import InvalidInput, NotFound


@pytest.fixture(autouse=True)
def _clear_cache():
    cache_module._caches.clear()
    yield


_SLUG = "30-day-stroke-in-hospital-mortality"
_EN_PAGE_URL = f"{constants.BASE_URL}{constants.INDICATOR_PATH}{_SLUG}"
_FR_PAGE_URL = f"{constants.BASE_URL}/fr/indicateurs/mortalite-avc"
_EN_FILE = (
    f"{constants.BASE_URL}/sites/default/files/document/data-file/955-{_SLUG}-data-table-en.xlsx"
)
_FR_FILE = _EN_FILE.replace("-en.xlsx", "-fr.xlsx")


def _page(file_url: str, name: str) -> str:
    return f"""<html><head><link hreflang="fr" href="{_FR_PAGE_URL}"></head><body><main>
<h1>{name}</h1>
<div class="view-metadata-summary"><div class="col-12"><div>Risk-adjusted rate of deaths.</div></div>
<ul class="indicator-meta-summary"><li>Data updated: <strong>September 2026</strong></li>
<li>Update frequency: <strong>Every year</strong></li></ul>
<div class="indicator-topics"><ul><li>Acute care </li><li>Health outcomes </li></ul></div></div>
<a href="{file_url}">download the data file</a></main></body></html>"""


def _workbook(header: list[str], rows: list[list[object]]) -> bytes:
    book = Workbook()
    book.create_sheet("Instructions", 0)
    book.remove(book.worksheets[1])
    sheet = book.create_sheet("Table 1")
    sheet.append(["Table 1 30-Day Stroke In-Hospital Mortality"])
    sheet.append([*header, None])
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


_HEADER = ["Reporting level", "Place or organization", "Time frame", "Risk-adjusted rate"]
_ROWS = [
    ["Province/territory", "Alberta", "2024–2025", 12.1],
    ["Province/territory", "Alberta", "2025–2026", 10.8],
    ["Facility", "University of Alberta Hospital (Alta.)", "2025–2026", 15.9],
    ["Province/territory", "Ontario", "2025–2026", 11.0],
]


async def test_search_pages_library_until_no_new_indicators(httpx_mock):
    page = '<main><a href="/en/indicators/30-day-stroke-in-hospital-mortality">30-Day Stroke In-Hospital Mortality</a><a href="/en/indicators/hospitalized-strokes">Hospitalized Strokes</a></main>'
    httpx_mock.add_response(url=f"{constants.LIBRARY_URL}?page=0", text=page)
    httpx_mock.add_response(url=f"{constants.LIBRARY_URL}?page=1", text=page)
    result = await client.search_indicators("stroke mortality")
    assert [r.slug for r in result.indicators] == [_SLUG]


async def test_get_indicator_reads_facts_and_data_file(httpx_mock):
    httpx_mock.add_response(url=_EN_PAGE_URL, text=_page(_EN_FILE, "30-Day Stroke"))
    detail = await client.get_indicator(_SLUG)
    assert detail.description == "Risk-adjusted rate of deaths."
    assert detail.facts["Data updated"] == "September 2026"
    assert detail.facts["Topics"] == "Acute care, Health outcomes"
    assert detail.data_file_url == _EN_FILE


async def test_get_indicator_data_filters_place_and_columns(httpx_mock):
    httpx_mock.add_response(url=_EN_PAGE_URL, text=_page(_EN_FILE, "30-Day Stroke"))
    httpx_mock.add_response(url=_EN_FILE, content=_workbook(_HEADER, _ROWS))
    data = await client.get_indicator_data(
        _SLUG, place="alberta", filters={"reporting level": "Province/territory"}, limit=1
    )
    assert data.tables == ["Table 1"]
    assert data.columns == _HEADER
    assert data.matching_rows == 2
    assert data.rows == [
        {
            "Reporting level": "Province/territory",
            "Place or organization": "Alberta",
            "Time frame": "2025–2026",
            "Risk-adjusted rate": "10.8",
        }
    ]


async def test_columns_pick_and_blank_columns_drop(httpx_mock):
    # Live rows have 33 columns, many blank for a given indicator.
    header = [*_HEADER, "Level 2 breakdown", "Comparison"]
    rows = [[*r, None, None] for r in _ROWS]
    httpx_mock.add_response(url=_EN_PAGE_URL, text=_page(_EN_FILE, "30-Day Stroke"))
    httpx_mock.add_response(url=_EN_FILE, content=_workbook(header, rows))
    data = await client.get_indicator_data(_SLUG, place="Ontario")
    assert data.columns == _HEADER
    assert data.empty_columns == ["Level 2 breakdown", "Comparison"]
    picked = await client.get_indicator_data(
        _SLUG, place="Ontario", columns=["time frame", "Risk-adjusted rate"]
    )
    assert picked.rows == [{"Time frame": "2025–2026", "Risk-adjusted rate": "11"}]
    with pytest.raises(InvalidInput, match="Unknown column"):
        await client.get_indicator_data(_SLUG, columns=["Nope"])


async def test_french_data_follows_hreflang(httpx_mock):
    httpx_mock.add_response(url=_EN_PAGE_URL, text=_page(_EN_FILE, "30-Day Stroke"))
    httpx_mock.add_response(url=_FR_PAGE_URL, text=_page(_FR_FILE, "Mortalité AVC"))
    french_header = ["Niveau de déclaration", "Lieu ou organisme", "Période", "Taux"]
    httpx_mock.add_response(url=_FR_FILE, content=_workbook(french_header, _ROWS))
    data = await client.get_indicator_data(_SLUG, place="Ontario", lang="fr")
    assert data.rows[0]["Lieu ou organisme"] == "Ontario"


async def test_input_validation(httpx_mock):
    with pytest.raises(InvalidInput):
        await client.get_indicator("../../etc")
    with pytest.raises(InvalidInput):
        await client.get_indicator_data(_SLUG, limit=0)
    httpx_mock.add_response(url=_EN_PAGE_URL, text=_page(_EN_FILE, "30-Day Stroke"))
    httpx_mock.add_response(url=_EN_FILE, content=_workbook(_HEADER, _ROWS))
    with pytest.raises(InvalidInput, match="Unknown column"):
        await client.get_indicator_data(_SLUG, filters={"Bogus": "x"})
    with pytest.raises(InvalidInput, match="Unknown table"):
        await client.get_indicator_data(_SLUG, table="Table 9")


async def test_missing_data_file_is_not_found(httpx_mock):
    httpx_mock.add_response(
        url=re.compile(re.escape(_EN_PAGE_URL)), text="<main><h1>No file</h1></main>"
    )
    with pytest.raises(NotFound, match="no downloadable data table"):
        await client.get_indicator_data(_SLUG)


# The French library and the sitemap that pairs it with the English one,
# shaped on live 2026-10-03 responses: a sitemap index (with a byte-order
# mark), then a child sitemap whose <url> entries carry both hreflang
# alternates.
_FR_SLUG = "mortalite-a-lhopital-dans-les-30-jours-accident-vasculaire-cerebral"
_FR_LIBRARY = f"""<main>
<h3 class="c-card__title"><a href="/fr/indicateurs/{_FR_SLUG}" hreflang="fr">Mortalité à
 l'hôpital dans les 30 jours, accident vasculaire cérébral</a></h3>
<h3 class="c-card__title"><a href="/fr/indicateurs/reserve-aux-francophones" hreflang="fr">Hôpitaux
 réservés</a></h3></main>"""
_EN_LIBRARY = f"""<main>
<h3><a href="/en/indicators/{_SLUG}" hreflang="en">30-Day Stroke In-Hospital Mortality</a></h3>
<h3><a href="/en/indicators/english-only-hospital-stays" hreflang="en">English-Only Hospital
 Stays</a></h3></main>"""
_SITEMAP_INDEX = """\N{ZERO WIDTH NO-BREAK SPACE}<?xml version="1.0" encoding="UTF-8"?>
<sitemapindex xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">
 <sitemap><loc>https://www.cihi.ca/sitemap.xml?page=1</loc></sitemap>
</sitemapindex>"""
_SITEMAP_PAGE = f"""<?xml version="1.0" encoding="UTF-8"?>
<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9"
 xmlns:xhtml="http://www.w3.org/1999/xhtml">
 <url>
  <loc>https://www.cihi.ca/en/indicators/{_SLUG}</loc>
  <xhtml:link rel="alternate" hreflang="en" href="https://www.cihi.ca/en/indicators/{_SLUG}"/>
  <xhtml:link rel="alternate" hreflang="fr" href="https://www.cihi.ca/fr/indicateurs/{_FR_SLUG}"/>
 </url>
 <url>
  <loc>https://www.cihi.ca/fr/indicateurs/{_FR_SLUG}</loc>
  <xhtml:link rel="alternate" hreflang="en" href="https://www.cihi.ca/en/indicators/{_SLUG}"/>
  <xhtml:link rel="alternate" hreflang="fr" href="https://www.cihi.ca/fr/indicateurs/{_FR_SLUG}"/>
 </url>
 <url>
  <loc>https://www.cihi.ca/en/about-cihi</loc>
  <xhtml:link rel="alternate" hreflang="en" href="https://www.cihi.ca/en/about-cihi"/>
  <xhtml:link rel="alternate" hreflang="fr" href="https://www.cihi.ca/fr/a-propos-de-licis"/>
 </url>
</urlset>"""


def _mock_pairing(httpx_mock) -> None:
    httpx_mock.add_response(url=constants.SITEMAP_URL, text=_SITEMAP_INDEX)
    httpx_mock.add_response(url=f"{constants.SITEMAP_URL}?page=1", text=_SITEMAP_PAGE)


async def test_french_search_reads_french_names_and_pairs_english_slugs(httpx_mock):
    for page in (0, 1):
        httpx_mock.add_response(url=f"{constants.FR_LIBRARY_URL}?page={page}", text=_FR_LIBRARY)
        httpx_mock.add_response(url=f"{constants.LIBRARY_URL}?page={page}", text=_EN_LIBRARY)
    _mock_pairing(httpx_mock)
    # Articles and the apostrophe are ignored; accents and case do not matter.
    result = await client.search_indicators("mortalite a l'HOPITAL accident vasculaire", "fr")
    assert [(r.slug, r.english_slug) for r in result.indicators] == [(_FR_SLUG, _SLUG)]
    assert result.indicators[0].name.startswith("Mortalité à l'hôpital")
    assert result.indicators[0].url.endswith(f"/fr/indicateurs/{_FR_SLUG}")
    # A plural matches the singular name; an indicator with no French page
    # is listed by its English name, and one with no English twin says so.
    plural = await client.search_indicators("hôpitaux", "fr")
    assert {r.slug for r in plural.indicators} == {_FR_SLUG, "reserve-aux-francophones"}
    english_only = await client.search_indicators("hospital stays", "fr")
    assert [r.english_slug for r in english_only.indicators] == ["english-only-hospital-stays"]
    assert "Aucune page en français" in (english_only.indicators[0].note or "")
    assert "1 indicateur(s) sans page en français" in (english_only.note or "")
    assert english_only.provenance.freshness == (
        "listes d'indicateurs mises en cache 7 jours, appariement anglais-français 1 jour"
    )
    assert (english_only.provenance.licence or "").startswith("Conditions d'utilisation de l'ICIS")
    french_only = await client.search_indicators("réservés", "fr")
    assert french_only.indicators[0].english_slug is None
    assert "aucune page en anglais" in (french_only.indicators[0].note or "")


async def test_french_slug_works_in_either_language(httpx_mock):
    french_url = f"{constants.BASE_URL}{constants.FR_INDICATOR_PATH}{_FR_SLUG}"
    httpx_mock.add_response(
        url=f"{constants.BASE_URL}{constants.INDICATOR_PATH}{_FR_SLUG}", status_code=404
    )
    _mock_pairing(httpx_mock)
    httpx_mock.add_response(url=_EN_PAGE_URL, text=_page(_EN_FILE, "30-Day Stroke"))
    httpx_mock.add_response(url=french_url, text=_page(_FR_FILE, "Mortalité AVC"))
    english = await client.get_indicator(_FR_SLUG)
    assert (english.slug, english.french_slug, english.name) == (_SLUG, _FR_SLUG, "30-Day Stroke")
    assert english.data_file_url == _EN_FILE
    french = await client.get_indicator(french_url, lang="fr")
    assert (french.slug, french.name, french.data_file_url) == (_SLUG, "Mortalité AVC", _FR_FILE)
    assert french.facts["Thèmes"] == "Acute care, Health outcomes"


async def test_unknown_slug_in_both_libraries_is_not_found(httpx_mock):
    for path in (constants.INDICATOR_PATH, constants.FR_INDICATOR_PATH):
        httpx_mock.add_response(url=f"{constants.BASE_URL}{path}nope-nope", status_code=404)
    _mock_pairing(httpx_mock)
    with pytest.raises(NotFound, match="English or French library"):
        await client.get_indicator("nope-nope")
    for path in (constants.INDICATOR_PATH, constants.FR_INDICATOR_PATH):
        httpx_mock.add_response(url=f"{constants.BASE_URL}{path}nope-nope", status_code=404)
    with pytest.raises(NotFound, match=r"^Aucune correspondance trouvée\xa0: cihi\xa0: aucun"):
        await client.get_indicator("nope-nope", lang="fr")
    with pytest.raises(InvalidInput, match="identifiant"):
        await client.get_indicator("Not a slug!", lang="fr")
