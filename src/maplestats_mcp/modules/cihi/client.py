"""Client for CIHI's Indicator Library pages and XLSX data tables.

Pages are parsed with BeautifulSoup; data tables with openpyxl in
read-only mode. Indicators are addressed by their English or French page
slug. The two libraries use different slugs; the site's sitemap pairs
them (each entry carries both hreflang alternates), and an English page's
own hreflang link leads to its French page and French file.
"""

from __future__ import annotations

import io
import re
import unicodedata
from typing import Any
from urllib.parse import urlparse

import httpx
from bs4 import BeautifulSoup, Tag
from defusedxml import ElementTree

from maplestats_mcp.modules.cihi import constants
from maplestats_mcp.modules.cihi.schemas import (
    IndicatorData,
    IndicatorDetail,
    IndicatorRef,
    IndicatorSearchResult,
)
from maplestats_mcp.shared.cache import cached_fetch
from maplestats_mcp.shared.envelope import make_provenance
from maplestats_mcp.shared.errors import InvalidInput, NotFound, UpstreamError, UpstreamUnavailable
from maplestats_mcp.shared.executor import run_parse
from maplestats_mcp.shared.fr_typography import fr_or_en, lang_error, truncation_note_lang
from maplestats_mcp.shared.http import get_raw
from maplestats_mcp.shared.limits import fit_to_budget
from maplestats_mcp.shared.rate_limiter import get_limiter

_LIMITER = get_limiter(
    constants.RATE_LIMIT_SOURCE,
    rate=constants.RATE_LIMIT_PER_SECOND,
    capacity=constants.RATE_LIMIT_CAPACITY,
)
_SLUG = re.compile(r"^[a-z0-9-]{3,200}$")
_DATA_FILE = re.compile(r"data-table-(en|fr)\.xlsx$")
_SITEMAP_NS = "http://www.sitemaps.org/schemas/sitemap/0.9"
_XHTML_NS = "http://www.w3.org/1999/xhtml"


async def _get(url: str, lang: str = "en") -> httpx.Response:
    await _LIMITER.acquire()
    try:
        return await get_raw(url, timeout=120.0)
    except httpx.HTTPStatusError as exc:
        status = exc.response.status_code
        if status == 404:
            raise lang_error(
                NotFound,
                lang,
                f"cihi: nothing published at {url}.",
                f"cihi : rien n'est publié à {url}.",
            ) from exc
        raise lang_error(
            UpstreamError,
            lang,
            f"cihi: {url} returned HTTP {status}.",
            f"cihi : {url} a renvoyé HTTP {status}.",
        ) from exc
    except httpx.HTTPError as exc:
        raise lang_error(
            UpstreamUnavailable,
            lang,
            f"cihi: {url} did not respond in time.",
            f"cihi : {url} n'a pas répondu à temps.",
        ) from exc


async def _page(url: str, lang: str = "en") -> tuple[str, bool]:
    async def fetch() -> str:
        return (await _get(url, lang)).text

    return await cached_fetch(f"cihi:page:{url}", constants.CACHE_TTL_PAGE_SECONDS, fetch)


def _fold(text: str) -> str:
    decomposed = unicodedata.normalize("NFKD", text.casefold())
    return "".join(c for c in decomposed if not unicodedata.combining(c))


def _clean(text: str) -> str:
    return " ".join(text.split())


def _parse_library(html: str, path: str = constants.INDICATOR_PATH) -> list[IndicatorRef]:
    soup = BeautifulSoup(html, "html.parser")
    main = soup.find("main") or soup
    refs: list[IndicatorRef] = []
    for link in main.find_all("a", href=True):
        href = str(link["href"])
        if href.startswith(path):
            slug = href[len(path) :].strip("/")
            refs.append(
                IndicatorRef(slug=slug, name=_clean(link.get_text()), url=constants.BASE_URL + href)
            )
    return refs


async def _library(lang: str = "en") -> tuple[list[IndicatorRef], bool]:
    french = lang == "fr"
    library_url = constants.FR_LIBRARY_URL if french else constants.LIBRARY_URL
    path = constants.FR_INDICATOR_PATH if french else constants.INDICATOR_PATH

    async def fetch() -> list[IndicatorRef]:
        seen: dict[str, IndicatorRef] = {}
        for page in range(constants.LIBRARY_MAX_PAGES):
            response = await _get(f"{library_url}?page={page}", lang)
            refs = _parse_library(response.text, path)
            new = [r for r in refs if r.slug not in seen]
            if not new:
                break
            seen.update({r.slug: r for r in new})
        return list(seen.values())

    key = "cihi:library:fr" if french else "cihi:library"
    return await cached_fetch(key, constants.CACHE_TTL_LIBRARY_SECONDS, fetch)


def _sitemap_pairs(body: str) -> tuple[list[str], dict[str, str]]:
    """(child sitemap URLs, English slug -> French slug) from one sitemap file."""
    root = ElementTree.fromstring(body.lstrip("\N{ZERO WIDTH NO-BREAK SPACE}").encode())
    children = [
        (loc.text or "").strip()
        for loc in root.iter(f"{{{_SITEMAP_NS}}}loc")
        if root.tag.endswith("sitemapindex")
    ]
    pairs: dict[str, str] = {}
    english_prefix = constants.BASE_URL + constants.INDICATOR_PATH
    french_prefix = constants.BASE_URL + constants.FR_INDICATOR_PATH
    for entry in root.iter(f"{{{_SITEMAP_NS}}}url"):
        alternates = {
            link.get("hreflang"): link.get("href") or ""
            for link in entry.iter(f"{{{_XHTML_NS}}}link")
        }
        english, french = alternates.get("en", ""), alternates.get("fr", "")
        if english.startswith(english_prefix) and french.startswith(french_prefix):
            english_slug = english[len(english_prefix) :].strip("/")
            pairs[english_slug] = french[len(french_prefix) :].strip("/")
    return children, pairs


async def _pairing(lang: str = "en") -> tuple[dict[str, str], bool]:
    """English slug -> French slug, as CIHI's own hreflang alternates pair them."""

    async def fetch() -> dict[str, str]:
        # A child sitemap is about 1 MB of XML: parse it off the event loop.
        index = (await _get(constants.SITEMAP_URL, lang)).text
        children, pairs = await run_parse(_sitemap_pairs, index)
        for url in children[: constants.SITEMAP_MAX_PAGES]:
            page = (await _get(url, lang)).text
            pairs.update((await run_parse(_sitemap_pairs, page))[1])
        if not pairs:
            raise lang_error(
                UpstreamError,
                lang,
                f"cihi: {constants.SITEMAP_URL} lists no indicator pages.",
                f"cihi : {constants.SITEMAP_URL} ne liste aucune page d'indicateur.",
            )
        return pairs

    return await cached_fetch("cihi:pairing", constants.CACHE_TTL_PAIRING_SECONDS, fetch)


def _french_stem(word: str) -> str:
    # Names are mostly singular ("service d'urgence", "hôpital"): a plural
    # query word matches by its stem ("urgences" -> "urgence", "hôpitaux"
    # -> "hopita", which matches "hopital" and "hopitaux").
    if len(word) > 4 and word.endswith("aux"):
        return word[:-2]
    if len(word) > 3 and word.endswith(("s", "x")):
        return word[:-1]
    return word


def _french_words(query: str) -> list[str]:
    words = re.findall(r"[a-z0-9]+", _fold(query))
    return [_french_stem(w) for w in words if w not in constants.FR_STOP_WORDS]


async def _french_refs() -> tuple[list[IndicatorRef], str | None, bool]:
    """French library entries with their English twin, then English-only entries."""
    french, french_cached = await _library("fr")
    english, english_cached = await _library("en")
    pairs, pairs_cached = await _pairing("fr")
    to_english = {fr: en for en, fr in pairs.items()}
    # These refs answer a French search, so their notes are in French.
    refs = [
        r.model_copy(
            update={
                "english_slug": to_english.get(r.slug),
                "note": None
                if r.slug in to_english
                else "L'ICIS ne publie aucune page en anglais pour cet indicateur.",
            }
        )
        for r in french
    ]
    listed = {r.slug for r in french}
    missing = [r for r in english if pairs.get(r.slug) not in listed]
    refs += [
        r.model_copy(
            update={
                "english_slug": r.slug,
                "note": fr_or_en(
                    "fr",
                    "",
                    "Aucune page en français dans la bibliothèque française de l'ICIS : "
                    "nom et page en anglais.",
                ),
            }
        )
        for r in missing
    ]
    note = (
        fr_or_en(
            "fr",
            "",
            f"{len(missing)} indicateur(s) sans page en français, listé(s) sous leur nom anglais.",
        )
        if missing
        else None
    )
    return refs, note, french_cached and english_cached and pairs_cached


async def search_indicators(query: str = "", lang: str = "en") -> IndicatorSearchResult:
    note = None
    if lang == "fr":
        refs, note, cached = await _french_refs()
        words = _french_words(query)
        url = constants.FR_LIBRARY_URL
    else:
        refs, cached = await _library()
        words = _fold(query).split()
        url = constants.LIBRARY_URL
    matches = [r for r in refs if all(w in _fold(r.name) for w in words)]
    return IndicatorSearchResult(
        indicators=matches,
        total_matches=len(matches),
        note=note,
        provenance=make_provenance(
            source=constants.RATE_LIMIT_SOURCE,
            url=url,
            cached=cached,
            schema_name="cihi.IndicatorSearchResult",
            freshness=fr_or_en(
                lang,
                "indicator lists cached 7 days, English-French pairing 1 day",
                "listes d'indicateurs mises en cache 7 jours, appariement anglais-français 1 jour",
            ),
            lang=lang,
        ),
    )


def _slug(indicator: str, lang: str = "en") -> tuple[str, str | None]:
    """The slug, and "en"/"fr" when a page URL says which library it is from."""
    value = indicator.strip().rstrip("/").lower()
    kind = None
    for path, library in ((constants.INDICATOR_PATH, "en"), (constants.FR_INDICATOR_PATH, "fr")):
        prefix = (constants.BASE_URL + path).lower()
        if value.startswith(prefix):
            value, kind = value[len(prefix) :], library
    if not _SLUG.match(value):
        raise lang_error(
            InvalidInput,
            lang,
            f"indicator must be a slug from cihi_search_indicators, got {indicator!r}.",
            f"indicator doit être un identifiant (slug) tiré de cihi_search_indicators ; reçu "
            f"{indicator!r}.",
        )
    return value, kind


def _english_url(slug: str) -> str:
    return f"{constants.BASE_URL}{constants.INDICATOR_PATH}{slug}"


def _french_url(slug: str) -> str:
    return f"{constants.BASE_URL}{constants.FR_INDICATOR_PATH}{slug}"


async def _locate(indicator: str, lang: str) -> tuple[str, str, str | None, str | None]:
    """(page URL to read, English slug or French if none, French slug, note).

    An English slug is read as before; the French page is the English
    page's hreflang link. A French slug is turned into its English twin
    through the pairing, so either identifier works in either language.
    """
    slug, kind = _slug(indicator, lang)
    if kind is None:
        try:
            await _page(_english_url(slug), lang)
            kind = "en"
        except NotFound:
            kind = "fr"
    if kind == "en":
        if lang != "fr":
            return _english_url(slug), slug, None, None
        html, _ = await _page(_english_url(slug), lang)
        alternate = BeautifulSoup(html, "html.parser").find("link", hreflang="fr")
        href = (
            str(alternate["href"]) if isinstance(alternate, Tag) and alternate.get("href") else ""
        )
        if not href or href.rstrip("/") == _english_url(slug):
            note = fr_or_en(
                lang,
                "",
                "L'ICIS ne publie aucune page en français pour cet indicateur ; la page en "
                "anglais est affichée.",
            )
            return _english_url(slug), slug, None, note
        return href, slug, href.rstrip("/").rsplit("/", 1)[-1], None
    pairs, _ = await _pairing(lang)
    to_english = {fr: en for en, fr in pairs.items()}
    english = to_english.get(slug)
    if english is None:
        try:
            # Neither an English page nor a paired French one: read it as a
            # French-only page if CIHI has one.
            await _page(_french_url(slug), lang)
        except NotFound as exc:
            raise lang_error(
                NotFound,
                lang,
                f"cihi: no indicator {slug!r} in the English or French library; "
                "use a slug from cihi_search_indicators.",
                f"cihi : aucun indicateur {slug!r} dans la bibliothèque anglaise ou française ; "
                "utilisez un identifiant tiré de cihi_search_indicators.",
            ) from exc
        note = None if lang == "fr" else "CIHI publishes no English page for this indicator."
        return _french_url(slug), slug, slug, note
    if lang == "fr":
        return _french_url(slug), english, slug, None
    return _english_url(english), english, slug, None


def _data_file(soup: BeautifulSoup) -> str | None:
    for link in soup.find_all("a", href=True):
        href = str(link["href"])
        if _DATA_FILE.search(href):
            return href if href.startswith("http") else constants.BASE_URL + href
    return None


async def get_indicator(indicator: str, lang: str = "en") -> IndicatorDetail:
    url, slug, french_slug, note = await _locate(indicator, lang)
    html, cached = await _page(url, lang)
    soup = BeautifulSoup(html, "html.parser")
    heading = soup.find("h1")
    summary = soup.find(class_="view-metadata-summary")
    description = None
    facts: dict[str, str] = {}
    if isinstance(summary, Tag):
        first = summary.find("div", class_="col-12")
        description = _clean(first.get_text()) if isinstance(first, Tag) else None
        for item in summary.select("ul.indicator-meta-summary li"):
            label, _, value = _clean(item.get_text()).partition(":")
            if value:
                facts[label.strip()] = value.strip()
        topics = [_clean(t.get_text()) for t in summary.select(".indicator-topics li")]
        if topics:
            in_french = constants.FR_INDICATOR_PATH in url
            facts["Thèmes" if in_french else "Topics"] = ", ".join(topics)
    return IndicatorDetail(
        slug=slug,
        french_slug=french_slug,
        name=_clean(heading.get_text()) if isinstance(heading, Tag) else slug,
        note=note,
        description=description,
        facts=facts,
        page_url=url,
        data_file_url=_data_file(soup),
        provenance=make_provenance(
            source=constants.RATE_LIMIT_SOURCE,
            url=url,
            cached=cached,
            schema_name="cihi.IndicatorDetail",
            lang=lang,
        ),
    )


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_workbook(body: bytes) -> dict[str, tuple[str, list[str], list[list[str]]]]:
    """Each Table sheet as (title, header, rows)."""
    # Imported here, not at module level: openpyxl takes ~0.7 s to import and
    # only this tool needs it, so server startup stays faster (2026-09-24).
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    tables: dict[str, tuple[str, list[str], list[list[str]]]] = {}
    for sheet in workbook.worksheets:
        if not sheet.title.lower().startswith(("table", "tableau")):
            continue
        rows = sheet.iter_rows(values_only=True)
        title_row = next(rows, ())
        header_row = next(rows, ())
        header = [_cell(v) for v in header_row]
        while header and not header[-1]:
            header.pop()
        width = len(header)
        data = [cells for row in rows if any(cells := [_cell(v) for v in row[:width]])]
        tables[sheet.title] = (_cell(title_row[0] if title_row else ""), header, data)
    workbook.close()
    return tables


async def _tables(
    url: str, lang: str = "en"
) -> tuple[dict[str, tuple[str, list[str], list[list[str]]]], bool]:
    if urlparse(url).hostname not in constants.ALLOWED_HOSTS:
        raise lang_error(
            UpstreamError,
            lang,
            f"cihi: unexpected data file host in {url}.",
            f"cihi : hôte inattendu pour le fichier de données {url}.",
        )

    async def fetch() -> dict[str, tuple[str, list[str], list[list[str]]]]:
        response = await _get(url, lang)
        if len(response.content) > constants.MAX_FILE_BYTES:
            raise lang_error(
                UpstreamError,
                lang,
                f"cihi: {url} is larger than this tool reads.",
                f"cihi : {url} dépasse la taille que cet outil peut lire.",
            )
        try:
            # Parsing a ~2 MB workbook held the event loop for ~1.8 s (measured
            # 2026-09-24), stalling every other request on the server meanwhile.
            return await run_parse(_parse_workbook, response.content)
        except Exception as exc:  # openpyxl raises several unrelated types
            raise lang_error(
                UpstreamError,
                lang,
                f"cihi: {url} is not a readable XLSX file.",
                f"cihi : {url} n'est pas un fichier XLSX lisible.",
            ) from exc

    return await cached_fetch(f"cihi:data:{url}", constants.CACHE_TTL_DATA_SECONDS, fetch)


async def get_indicator_data(
    indicator: str,
    *,
    place: str | None = None,
    filters: dict[str, str] | None = None,
    table: str | None = None,
    columns: list[str] | None = None,
    limit: int = constants.ROWS_DEFAULT,
    lang: str = "en",
) -> IndicatorData:
    """Rows of an indicator's data table; the latest rows come last in the file."""
    if limit < 1 or limit > constants.ROWS_MAX:
        raise lang_error(
            InvalidInput,
            lang,
            f"limit must be between 1 and {constants.ROWS_MAX}, got {limit}.",
            f"limit doit être compris entre 1 et {constants.ROWS_MAX} ; reçu {limit}.",
        )
    detail = await get_indicator(indicator, lang)
    if not detail.data_file_url:
        raise lang_error(
            NotFound,
            lang,
            f"cihi: {detail.name!r} has no downloadable data table.",
            f"cihi : {detail.name!r} n'a aucun tableau de données téléchargeable.",
        )
    tables, tables_cached = await _tables(detail.data_file_url, lang)
    if not tables:
        raise lang_error(
            UpstreamError,
            lang,
            f"cihi: {detail.data_file_url} has no data table sheets.",
            f"cihi : {detail.data_file_url} n'a aucune feuille de tableau de données.",
        )
    sheet = table or next(iter(tables))
    if sheet not in tables:
        raise lang_error(
            InvalidInput,
            lang,
            f"Unknown table {sheet!r}; tables are {list(tables)}.",
            f"tableau {sheet!r} inconnu ; les tableaux sont {list(tables)}.",
        )
    title, header, data = tables[sheet]
    by_lower = {h.lower(): i for i, h in enumerate(header)}

    def index(name: str) -> int:
        position = by_lower.get(name.strip().lower())
        if position is None:
            raise lang_error(
                InvalidInput,
                lang,
                f"Unknown column {name!r}; columns are {header}.",
                f"colonne {name!r} inconnue ; les colonnes sont {header}.",
            )
        return position

    wanted = [(index(k), v.strip().lower()) for k, v in (filters or {}).items()]
    place_index = next(
        (
            i
            for name, i in by_lower.items()
            if name in ("place or organization", "lieu ou organisme", "lieu ou organisation")
        ),
        None,
    )
    needle = (place or "").strip().lower()
    if needle and place_index is None:
        raise lang_error(
            InvalidInput,
            lang,
            f"This table has no place column; columns are {header}.",
            f"ce tableau n'a pas de colonne de lieu ; les colonnes sont {header}.",
        )

    def keep(row: list[str]) -> bool:
        cells = row + [""] * (len(header) - len(row))
        if needle and place_index is not None and needle not in cells[place_index].lower():
            return False
        return all(cells[i].lower() == v for i, v in wanted)

    matching = [r for r in data if keep(r)]
    kept = [r + [""] * (len(header) - len(r)) for r in matching[-limit:]]
    # Rows carry 33 columns (about 1.35 KB each, live 2026-10-03), many of
    # them blank for a given indicator; pick columns, or drop the blank ones.
    if columns:
        chosen = [index(name) for name in columns]
        empty: list[str] = []
    else:
        chosen = [i for i in range(len(header)) if any(row[i] for row in kept)]
        empty = [header[i] for i in range(len(header)) if i not in chosen] if kept else []
        if not kept:
            chosen = list(range(len(header)))
    shown = [header[i] for i in chosen]

    # The latest rows come last in the file: fill the byte budget newest
    # first, then put them back in file order.
    rows = fit_to_budget([{header[i]: row[i] for i in chosen} for row in kept][::-1])[::-1]
    return IndicatorData(
        slug=detail.slug,
        table=title or sheet,
        tables=list(tables),
        columns=shown,
        empty_columns=empty,
        rows=rows,
        total_rows=len(data),
        matching_rows=len(matching),
        returned_count=len(rows),
        data_file_url=detail.data_file_url,
        note=detail.note,
        provenance=make_provenance(
            source=constants.RATE_LIMIT_SOURCE,
            url=detail.data_file_url,
            cached=detail.provenance.cached and tables_cached,
            schema_name="cihi.IndicatorData",
            coverage=fr_or_en(
                lang,
                f"last {len(rows)} of {len(matching)} matching rows",
                f"{len(rows)} dernières lignes sur {len(matching)} correspondantes",
            ),
            limits=truncation_note_lang(
                lang,
                returned=len(rows),
                total=len(matching),
                unit="matching rows",
                unit_fr="lignes correspondantes",
                order="latest",
                how_to_get_more="filter by place or column values, pick fewer columns, or raise "
                f"limit (max {constants.ROWS_MAX}; responses are capped near 200 KB)",
                how_to_get_more_fr="filtrez par lieu ou par valeurs de colonne, choisissez moins "
                f"de colonnes ou augmentez limit (max. {constants.ROWS_MAX} ; les réponses sont "
                "plafonnées à environ 200 Ko)",
            ),
            lang=lang,
        ),
    )
