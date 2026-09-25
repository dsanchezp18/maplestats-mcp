from __future__ import annotations

import io
import re
import zipfile
from datetime import date

import openpyxl
import pytest

from maplestats_mcp.modules.ised.ip_horizons import client, constants
from maplestats_mcp.shared import cache as cache_module
from maplestats_mcp.shared.errors import InvalidInput, UpstreamError


@pytest.fixture(autouse=True)
def _clear_cache():
    cache_module._caches.clear()
    yield


_BASE = "https://opic-cipo.ca/cipo/client_downloads"
_CKAN = re.compile(re.escape(constants.CKAN_BASE_URL) + r"\?id=.*")

# Real resource shapes from the patent package, checked live 2026-09-25:
# an older unsplit release, the newest split release, a CKAN name that
# says txt_format for a CSV file, and the dictionary link.
_RESOURCES = [
    {
        "name": "Patent Data Dictionary",
        "url": f"{_BASE}/IP_Horizon_Resources/PT_Data_Dictionary.zip",
    },
    {"name": "PT_claim", "url": f"{_BASE}/Patent_CSV_2024_05_22/PT_claim.zip"},
    {
        "name": "PT_claim_1_to_2000000_2024-10-11",
        "url": f"{_BASE}/Patent_CSV_2024_10_11/PT_claim_1_to_2000000_2024-10-11.zip",
    },
    {
        "name": "PT_claim_2000001_to_4000000_2024-10-11",
        "url": f"{_BASE}/Patent_CSV_2024_10_11/PT_claim_2000001_to_4000000_2024-10-11.zip",
    },
    {
        "name": "PT_abstract_txt_format_2000001_to_4000000_2024-05-27",
        "url": f"{_BASE}/Patent_CSV_2024_05_22/PT_abstract_2000001_to_4000000_2024-05-27.zip",
    },
    {
        "name": "PT_IPC_classification",
        "url": f"{_BASE}/Patent_CSV_2024_05_22/PT_IPC_classification.zip",
    },
]


def test_parse_file_url_reads_table_range_and_release():
    item = client.parse_file_url(
        "patent",
        f"{_BASE}/Patent_CSV_2024_05_22/PT_disclosure_txt_format_1_to_2000000_2024-05-27.zip",
        "x",
    )
    assert item is not None
    assert item.table == "disclosure"
    assert item.text_format is True
    assert (item.number_from, item.number_to) == (1, 2_000_000)
    # the folder date wins over the file's own date
    assert item.release_date == date(2024, 5, 22)


def test_parse_file_url_ignores_the_folder_prefix():
    # ID_CSV_... and TM_CSV_... folders share the file prefix, checked live.
    item = client.parse_file_url(
        "industrial_design", f"{_BASE}/ID_CSV_2024_03_07/ID_application_main.zip", "x"
    )
    assert item is not None
    assert item.table == "application_main"
    assert item.number_from is None
    assert item.release_date == date(2024, 3, 7)


async def test_list_files_latest_only_drops_older_releases(httpx_mock):
    httpx_mock.add_response(url=_CKAN, json={"result": {"resources": _RESOURCES}})
    result = await client.list_files("patent")
    claims = [f for f in result.files if f.table == "claim"]
    assert [(f.number_from, f.release_date) for f in claims] == [
        (1, date(2024, 10, 11)),
        (2_000_001, date(2024, 10, 11)),
    ]
    assert result.tables == ["abstract", "claim", "ipc_classification"]
    # the CKAN name says txt_format but the URL is the CSV variant
    abstract = next(f for f in result.files if f.table == "abstract")
    assert abstract.text_format is False
    assert all("IP_Horizon_Resources" not in f.url for f in result.files)


async def test_list_files_all_releases_and_table_filter(httpx_mock):
    httpx_mock.add_response(url=_CKAN, json={"result": {"resources": _RESOURCES}})
    result = await client.list_files("patent", table="CLAIM", latest_only=False)
    assert result.returned_count == 3


async def test_list_files_unknown_table_raises(httpx_mock):
    httpx_mock.add_response(url=_CKAN, json={"result": {"resources": _RESOURCES}})
    with pytest.raises(InvalidInput, match="tables are"):
        await client.list_files("patent", table="inventors")


async def test_list_files_null_resources(httpx_mock):
    httpx_mock.add_response(url=_CKAN, json={"result": {"resources": None}})
    result = await client.list_files("trademark")
    assert result.files == []


async def test_list_files_missing_result_raises(httpx_mock):
    httpx_mock.add_response(url=_CKAN, json={"success": False})
    with pytest.raises(UpstreamError):
        await client.list_files("patent")


def _dictionary_zip() -> bytes:
    workbook = openpyxl.Workbook()
    overview = workbook.active
    assert overview is not None
    overview.title = "PT"
    overview.append(["Dataset coverage period - Période", None])
    overview.append(
        [
            "The patent database covers all patent applications filed in Canada from 1869.",
            "La base de données...",
        ]
    )
    workbook.create_sheet("PT_Road_Map")
    main = workbook.create_sheet("PT_Main")
    main.append(["PT_Main dictionary - Dictionnaire PT_principal"])
    main.append(["Variable Name", "Nom de variable", "Variable Name - Nom de variable"])
    main.append(
        ["Patent Number ", "Numéro du brevet", "=A3", "Numeric", "Numérique", "An id.", "Un id."]
    )
    ipc = workbook.create_sheet("PT_IPC Classification")
    ipc.append([None, "PT_IPC Classification dictionary"])
    ipc.append(["Variable Name", "Nom de variable"])
    ipc.append(["IPC Section", "Section CIB", None, "Text", "Texte"])
    book = io.BytesIO()
    workbook.save(book)
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as handle:
        handle.writestr("PT_Data Dictionary.xlsx", book.getvalue())
    return archive.getvalue()


async def test_get_dictionary_parses_sheets(httpx_mock):
    httpx_mock.add_response(url=constants.DICTIONARY_URLS["patent"], content=_dictionary_zip())
    result = await client.get_dictionary("patent")
    assert result.tables == ["ipc_classification", "main"]
    first = result.fields[0]
    assert first.table == "main"
    assert first.name_en == "Patent Number"  # trailing space in the sheet
    assert first.type_en == "Numeric"
    assert first.description_fr == "Un id."
    assert result.notes and result.notes[0].startswith("The patent database covers")

    only_ipc = await client.get_dictionary("patent", table="ipc_classification")
    assert [f.name_en for f in only_ipc.fields] == ["IPC Section"]
    assert only_ipc.provenance.cached is True


async def test_get_dictionary_html_404_page_raises(httpx_mock):
    # opic-cipo.ca serves an HTML "404 - File or directory not found" page.
    httpx_mock.add_response(
        url=constants.DICTIONARY_URLS["industrial_design"], content=b"<!DOCTYPE html><html>"
    )
    with pytest.raises(UpstreamError, match="not a dictionary ZIP"):
        await client.get_dictionary("industrial_design")


async def test_get_dictionary_trademark_has_none():
    with pytest.raises(InvalidInput):
        await client.get_dictionary("trademark")
