"""Manitoba general election results from Elections Manitoba's downloads (1999 to 2023).

Checked live 2026-10-03 on https://www.electionsmanitoba.ca/en/Results/Elections1999AndLater.
Each general election's page links, under /downloads/:

- "<n>GE Summary of Votes Received" (.xls, .xlsx for 2023): one row per candidate. The
  division name is on the first candidate's row only (1999 to 2007, 2019, 2023) or on every
  row (2011, 2016). In 1999 the second row of a bilingual division carries its French name
  ("Brandon Est", "Saint-Boniface"); from 2003 the name cell holds both ("Brandon East /
  Brandon Est", "Brandon East\\nBrandon-Est"). The 2011 and 2016 header spells "Canidate".
  2023 adds "Declined" and "Rejected" rows per division. The 2019 workbook keeps an earlier
  sheet named "Old" beside the final one.
- "Summary_of_Results_GE<year>" (.xls; the 1999 to 2007 files are .xlsx workbooks with an
  .xls name; .xlsx for 2023): one row per division with the member elected, registered
  voters, ballots cast, votes by party, rejected and declined ballots. Its division names run
  words together in places ("BrandonWest", "LaVerendrye", "St.Vital"), so names are matched
  on their letters and digits only. The 2019 workbook again has an "Old" sheet.
- "<n>GE.zip": results by voting area in one of two shapes. Long (one row per voting area
  and candidate): a single workbook for 1999 (.xlsx) and 2003 (.xls), one workbook per
  division for 2007 with last and first names in separate columns. Wide (one row per voting
  area, one column per candidate headed "LAST, First (PARTY)"): a single .xls with one sheet
  per division for 2011, one workbook per division for 2016, 2019 and 2023 (2023's workbook
  also has a sheet by polling place, which is skipped; it names no voting place per area).

Summing the candidate rows gives NDP 32, PC 24 and Liberal 1 in 1999, 35, 20 and 2 in 2003,
36, 19 and 2 in 2007, 37, 19 and 1 in 2011, PC 40, NDP 14 and Liberal 3 in 2016, 36, 18 and
3 in 2019, and NDP 34, PC 22 and Liberal 1 in 2023: the seats the legislature was sworn in
with. The voting-area rows add up to the division's votes in the summary of votes received
for 387 of the 399 divisions; in the other twelve (three in 2003, six in 2007, two in 2016,
one in 2019) the voting-area file itself differs, so responses give both sums. Candidate names keep the capitals the files use ("Stu BRIESE" in 2011).
"""

from __future__ import annotations

import io
import re
import unicodedata
from collections.abc import Iterator
from dataclasses import dataclass, field

from maplestats_mcp.modules.elections_provincial import constants
from maplestats_mcp.modules.elections_provincial.common import (
    Candidate,
    District,
    fetch_bytes,
    finish_shares,
    fold,
    mark_winners,
)
from maplestats_mcp.shared import remote_zip
from maplestats_mcp.shared.errors import UpstreamError
from maplestats_mcp.shared.executor import run_parse
from maplestats_mcp.shared.rate_limiter import get_limiter

_LIMITER = get_limiter(
    constants.MB_RATE_LIMIT_SOURCE,
    rate=constants.MB_RATE_LIMIT_PER_SECOND,
    capacity=constants.MB_RATE_LIMIT_CAPACITY,
)

# The files abbreviate parties several ways over the years ("NDP/NPD", "NDP / NPD / NPD",
# "PC Manitoba", "Liberal", "The Manitoba Greens"); each is reduced to one code here.
_CODE_ALIASES = {
    "PC MANITOBA": "PC",
    "LIBERAL": "Lib.",
    "LIB": "Lib.",
    "THE MANITOBA GREENS": "GPM",
    "IND": "Ind.",
}
PARTY_NAMES: dict[str, str] = {
    "CPC-M": "Communist Party of Canada - Manitoba",
    "GPM": "Green Party of Manitoba",
    "Ind.": "Independent",
    "KP": "Keystone Party of Manitoba",
    "Lib.": "Manitoba Liberal Party",
    "LPM": "Libertarian Party of Manitoba",
    "MBFWD": "Manitoba Forward",
    "MF": "Manitoba First",
    "MLP": "Manitoba Liberal Party",
    "MP": "Manitoba Party",
    "NDP": "New Democratic Party of Manitoba",
    "PC": "Progressive Conservative Party of Manitoba",
}

# The 2011 summary of results heads the Communist column "CPC".
_HEADER_CODES = {"CPC": "CPC-M", "CPC_M": "CPC-M"}


@dataclass
class AreaVote:
    name: str | None
    party: str
    party_code: str
    votes: int


@dataclass
class Area:
    """One voting area (or special poll: advance, absentee, homebound, write-in)."""

    label: str
    place: str | None
    electors: int | None
    rejected: int
    declined: int
    votes: list[AreaVote] = field(default_factory=list)


def files(number: str) -> constants.ManitobaFiles:
    return constants.MB_FILES[number]


def key(text: str) -> str:
    """Letters and digits of the English part of a name, folded: 'St.Vital' -> 'stvital'."""
    return re.sub(r"[^a-z0-9]", "", fold(english(text)))


def english(text: str) -> str:
    """The English part of a bilingual label ('Brandon East / Brandon Est', 'A\\nB')."""
    first = re.split(r"\s*/\s*|\n", str(text).strip(), maxsplit=1)[0]
    return re.sub(r"\s+", " ", first).strip()


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _count(value: object) -> int:
    """A count: blank (or the single space some summaries hold) is zero."""
    text = _cell(value).replace(",", "")
    if not text:
        return 0
    try:
        return int(float(text))
    except ValueError as exc:
        raise UpstreamError(
            f"elections_provincial: Manitoba file has the count {value!r}."
        ) from exc


def _is_number(value: object) -> bool:
    if isinstance(value, int | float):
        return True
    return bool(re.fullmatch(r"\d+(\.0+)?", _cell(value)))


def _person(name: str) -> str:
    """'First Last' from 'LAST, First' or 'First Last', spaces tidied."""
    name = re.sub(r"\s+", " ", name).strip()
    last, comma, first = name.partition(",")
    return f"{first.strip()} {last.strip()}".strip() if comma else name


def party(raw: str) -> tuple[str, str]:
    """(name, code) for a party label as the files write it."""
    text = re.sub(r"\s+", " ", raw).strip()
    code = text.split("/")[0].strip()
    code = _CODE_ALIASES.get(code.upper().rstrip("."), code)
    code = _HEADER_CODES.get(code, code)
    return PARTY_NAMES.get(code, code), code


def _header_key(value: object) -> str:
    stripped = unicodedata.normalize("NFKD", _cell(value))
    return re.sub(r"[^a-z0-9]", "", stripped.casefold())


def _sheets(body: bytes) -> Iterator[tuple[str, list[list[object]]]]:
    """(sheet name, rows) of an .xls or .xlsx workbook, told apart by content, not name.

    The 1999 to 2007 summaries of results are .xlsx workbooks saved with an .xls name.
    """
    if body[:2] == b"PK":
        # openpyxl takes ~0.7 s to import, so it loads on first use.
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
        try:
            for sheet in workbook:
                yield sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
        return
    import xlrd

    book = xlrd.open_workbook(file_contents=body)
    for sheet in book.sheets():
        rows: list[list[object]] = [list(sheet.row_values(i)) for i in range(sheet.nrows)]
        yield sheet.name, rows


def _main_sheet(body: bytes, what: str) -> list[list[object]]:
    """The sheet holding the data: the first with rows, skipping 2019's 'Old' sheet."""
    for name, rows in _sheets(body):
        if name.strip().lower() == "old":
            continue
        if any(_cell(c) for row in rows for c in row):
            return rows
    raise UpstreamError(f"elections_provincial: Manitoba {what} has no data sheet.")


def _header_row(rows: list[list[object]], what: str) -> int:
    for i, row in enumerate(rows):
        if row and _header_key(row[0]).startswith("electoraldivision"):
            return i
    raise UpstreamError(
        f"elections_provincial: Manitoba {what} has no 'Electoral Division' header."
    )


def _column(header: list[object], *prefixes: str) -> int | None:
    keys = [_header_key(c) for c in header]
    for i, k in enumerate(keys):
        if any(k.startswith(p) for p in prefixes):
            return i
    return None


@dataclass
class _Summary:
    name: str
    electors: int | None
    ballots: int | None
    rejected: int | None
    member: str


def parse_summary(body: bytes) -> dict[str, _Summary]:
    """Division rows of a summary of results, keyed by the name's letters and digits."""
    rows = _main_sheet(body, "summary of results")
    top = _header_row(rows, "summary of results")
    header = rows[top]
    member_col = _column(header, "member")
    registered_col = _column(header, "reg")
    ballots_col = _column(header, "votescast")
    rejected_col = _column(header, "rejected")
    if None in (member_col, registered_col, ballots_col, rejected_col):
        raise UpstreamError(
            "elections_provincial: Manitoba summary of results lacks the expected columns "
            f"(has {[_cell(c) for c in header]})."
        )
    assert member_col is not None and registered_col is not None
    assert ballots_col is not None and rejected_col is not None
    found: dict[str, _Summary] = {}
    for row in rows[top + 1 :]:
        if len(row) <= max(member_col, registered_col, ballots_col, rejected_col):
            continue
        name, member = _cell(row[0]), _cell(row[member_col])
        if not name or not member:
            continue
        found[key(name)] = _Summary(
            name=english(name),
            electors=_count(row[registered_col]) or None,
            ballots=_count(row[ballots_col]) or None,
            rejected=_count(row[rejected_col]),
            member=member,
        )
    if not found:
        raise UpstreamError("elections_provincial: Manitoba summary of results has no rows.")
    return found


def parse_results(votes_body: bytes, summary_body: bytes) -> list[District]:
    """Districts from one election's summary of votes received and summary of results."""
    summary = parse_summary(summary_body)
    rows = _main_sheet(votes_body, "summary of votes received")
    top = _header_row(rows, "summary of votes received")
    header = rows[top]
    candidate_col = _column(header, "can")  # "Candidate", and "Canidate" in 2011 and 2016
    party_col = _column(header, "party", "registeredparty", "political")
    votes_col = _column(header, "votes")
    if None in (candidate_col, party_col, votes_col):
        raise UpstreamError(
            "elections_provincial: Manitoba summary of votes received lacks the expected "
            f"columns (has {[_cell(c) for c in header]})."
        )
    assert candidate_col is not None and party_col is not None and votes_col is not None

    districts: dict[str, District] = {}
    current: District | None = None
    for row in rows[top + 1 :]:
        if len(row) <= max(candidate_col, party_col, votes_col):
            continue
        name = _cell(row[candidate_col])
        division = _cell(row[0])
        # A division cell names a new division only if the summary of results knows it;
        # otherwise it is the French name on a division's second row (1999).
        if division and key(division) in summary:
            division_key = key(division)
            if division_key not in districts:
                info = summary[division_key]
                districts[division_key] = District(
                    name=english(division),
                    number=None,
                    electors=info.electors,
                    rejected_ballots=info.rejected,
                    turnout=(
                        round(100 * info.ballots / info.electors, 2)
                        if info.ballots and info.electors
                        else None
                    ),
                )
            current = districts[division_key]
        # Totals, notes, and 2023's "Declined" and "Rejected" rows are not candidates.
        if not name or name.lower() in ("declined", "rejected"):
            continue
        if current is None:
            raise UpstreamError(
                f"elections_provincial: Manitoba candidate {name!r} comes before any division."
            )
        party_name, code = party(_cell(row[party_col]))
        current.candidates.append(
            Candidate(
                name=_person(name),
                party=party_name,
                party_code=code,
                votes=_count(row[votes_col]),
                share=None,
            )
        )

    missing = [summary[k].name for k in summary if k not in districts]
    if missing:
        raise UpstreamError(
            f"elections_provincial: Manitoba summary of votes received lacks divisions {missing}."
        )
    built = list(districts.values())
    for district in built:
        district.candidates.sort(key=lambda c: -c.votes)
        mark_winners(district)
        finish_shares(district)
    return built


async def fetch(number: str) -> list[District]:
    """Districts for one Manitoba general election, by its number (37 to 43)."""
    source = files(number)
    context = f"elections_provincial:mb:{number}"
    votes = await fetch_bytes(
        source.votes, context=context, max_bytes=constants.MB_MAX_BYTES, limiter=_LIMITER
    )
    summary = await fetch_bytes(
        source.summary, context=context, max_bytes=constants.MB_MAX_BYTES, limiter=_LIMITER
    )
    try:
        return await run_parse(parse_results, votes, summary)
    except UpstreamError:
        raise
    except Exception as exc:  # xlrd and openpyxl raise several unrelated types
        raise UpstreamError(f"elections_provincial: could not read {source.votes}: {exc}") from exc


# --- Results by voting area -------------------------------------------------------------


def _candidate_heading(text: str) -> AreaVote:
    """'BRIESE,\\n Stu (PC)' or 'DAVIES, Richard\\n(MLP/PLM)' -> name and party."""
    flat = re.sub(r"\s+", " ", text).strip()
    match = re.fullmatch(r"(.*?)\s*\(([^()]*)\)", flat)
    if not match:
        raise UpstreamError(f"elections_provincial: Manitoba candidate heading {text!r}.")
    party_name, code = party(match.group(2))
    return AreaVote(name=_person(match.group(1)), party=party_name, party_code=code, votes=0)


def _wide(rows: list[list[object]]) -> list[Area] | None:
    """Areas of one division's wide sheet, or None if the sheet is not by voting area."""
    top = next(
        (
            i
            for i, row in enumerate(rows)
            if any(_header_key(c).startswith("rejected") for c in row)
        ),
        None,
    )
    if top is None:
        return None
    header = rows[top]
    area_col = _column(header, "votingarea")
    if area_col is None:
        return None  # 2023's sheet by polling place
    place_col = _column(header, "votingplace")
    # Some 2011 sheets head the place column "Voting Area" and leave the number's blank;
    # two (Spruce Woods, St. Vital) head both columns "Voting Area".
    if place_col is None:
        if area_col > 0 and not _cell(header[area_col - 1]):
            area_col, place_col = area_col - 1, area_col
        elif area_col + 1 < len(header) and _header_key(header[area_col + 1]).startswith(
            "votingarea"
        ):
            place_col = area_col + 1
    rejected_col = _column(header, "rejected")
    declined_col = _column(header, "declined")
    keys = [_header_key(c) for c in header]
    registered_col = next((i for i, k in enumerate(keys) if "registered" in k), None)
    assert rejected_col is not None
    first = max(area_col, place_col if place_col is not None else -1) + 1
    heads = [
        (i, _candidate_heading(_cell(header[i])))
        for i in range(first, rejected_col)
        if _cell(header[i])
    ]
    if not heads:
        raise UpstreamError("elections_provincial: Manitoba voting-area sheet has no candidates.")

    def at(row: list[object], col: int | None) -> object:
        return row[col] if col is not None and col < len(row) else None

    areas: list[Area] = []
    for row in rows[top + 1 :]:
        label = _cell(at(row, area_col))
        place = _cell(at(row, place_col))
        first_label = label or place
        if not first_label:
            continue
        if re.match(r"(final )?totals?\b", first_label.lower()):
            break
        # 2019 notes that ballots cast outside the division are counted with the write-in
        # ballots, in a text cell across the vote columns; such a row holds no counts.
        if any(_cell(at(row, col)) and not _is_number(at(row, col)) for col, _h in heads):
            continue
        area = Area(
            label=label or place,
            place=place or None,
            electors=_count(at(row, registered_col)) if registered_col is not None else None,
            rejected=_count(at(row, rejected_col)),
            declined=_count(at(row, declined_col)),
        )
        for col, head in heads:
            area.votes.append(
                AreaVote(head.name, head.party, head.party_code, _count(at(row, col)))
            )
        areas.append(area)
    return areas


def _long(rows: list[list[object]]) -> dict[str, list[Area]]:
    """Areas by division name from a long sheet (one row per voting area and candidate)."""
    header = rows[0]
    division_col = _column(header, "electoraldivision")
    area_col = _column(header, "pollid", "votingarea")
    place_col = _column(header, "pollname", "votingplace")
    name_col = _column(header, "candidatename")
    last_col = _column(header, "candidatelastname")
    first_col = _column(header, "candidatefirstname")
    party_col = _column(header, "partyabbrev", "politicalaffiliation")
    votes_col = _column(header, "totalvotes", "ballotscastforcandidate")
    rejected_col = _column(header, "rejected")
    declined_col = _column(header, "declined")
    registered_col = _column(header, "registered")
    if None in (division_col, area_col, party_col, votes_col) or (
        name_col is None and last_col is None
    ):
        raise UpstreamError(
            "elections_provincial: Manitoba voting-area file lacks the expected columns "
            f"(has {[_cell(c) for c in header]})."
        )
    assert division_col is not None and area_col is not None
    assert party_col is not None and votes_col is not None

    def at(row: list[object], col: int | None) -> object:
        return row[col] if col is not None and col < len(row) else None

    found: dict[str, dict[str, Area]] = {}
    for row in rows[1:]:
        # The 2007 files end with a legend of party names, which has no vote count.
        if not _is_number(at(row, votes_col)):
            continue
        # The 2003 file has rows with no candidate and no party (Interlake); the official
        # division totals do not count them.
        if not _cell(at(row, party_col)) and not _cell(at(row, name_col or last_col)):
            continue
        division = english(_cell(at(row, division_col)))
        label = _cell(at(row, area_col))
        if name_col is not None:
            name = _person(_cell(at(row, name_col)))
        else:
            name = f"{_cell(at(row, first_col))} {_cell(at(row, last_col))}".strip()
        by_label = found.setdefault(division, {})
        area = by_label.get(label)
        if area is None:
            # Rejected, declined and registered repeat on each candidate's row of an area.
            electors = at(row, registered_col)
            area = by_label[label] = Area(
                label=label,
                place=_cell(at(row, place_col)) or None,
                electors=_count(electors) if registered_col is not None else None,
                rejected=_count(at(row, rejected_col)),
                declined=_count(at(row, declined_col)),
            )
        party_name, code = party(_cell(at(row, party_col)))
        area.votes.append(AreaVote(name or None, party_name, code, _count(at(row, votes_col))))
    return {division: list(areas.values()) for division, areas in found.items()}


def parse_areas(body: bytes) -> dict[str, list[Area]]:
    """Areas by division from one workbook of results by voting area.

    A long sheet gives its division names; a wide workbook with several sheets gives one
    division per sheet, named by the sheet; a wide workbook for one division is keyed "".
    """
    found: dict[str, list[Area]] = {}
    wide: list[tuple[str, list[Area]]] = []
    for name, rows in _sheets(body):
        rows = [row for row in rows if any(_cell(c) for c in row)]
        if not rows:
            continue
        if _header_key(rows[0][0]).startswith("electoraldivision"):
            found.update(_long(rows))
            continue
        areas = _wide(rows)
        if areas is not None:
            wide.append((name, areas))
    if len(wide) == 1:
        found[""] = wide[0][1]
    else:
        found.update(wide)
    if not found:
        raise UpstreamError("elections_provincial: Manitoba voting-area file has no results.")
    return found


def member_key(path: str) -> str:
    """Division key of a zip member: 'Brandon_East_Candidate_by_VA.xlsx' -> 'brandoneast'."""
    stem = path.rsplit("/", 1)[-1]
    stem = re.split(r"_(?:votesreceived|candidate)", stem, maxsplit=1, flags=re.IGNORECASE)[0]
    return re.sub(r"[^a-z0-9]", "", fold(stem))


async def list_area_files(number: str) -> list[remote_zip.ZipMember]:
    """The workbooks in an election's voting-area zip, read from its directory by range."""
    await _LIMITER.acquire()
    members, _total = await remote_zip.list_members(files(number).by_area)
    return [m for m in members if not m.name.endswith("/")]


async def read_area_file(number: str, member: remote_zip.ZipMember) -> dict[str, list[Area]]:
    """One workbook of the voting-area zip, read by range and parsed."""
    await _LIMITER.acquire()
    body = await remote_zip.read_member(
        files(number).by_area, member, max_bytes=constants.MB_MAX_BYTES
    )
    try:
        return await run_parse(parse_areas, body)
    except UpstreamError:
        raise
    except Exception as exc:  # xlrd and openpyxl raise several unrelated types
        raise UpstreamError(
            f"elections_provincial: could not read {member.name} in {files(number).by_area}: {exc}"
        ) from exc
