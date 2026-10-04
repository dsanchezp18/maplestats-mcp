"""Saskatchewan general election results from Elections Saskatchewan's poll-by-poll files.

Checked live 2026-10-02 on https://www.elections.sk.ca/reports-data/election-results/,
which links one poll-by-poll file per general election on cdn.elections.sk.ca:

- 2024, 2020 and 2016: CSV, one row per poll (or advance, mobile, mail-in or care-home
  poll) with the registered voters, one vote column per party, the rejected ballots and
  one candidate-name column per party. The 2024 and 2016 files are Windows-1252, the 2020
  file is UTF-8 with a byte order mark. The 2024 file writes names as "Last, First", the
  others as "First Last". The 2020 file has vote counts with a thousands comma ("1,489").
  Column names differ by year ("Row Order" and "Row Ordering", "Poll Name" and "PollName",
  "Rejected" and "RejectedBallots", "BPSK" and "BP").
- 2011: one Excel workbook with one sheet per constituency (Appendix VI of the statement
  of votes): a header row starting "Poll", the candidate names, the party codes on the next
  row, one row per poll, then a "Totals" row whose cells are formulas. Totals are summed
  here from the poll rows instead.

Summing the poll rows by constituency gives 34 Saskatchewan Party and 27 NDP winners in 2024,
48 and 13 in 2020, 51 and 10 in 2016 and 49 and 9 in 2011, the seats the legislature was
sworn in with. Registered voters are not summed because split polls ("2 A/B") repeat the
same electors on each row, so turnout is not computed. By-elections have their own files and
are not read.
"""

from __future__ import annotations

import csv
import io
import re
from collections import defaultdict

from maplestats_mcp.modules.elections_provincial import constants
from maplestats_mcp.modules.elections_provincial.common import (
    Candidate,
    District,
    fetch_bytes,
    finish_shares,
    mark_winners,
)
from maplestats_mcp.shared.errors import UpstreamError
from maplestats_mcp.shared.executor import run_parse

# Names for the party codes in the files. Elections Saskatchewan's own notes spell out GP,
# LIB, NDP, PC, SP and WIP; the other codes are the party names it registers.
PARTY_NAMES: dict[str, str] = {
    "BP": "Buffalo Party of Saskatchewan",
    "BPSK": "Buffalo Party of Saskatchewan",
    "GP": "Green Party",
    "IND": "Independent",
    "LIB": "Liberal Party",
    "NDP": "New Democratic Party",
    "PC": "Progressive Conservative Party of Saskatchewan",
    "SGP": "Saskatchewan Green Party",
    "SP": "Saskatchewan Party",
    "SPP": "Saskatchewan Progress Party",
    "SUP": "Saskatchewan United Party",
    "WIP": "Western Independence Party",
}

# The 2011 workbook misspells one sheet's constituency.
_NAME_FIXES = {"Sasktoon Nutana": "Saskatoon Nutana"}

_TITLE = "vote summary by constituency polling division:"


def file_url(year: str) -> str:
    return constants.SK_FILES[year]


def _text(body: bytes) -> str:
    try:
        return body.decode("utf-8-sig")
    except UnicodeDecodeError:
        return body.decode("cp1252")


def _count(value: object) -> int:
    """A vote count: blank is zero and a thousands comma is dropped."""
    if value is None:
        return 0
    text = str(value).replace(",", "").strip()
    if not text:
        return 0
    try:
        return int(float(text))
    except ValueError as exc:
        raise UpstreamError(
            f"elections_provincial: Saskatchewan file has the count {value!r}."
        ) from exc


def _person(name: str) -> str:
    """'First Last' from either 'Last, First' or 'First Last', with spaces tidied."""
    last, comma, first = name.partition(",")
    flipped = f"{first.strip()} {last.strip()}" if comma else name
    return re.sub(r"\s+", " ", flipped).strip()


def _district_name(raw: str) -> str:
    name = re.sub(r"\s+", " ", raw).strip()
    if name.isupper():
        name = name.title()
    return _NAME_FIXES.get(name, name)


def _party(code: str) -> tuple[str, str]:
    return PARTY_NAMES.get(code, code), code


def parse_csv(body: bytes) -> list[District]:
    """Districts from one poll-by-poll CSV (2016, 2020 or 2024)."""
    reader = csv.DictReader(io.StringIO(_text(body)))
    # Header names with their spaces removed, so 'Poll Name' and 'PollName' match.
    names = {name: name.replace(" ", "") for name in reader.fieldnames or []}
    ordered = list(names.values())
    keys = set(ordered)
    district_key = "Constituency" if "Constituency" in keys else "ConstituencyName"
    rejected_key = "Rejected" if "Rejected" in keys else "RejectedBallots"
    parties = [k.removesuffix("Candidate") for k in ordered if k.endswith("Candidate")]
    needed = {district_key, rejected_key, *parties, *(f"{p}Candidate" for p in parties)}
    if not parties or not needed <= keys:
        raise UpstreamError(
            "elections_provincial: Saskatchewan poll file lacks the expected columns "
            f"(has {sorted(keys)})."
        )
    code_key = "ConstituencyCode"

    votes: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    candidates: dict[str, dict[str, str]] = defaultdict(dict)
    rejected: dict[str, int] = defaultdict(int)
    codes: dict[str, str] = {}
    for raw in reader:
        row = {names[k]: v for k, v in raw.items() if k in names}
        if not (row[district_key] or "").strip():
            continue
        district = _district_name(row[district_key])
        codes.setdefault(district, (row.get(code_key) or "").strip())
        rejected[district] += _count(row[rejected_key])
        for party in parties:
            votes[district][party] += _count(row[party])
            name = (row[f"{party}Candidate"] or "").strip()
            if name:
                candidates[district].setdefault(party, _person(name))

    return _build(votes, candidates, rejected, codes)


def _build(
    votes: dict[str, dict[str, int]],
    candidates: dict[str, dict[str, str]],
    rejected: dict[str, int],
    codes: dict[str, str],
) -> list[District]:
    built: list[District] = []
    for name, by_party in votes.items():
        district = District(
            name=name, number=codes.get(name) or None, rejected_ballots=rejected.get(name)
        )
        for code, total in by_party.items():
            person = candidates[name].get(code)
            # A party column with no candidate name and no votes means it ran nobody here.
            if person is None and total == 0:
                continue
            party, party_code = _party(code)
            district.candidates.append(
                Candidate(name=person, party=party, party_code=party_code, votes=total, share=None)
            )
        district.candidates.sort(key=lambda c: -c.votes)
        mark_winners(district)
        finish_shares(district)
        built.append(district)
    if not built:
        raise UpstreamError("elections_provincial: Saskatchewan poll file has no rows.")
    return built


def parse_xlsx(body: bytes) -> list[District]:
    """Districts from the 2011 poll results workbook (one sheet per constituency)."""
    # openpyxl takes ~0.7 s to import, so it loads on first use.
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    votes: dict[str, dict[str, int]] = {}
    candidates: dict[str, dict[str, str]] = {}
    rejected: dict[str, int] = {}
    try:
        for sheet in workbook:
            rows = list(sheet.iter_rows(values_only=True))
            title = next(
                (
                    str(cell)
                    for row in rows[:4]
                    for cell in row
                    if cell and _TITLE in str(cell).lower()
                ),
                None,
            )
            header = next((i for i, r in enumerate(rows) if str(r[0]).strip() == "Poll"), None)
            if title is None or header is None:
                raise UpstreamError(
                    f"elections_provincial: Saskatchewan 2011 sheet {sheet.title!r} has no "
                    "title or header row."
                )
            name = _district_name(title.split(":", 1)[1])
            names, codes = rows[header], rows[header + 1]
            # Candidates run from column E up to the "Total Counted" column; the rejected
            # ballots column follows it.
            end = next(i for i, c in enumerate(names) if c and "counted" in str(c).lower())
            party_totals: dict[str, int] = defaultdict(int)
            people: dict[str, str] = {}
            rejected_total = 0
            for row in rows[header + 2 :]:
                if str(row[1]).strip() == "Totals":
                    break
                if row[0] is None:
                    continue
                for i in range(4, end):
                    party_totals[str(codes[i]).strip()] += _count(row[i])
                rejected_total += _count(row[end + 1])
            for i in range(4, end):
                people[str(codes[i]).strip()] = _person(str(names[i]))
            votes[name], candidates[name], rejected[name] = party_totals, people, rejected_total
    finally:
        workbook.close()
    return _build(votes, candidates, rejected, {})


async def fetch(year: str) -> list[District]:
    """Districts for one Saskatchewan general election, by year."""
    url = file_url(year)
    body = await fetch_bytes(
        url, context=f"elections_provincial:sk:{year}", max_bytes=constants.SK_MAX_BYTES
    )
    parser = parse_xlsx if url.endswith(".xlsx") else parse_csv
    try:
        return await run_parse(parser, body)
    except UpstreamError:
        raise
    except Exception as exc:  # csv and openpyxl raise several unrelated types
        raise UpstreamError(f"elections_provincial: could not read {url}: {exc}") from exc
