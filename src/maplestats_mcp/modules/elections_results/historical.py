"""Constituency-level results of every federal general election, 1867 to 2015.

Source: Stanley L. Winer and J. Stephen Ferris (with Haizhen Mou, Derek E. H.
Olmstead and Jérôme Archambault), "Data Set on Federal Elections, With
Superconstituencies, Canada 1867 - 2015, Elections 1 - 42", Scholars Portal
Dataverse (Borealis), https://doi.org/10.5683/SP2/1N4Y1G, licence CC0 1.0.
Checked live 2026-09-30: one 3.3 MB workbook whose sheet "Canada 1867-2015
(raw data)" holds 10,585 constituency-election rows, each with up to 13
party and vote pairs (k1/v1 to k13/v13, largest first), and whose sheet
"Party Names and History" decodes the party mnemonics (125 of them).
"""

from __future__ import annotations

import io
import unicodedata
import warnings
from datetime import date

import httpx
from pydantic import BaseModel, Field

from maplestats_mcp.shared.cache import cached_fetch
from maplestats_mcp.shared.envelope import make_provenance
from maplestats_mcp.shared.errors import InvalidInput, UpstreamError, UpstreamUnavailable
from maplestats_mcp.shared.executor import run_parse
from maplestats_mcp.shared.fr_typography import fr_or_en, lang_error
from maplestats_mcp.shared.http import get_raw
from maplestats_mcp.shared.limits import join_limits
from maplestats_mcp.shared.models import Provenance
from maplestats_mcp.shared.rate_limiter import get_limiter

DATASET_DOI = "https://doi.org/10.5683/SP2/1N4Y1G"
DATASET_FILE_URL = "https://borealisdata.ca/api/access/datafile/86556"
RAW_SHEET = "Canada 1867-2015 (raw data)"
PARTY_SHEET = "Party Names and History"
FIRST_ELECTION = 1
LAST_ELECTION = 42
CACHE_TTL_SECONDS = 24 * 60 * 60
MAX_FILE_BYTES = 30 * 1024 * 1024
ROWS_LIMIT_DEFAULT = 50
ROWS_LIMIT_MAX = 500
_PAIRS = 13

_LIMITER = get_limiter("borealis-elections", rate=0.5, capacity=2.0)


class PartyVotes(BaseModel):
    party: str = Field(description="The dataset's party mnemonic, e.g. 'Lib', 'C', 'NDP'.")
    party_name: str | None = Field(description="Full name from the dataset's party sheet.")
    votes: int


class HistoricalRiding(BaseModel):
    election: int = Field(description="General election number, 1 (1867) to 42 (2015).")
    election_date: date
    province: str
    constituency: str
    electors: int | None
    rejected_ballots: int | None
    ballots_cast: int | None
    candidates: int | None
    seats: int | None = Field(description="Seats filled here (2 in multi-member ridings).")
    acclamation: bool
    leading_party: str | None = Field(
        description="Party with the most votes: the winner in a single-member riding with one "
        "candidate per party; null for an acclamation. Not a recorded winner."
    )
    results: list[PartyVotes] = Field(description="Parties and votes, largest first.")


class HistoricalResult(BaseModel):
    election: int | None
    ridings: list[HistoricalRiding]
    total_ridings: int
    offset: int
    truncated: bool
    licence: str
    citation: str
    provenance: Provenance


def _fold(text: str) -> str:
    stripped = unicodedata.normalize("NFKD", text)
    return "".join(c for c in stripped if not unicodedata.combining(c)).casefold().strip()


def _int(value: object) -> int | None:
    return None if value is None or value == "" else int(float(str(value)))


def _parse(body: bytes) -> tuple[list[dict[str, object]], dict[str, str]]:
    # openpyxl takes ~0.7 s to import, so it loads on first use.
    from openpyxl import load_workbook

    with warnings.catch_warnings():
        # The workbook carries an invalid print-area name that is irrelevant here.
        warnings.simplefilter("ignore", UserWarning)
        workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    try:
        names: dict[str, str] = {}
        for row in list(workbook[PARTY_SHEET].iter_rows(values_only=True))[3:]:
            if len(row) > 7 and row[6] and row[7]:
                names[str(row[6]).strip()] = str(row[7]).strip()

        rows = [
            r
            for r in workbook[RAW_SHEET].iter_rows(values_only=True)
            if any(c is not None for c in r)
        ]
    finally:
        workbook.close()
    header = [str(c) if c is not None else "" for c in rows[0]]
    return [dict(zip(header, r, strict=False)) for r in rows[1:]], names


async def _dataset(
    lang: str = "en",
) -> tuple[tuple[list[dict[str, object]], dict[str, str]], bool]:
    async def fetch() -> tuple[list[dict[str, object]], dict[str, str]]:
        await _LIMITER.acquire()
        # Borealis answers 303 with a signed storage URL; the shared client does
        # not follow redirects.
        current = DATASET_FILE_URL
        for _ in range(4):
            try:
                response = await get_raw(current, timeout=180.0)
                break
            except httpx.HTTPStatusError as exc:
                location = exc.response.headers.get("location")
                if exc.response.status_code in (301, 302, 303, 307, 308) and location:
                    current = str(exc.response.url.join(location))
                    if not current.startswith("https://"):
                        raise lang_error(
                            UpstreamError,
                            lang,
                            "elections_results: redirect to a non-https URL.",
                            "elections_results : redirection vers une adresse non https.",
                        ) from exc
                    continue
                raise lang_error(
                    UpstreamError,
                    lang,
                    f"elections_results: Borealis returned HTTP {exc.response.status_code}.",
                    f"elections_results : Borealis a renvoyé HTTP {exc.response.status_code}.",
                ) from exc
            except httpx.HTTPError as exc:
                raise lang_error(
                    UpstreamUnavailable,
                    lang,
                    "elections_results: Borealis did not respond in time.",
                    "elections_results : Borealis n'a pas répondu à temps.",
                ) from exc
        else:
            raise lang_error(
                UpstreamError,
                lang,
                "elections_results: Borealis redirected too many times.",
                "elections_results : Borealis a redirigé trop de fois.",
            )
        if len(response.content) > MAX_FILE_BYTES:
            raise lang_error(
                UpstreamError,
                lang,
                "elections_results: the historical workbook is larger than expected.",
                "elections_results : le classeur historique est plus volumineux que prévu.",
            )
        try:
            return await run_parse(_parse, response.content)
        except Exception as exc:  # openpyxl raises several unrelated types
            raise lang_error(
                UpstreamError,
                lang,
                f"elections_results: could not read the workbook: {exc}",
                f"elections_results : lecture du classeur impossible : {exc}",
            ) from exc

    return await cached_fetch("elections_results:historical", CACHE_TTL_SECONDS, fetch)


def _riding(row: dict[str, object], names: dict[str, str]) -> HistoricalRiding:
    results: list[PartyVotes] = []
    for i in range(1, _PAIRS + 1):
        party, votes = row.get(f"k{i}"), row.get(f"v{i}")
        if party is None or _int(votes) in (None, 0):
            continue
        key = str(party).strip()
        results.append(PartyVotes(party=key, party_name=names.get(key), votes=_int(votes) or 0))
    results.sort(key=lambda r: -r.votes)
    stamp = str(_int(row["election_date"]))
    acclaimed = bool(_int(row.get("acclamation")))
    return HistoricalRiding(
        election=_int(row["parliament_number"]) or 0,
        election_date=date(int(stamp[:4]), int(stamp[4:6]), int(stamp[6:8])),
        province=str(row.get("province_name") or ""),
        constituency=str(row.get("constituency_name") or ""),
        electors=_int(row.get("electors_local")),
        rejected_ballots=_int(row.get("rejected_local")),
        ballots_cast=_int(row.get("ballots_local")),
        candidates=_int(row.get("candidates_local")),
        seats=_int(row.get("seats_local")),
        acclamation=acclaimed,
        leading_party=None if acclaimed or not results else results[0].party,
        results=results,
    )


async def get_historical(
    election: int | None = None,
    province: str | None = None,
    constituency: str | None = None,
    party: str | None = None,
    limit: int = ROWS_LIMIT_DEFAULT,
    offset: int = 0,
    lang: str = "en",
) -> HistoricalResult:
    if election is not None and not FIRST_ELECTION <= election <= LAST_ELECTION:
        raise lang_error(
            InvalidInput,
            lang,
            f"elections_results: election must be {FIRST_ELECTION} to {LAST_ELECTION} "
            "(1867 to 2015) for historical results.",
            f"elections_results : election doit être entre {FIRST_ELECTION} et {LAST_ELECTION} "
            "(1867 à 2015) pour les résultats historiques.",
        )
    if not 1 <= limit <= ROWS_LIMIT_MAX:
        raise lang_error(
            InvalidInput,
            lang,
            f"elections_results: limit must be 1 to {ROWS_LIMIT_MAX}.",
            f"elections_results : limit doit être compris entre 1 et {ROWS_LIMIT_MAX}.",
        )
    if offset < 0:
        raise lang_error(
            InvalidInput,
            lang,
            "elections_results: offset must be 0 or more.",
            "elections_results : offset doit être 0 ou plus.",
        )

    (rows, names), cached = await _dataset(lang)
    if election is not None:
        rows = [r for r in rows if _int(r.get("parliament_number")) == election]
    if province:
        wanted = _fold(province)
        rows = [r for r in rows if wanted in _fold(str(r.get("province_name") or ""))]
    if constituency:
        wanted = _fold(constituency)
        rows = [
            r
            for r in rows
            if wanted in _fold(str(r.get("constituency_name") or ""))
            or wanted in _fold(str(r.get("old_name") or ""))
        ]
    if party:
        wanted = _fold(party)

        def has_party(row: dict[str, object]) -> bool:
            for i in range(1, _PAIRS + 1):
                key = row.get(f"k{i}")
                if key is None:
                    continue
                mnemonic = str(key).strip()
                if wanted == _fold(mnemonic) or wanted in _fold(names.get(mnemonic, "")):
                    return True
            return False

        rows = [r for r in rows if has_party(r)]

    total = len(rows)
    page = [_riding(r, names) for r in rows[offset : offset + limit]]
    return HistoricalResult(
        election=election,
        ridings=page,
        total_ridings=total,
        offset=offset,
        truncated=offset + limit < total,
        licence=fr_or_en(
            lang,
            "CC0 1.0 (public domain dedication)",
            "CC0 1.0 (dévouement au domaine public)",
        ),
        citation=(
            "Winer, S. L. and Ferris, J. S., with Mou, H., Olmstead, D. E. H. and Archambault, J. "
            "Data Set on Federal Elections, With Superconstituencies, Canada 1867 - 2015, "
            f"Elections 1 - 42. Scholars Portal Dataverse, {DATASET_DOI}."
        ),
        provenance=make_provenance(
            source="borealis-winer-ferris-federal-elections",
            url=DATASET_DOI,
            cached=cached,
            schema_name="elections_results.HistoricalResult",
            freshness=fr_or_en(
                lang,
                "A 2019 release covering the 1st to 42nd general elections; no later "
                "elections. Use the 38th to 45th tables for official Elections Canada results.",
                "Version de 2019 couvrant les 1re à 42e élections générales, sans élections "
                "ultérieures. Utilisez les tableaux des 38e à 45e pour les résultats officiels "
                "d'Élections Canada.",
            ),
            coverage=fr_or_en(
                lang,
                "Party totals per constituency, not candidate names. Vote counts were "
                "corrected by the authors where the source figures were inconsistent.",
                "Totaux par parti et par circonscription, sans noms de candidats. Les auteurs "
                "ont corrigé les votes là où les chiffres de la source étaient incohérents.",
            ),
            limits=join_limits(
                fr_or_en(
                    lang,
                    f"Showing rows {offset + 1} to {offset + len(page)} of {total}.",
                    f"Lignes {offset + 1} à {offset + len(page)} sur {total}.",
                )
                if offset + limit < total
                else None,
                fr_or_en(
                    lang,
                    "",
                    "Les noms de circonscriptions, de partis et de provinces viennent du jeu de données, en anglais seulement",
                )
                or None,
            ),
            lang=lang,
        ),
    )
