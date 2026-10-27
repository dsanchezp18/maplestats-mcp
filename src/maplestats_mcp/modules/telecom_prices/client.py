"""Client for the CRTC Communications Market Reports workbooks.

Checked live 2026-10-27:

1. The CRTC's 95 open.canada.ca datasets include the year-end plan price
   survey (2016-2025, version 11 of 2026-05-14, although the dataset title
   still says 2016-2022), the retail mobile and retail fixed Internet
   sector tables (2013-2024, titles say 2013-2022) and the quarterly
   current trends (2014 Q1 to 2026 Q1, version 17 of 2026-08-21). None has
   a DataStore. Their "CSV" resources are ZIP archives of one CSV per sheet
   in Windows-1252, which ckan_read_resource refuses; the XLSX resources
   are read here.
2. The files are on applications.crtc.gc.ca, which answers without the
   Cloudflare challenge that blocks crtc.gc.ca (403 "Security Check" on
   every page tried, including the low-cost plan page and the terms). The
   host reset a connection once in five downloads, so downloads are
   retried by the shared reader's caller (cached once read).
3. Resource ids have stayed the same across updates (the price XLSX was
   replaced in place in February and May 2026); a resource that is gone is
   looked up again by format and language in the dataset.
4. The sector workbooks exist in English and French with different tab
   names for figures (MB-F12 and MB-G12); the price and quarterly
   workbooks are bilingual, one file for both languages.
"""

from __future__ import annotations

import io
import re
import unicodedata
from collections.abc import Callable
from datetime import UTC, date, datetime
from typing import Any, NoReturn

import httpx

from maplestats_mcp.modules.telecom_prices import constants, parsing
from maplestats_mcp.modules.telecom_prices.schemas import (
    CityPlanPriceAverage,
    CityPlanPriceRow,
    CityPlanPrices,
    CmaCoverage,
    MarketObservation,
    MarketTable,
    MarketTableList,
    MarketTableRef,
    PlanPricePoint,
    PlanPrices,
    QuarterlyTrends,
    TrendObservation,
)
from maplestats_mcp.shared import file_download
from maplestats_mcp.shared.cache import cached_fetch
from maplestats_mcp.shared.envelope import make_provenance, raise_localized
from maplestats_mcp.shared.errors import InvalidInput, NotFound, UpstreamError, UpstreamUnavailable
from maplestats_mcp.shared.executor import run_parse
from maplestats_mcp.shared.http import api_get
from maplestats_mcp.shared.i18n import french_spacing, normalize_lang
from maplestats_mcp.shared.json_utils import list_or_empty
from maplestats_mcp.shared.limits import join_limits
from maplestats_mcp.shared.rate_limiter import get_limiter

_LIMITER = get_limiter(
    constants.RATE_LIMIT_SOURCE,
    rate=constants.RATE_LIMIT_PER_SECOND,
    capacity=constants.RATE_LIMIT_CAPACITY,
)


def _raise(exc_cls: type[ValueError], en: str, fr: str, lang: str) -> NoReturn:
    raise_localized(exc_cls, en, fr, lang)


def _text(en: str, fr: str, lang: str) -> str:
    return french_spacing(fr) if lang == "fr" else en


# Fetching ---------------------------------------------------------------------


async def _package(workbook: constants.Workbook, lang: str) -> tuple[dict[str, Any], bool]:
    async def fetch() -> dict[str, Any]:
        await _LIMITER.acquire()
        try:
            body = await api_get(
                constants.CKAN_PACKAGE_SHOW, params={"id": workbook.dataset_id}, timeout=30.0
            )
        except httpx.HTTPStatusError as exc:
            if exc.response.status_code == 404:
                _raise(
                    NotFound,
                    f"telecom_prices: the CRTC dataset {workbook.dataset_id} is no longer on "
                    "open.canada.ca.",
                    f"telecom_prices : le jeu de données du CRTC {workbook.dataset_id} n'est "
                    "plus sur ouvert.canada.ca.",
                    lang,
                )
            _raise(
                UpstreamError,
                f"telecom_prices: open.canada.ca answered HTTP {exc.response.status_code}.",
                f"telecom_prices : ouvert.canada.ca a répondu HTTP {exc.response.status_code}.",
                lang,
            )
        except httpx.HTTPError as exc:
            _raise(
                UpstreamUnavailable,
                f"telecom_prices: open.canada.ca did not respond ({type(exc).__name__}).",
                f"telecom_prices : ouvert.canada.ca n'a pas répondu ({type(exc).__name__}).",
                lang,
            )
        result = body.get("result") if isinstance(body, dict) else None
        if not isinstance(result, dict):
            _raise(
                UpstreamError,
                "telecom_prices: open.canada.ca returned no dataset record.",
                "telecom_prices : ouvert.canada.ca n'a renvoyé aucune fiche de jeu de données.",
                lang,
            )
        return result

    return await cached_fetch(
        f"telecom_prices:package:{workbook.dataset_id}", constants.PACKAGE_TTL_SECONDS, fetch
    )


def pick_resource(package: dict[str, Any], resource_id: str, lang: str) -> dict[str, Any]:
    """The XLSX resource by id, or else the dataset's XLSX in that language."""
    resources = [r for r in list_or_empty(package, "resources") if isinstance(r, dict)]
    for resource in resources:
        if resource.get("id") == resource_id:
            return resource
    xlsx = [r for r in resources if str(r.get("format") or "").upper() in ("XLSX", "XLS")]
    for resource in xlsx:
        if lang in list_or_empty(resource, "language"):
            return resource
    if xlsx:
        return xlsx[0]
    _raise(
        NotFound,
        "telecom_prices: the CRTC dataset no longer lists an Excel file.",
        "telecom_prices : le jeu de données du CRTC ne contient plus de fichier Excel.",
        lang,
    )


def _title(package: dict[str, Any], workbook: constants.Workbook, lang: str) -> str:
    titles = package.get("title_translated") or {}
    found = titles.get(lang) if isinstance(titles, dict) else None
    return str(found or (workbook.title_fr if lang == "fr" else workbook.title_en))


async def _download(url: str, lang: str) -> bytes:
    async def fetch() -> file_download.Downloaded:
        return await file_download.download(
            url,
            allow_host=lambda host: host == constants.FILE_HOST,
            limiter_for=lambda _host: _LIMITER,
            max_bytes=constants.MAX_FILE_BYTES,
            context="telecom_prices",
            timeout=90.0,
        )

    try:
        downloaded, _ = await file_download.cached_download(
            file_download.cache_key(url), constants.FILE_TTL_SECONDS, fetch
        )
    except UpstreamUnavailable:
        # The file host reset one connection in five on 2026-10-27 and then
        # answered the same URL at once; one more try is cheaper than a failure.
        downloaded, _ = await file_download.cached_download(
            file_download.cache_key(url), constants.FILE_TTL_SECONDS, fetch
        )
    if not downloaded.body.startswith(b"PK"):
        _raise(
            UpstreamError,
            f"telecom_prices: {url} is not an Excel workbook.",
            f"telecom_prices : {url} n'est pas un classeur Excel.",
            lang,
        )
    return downloaded.body


def read_sheets(body: bytes, wanted: Callable[[str], bool]) -> dict[str, list[tuple[Any, ...]]]:
    """The rows of every sheet whose name `wanted` accepts; key "" holds the sheet names."""
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    try:
        sheets: dict[str, list[tuple[Any, ...]]] = {"": [tuple(book.sheetnames)]}
        for name in book.sheetnames:
            if wanted(name):
                sheets[name] = [tuple(r) for r in book[name].iter_rows(values_only=True)]
        return sheets
    finally:
        book.close()


async def _workbook(
    workbook: constants.Workbook,
    lang: str,
    wanted: Callable[[str], bool],
    key: str,
) -> tuple[dict[str, list[tuple[Any, ...]]], dict[str, Any], dict[str, Any], bool]:
    package, package_cached = await _package(workbook, lang)
    resource_id = workbook.resource_fr if lang == "fr" else workbook.resource_en
    resource = pick_resource(package, resource_id, lang)
    url = str(resource.get("url") or "")

    async def fetch() -> dict[str, list[tuple[Any, ...]]]:
        body = await _download(url, lang)
        try:
            return await run_parse(read_sheets, body, wanted)
        except UpstreamError:
            raise  # the parse time budget (shared/executor.py) says so itself
        except Exception as exc:  # noqa: BLE001 (openpyxl raises several unrelated types)
            _raise(
                UpstreamError,
                f"telecom_prices: could not read {url} ({type(exc).__name__}).",
                f"telecom_prices : impossible de lire {url} ({type(exc).__name__}).",
                lang,
            )

    sheets, cached = await cached_fetch(
        f"telecom_prices:sheets:{resource.get('id')}:{key}", constants.FILE_TTL_SECONDS, fetch
    )
    return sheets, package, resource, cached and package_cached


def _as_of(released: date | None, resource: dict[str, Any]) -> datetime | None:
    if released is not None:
        return datetime(released.year, released.month, released.day, tzinfo=UTC)
    stamp = str(resource.get("last_modified") or resource.get("created") or "")
    try:
        return datetime.fromisoformat(stamp).replace(tzinfo=UTC) if stamp else None
    except ValueError:
        return None


def _provenance(
    url: str,
    cached: bool,
    schema: str,
    lang: str,
    *,
    as_of: datetime | None,
    freshness: str,
    coverage: str | None = None,
    limits: str | None = None,
) -> Any:
    return make_provenance(
        source=constants.SOURCE,
        url=url,
        cached=cached,
        schema_name=f"telecom_prices.{schema}",
        as_of=as_of,
        freshness=freshness,
        coverage=coverage,
        limits=limits,
        lang=lang,
    )


def _dataset_url(workbook: constants.Workbook, lang: str) -> str:
    return constants.DATASET_PAGE.format(lang=lang, dataset_id=workbook.dataset_id)


def _check_limit(limit: int, lang: str) -> None:
    if not 1 <= limit <= constants.ROWS_MAX_LIMIT:
        _raise(
            InvalidInput,
            f"telecom_prices: limit must be 1 to {constants.ROWS_MAX_LIMIT}.",
            f"telecom_prices : limit doit être compris entre 1 et {constants.ROWS_MAX_LIMIT}.",
            lang,
        )


def _check_years(start: int | None, end: int | None, lang: str) -> None:
    if start is not None and end is not None and start > end:
        _raise(
            InvalidInput,
            f"telecom_prices: from_year ({start}) is after to_year ({end}).",
            f"telecom_prices : from_year ({start}) est postérieure à to_year ({end}).",
            lang,
        )


# Plan prices -----------------------------------------------------------------

_GEOGRAPHY_ALIASES = {
    "canada": "Canada",
    "ca": "Canada",
    "national": "Canada",
    "provinces": "Provinces",
    "territories": "Territories",
    "territoires": "Territories",
    "north": "North",
    "nord": "North",
    "south": "South",
    "sud": "South",
}


def _split(value: str | list[str] | None) -> list[str]:
    """A list, or a comma-separated string, as its non-empty items."""
    items = value if isinstance(value, list) else (value or "").split(",")
    return [" ".join(str(item).split()) for item in items if str(item).strip()]


def _geography(value: str, lang: str) -> str:
    wanted = " ".join(value.split())
    if wanted.upper() in constants.PROVINCES:
        return wanted.upper()
    folded = parsing.text(wanted).casefold()
    if folded in _GEOGRAPHY_ALIASES:
        return _GEOGRAPHY_ALIASES[folded]
    for code, names in constants.PROVINCES.items():
        if folded in (names[0].casefold(), names[1].casefold()):
            return code
    _raise(
        InvalidInput,
        f"telecom_prices: unknown province {value!r}; use a code like ON, 'Canada', "
        "'Provinces', 'Territories', 'North' or 'South'.",
        f"telecom_prices : province inconnue {value!r} ; utilisez un code comme ON, "
        "'Canada', 'Provinces', 'Territories', 'North' ou 'South'.",
        lang,
    )


def plan_key(value: str) -> str:
    """'5GB', '5 Go', '5' and '5 gb' are one plan; 'talk', 'voix' the talk-only plan."""
    folded = parsing.text(value).casefold().replace(",", "")
    if any(word in folded for word in ("talk", "voice", "voix")):
        return "talk"
    digits = "".join(ch for ch in folded if ch.isdigit() or ch in "/.")
    return digits or folded


_METHOD = (
    "Each year the CRTC asks providers for the lowest non-promotional monthly price, before "
    "taxes and fees, that meets each service basket (a minimum data allowance, minutes and "
    "texts for mobile, bring your own device; a speed tier for Internet) in 24 urban centres "
    "and 55 rural communities. Provincial averages are simple averages over communities; the "
    "national average is the mean of the urban and rural averages. Baskets changed: before "
    "2019 the 1 GB plan required 450 minutes and 300 texts, and plans of 10 GB or more were "
    "added from 2019."
)
_METHOD_FR = (
    "Chaque année, le CRTC demande aux fournisseurs le prix mensuel non promotionnel le plus "
    "bas, avant taxes et frais, qui répond à chaque ensemble de services (une quantité "
    "minimale de données, de minutes et de textos pour le mobile, avec votre propre "
    "appareil; une vitesse pour Internet) dans 24 centres urbains et 55 collectivités "
    "rurales. Les moyennes provinciales sont des moyennes simples des collectivités; la "
    "moyenne nationale est la moyenne des moyennes urbaine et rurale. Les ensembles ont "
    "changé : avant 2019, le forfait de 1 Go exigeait 450 minutes et 300 textos, et les "
    "forfaits de 10 Go et plus ont été ajoutés à partir de 2019."
)


async def get_plan_prices(
    service: str = "mobile",
    *,
    plan: str | None = None,
    province: str | list[str] = "Canada",
    region: str = "national",
    technology: str = "lowest",
    from_year: int | None = None,
    to_year: int | None = None,
    limit: int = constants.ROWS_DEFAULT_LIMIT,
    lang: str = "en",
) -> PlanPrices:
    """Year-end average advertised prices of mobile or Internet plans."""
    lang = normalize_lang(lang)
    _check_limit(limit, lang)
    _check_years(from_year, to_year, lang)
    if service not in constants.PRICE_SHEETS:
        _raise(
            InvalidInput,
            "telecom_prices: service must be 'mobile' or 'internet'.",
            "telecom_prices : service doit être 'mobile' ou 'internet'.",
            lang,
        )
    if region not in ("national", "urban", "rural", "all"):
        _raise(
            InvalidInput,
            "telecom_prices: region must be 'national', 'urban', 'rural' or 'all'.",
            "telecom_prices : region doit être 'national', 'urban', 'rural' ou 'all'.",
            lang,
        )
    tech = technology.strip().lower()
    if service == "mobile" and tech not in ("lowest", "3g", "lte", "5g", "all"):
        _raise(
            InvalidInput,
            "telecom_prices: technology must be 'lowest', '3G', 'LTE', '5G' or 'all'.",
            "telecom_prices : technology doit être 'lowest', '3G', 'LTE', '5G' ou 'all'.",
            lang,
        )
    wanted_places = _split(province)
    if not wanted_places:
        _raise(
            InvalidInput,
            "telecom_prices: province is empty; give a code like ON, a list, or 'all'.",
            "telecom_prices : province est vide ; donnez un code comme ON, une liste ou 'all'.",
            lang,
        )
    geographies = (
        None
        if any(p.lower() == "all" for p in wanted_places)
        else {_geography(p, lang) for p in wanted_places}
    )
    sheet = constants.PRICE_SHEETS[service]
    sheets, _, resource, cached = await _workbook(
        constants.PRICES, lang, lambda n: n in (sheet, "Version"), service
    )
    if sheet not in sheets:
        _raise(
            UpstreamError,
            f"telecom_prices: the price workbook has no {sheet!r} sheet.",
            f"telecom_prices : le classeur des prix n'a pas de feuille {sheet!r}.",
            lang,
        )
    try:
        rows = await run_parse(parsing.parse_prices, sheets[sheet], service)
    except ValueError as exc:
        _raise(
            UpstreamError,
            f"telecom_prices: {exc}.",
            f"telecom_prices : la feuille des prix n'a pas la forme attendue ({exc}).",
            lang,
        )
    version, released = parsing.parse_version(sheets.get("Version", []))
    plans_en = sorted({r.plan_en for r in rows}, key=_plan_order)
    wanted_plan = plan_key(plan) if plan else None
    if wanted_plan and wanted_plan not in {plan_key(p) for p in plans_en}:
        _raise(
            NotFound,
            f"telecom_prices: no {service} plan {plan!r}; plans are {plans_en}.",
            f"telecom_prices : aucun forfait {service} {plan!r} ; les forfaits sont {plans_en}.",
            lang,
        )
    matched = [
        r
        for r in rows
        if (wanted_plan is None or plan_key(r.plan_en) == wanted_plan)
        and (geographies is None or r.geography in geographies)
        and (region == "all" or r.region == region)
        and (service != "mobile" or tech == "all" or (r.technology or "").lower() == tech)
        and (from_year is None or r.year >= from_year)
        and (to_year is None or r.year <= to_year)
    ]
    matched.sort(key=lambda r: (_plan_order(r.plan_en), r.geography, r.region, r.year))
    points = [
        PlanPricePoint(
            year=r.year,
            plan=r.plan_fr if lang == "fr" else r.plan_en,
            geography=r.geography,
            region=r.region,  # type: ignore[arg-type]
            technology=r.technology,
            average_price=r.value,
            as_published=r.as_published,
        )
        for r in matched[:limit]
    ]
    years = sorted({r.year for r in rows})
    cut = len(matched) - len(points)
    return PlanPrices(
        service=service,  # type: ignore[arg-type]
        dataset=_text(constants.PRICES.title_en, constants.PRICES.title_fr, lang),
        dataset_url=_dataset_url(constants.PRICES, lang),
        version=version,
        released=released,
        points=points,
        returned_count=len(points),
        total_matched=len(matched),
        plans=[(_plan_fr_label(p) if lang == "fr" else p) for p in plans_en],
        years=years,
        method=_text(_METHOD, _METHOD_FR, lang),
        provenance=_provenance(
            str(resource.get("url") or ""),
            cached,
            "PlanPrices",
            lang,
            as_of=_as_of(released, resource),
            freshness=_text(
                "yearly: prices at December 31, published the following spring",
                "annuelle : prix au 31 décembre, publiés le printemps suivant",
                lang,
            ),
            coverage=_text(
                f"{len(rows):,} {service} averages, {years[0]}-{years[-1]}" if years else "no rows",
                f"{len(rows):,} moyennes ({service}), {years[0]}-{years[-1]}"
                if years
                else "aucune ligne",
                lang,
            ),
            limits=join_limits(
                _text(
                    f"first {len(points)} of {len(matched)} matching averages; narrow the "
                    "filters or raise limit",
                    f"{len(points)} premières moyennes sur {len(matched)} ; précisez les "
                    "filtres ou augmentez limit",
                    lang,
                )
                if cut
                else None,
                _text(
                    "Averages of advertised list prices, not what subscribers pay: no "
                    "promotions, bundles or taxes.",
                    "Moyennes de prix annoncés, et non de ce que paient les abonnés : sans "
                    "promotions, offres groupées ni taxes.",
                    lang,
                ),
            ),
        ),
    )


def _plan_order(plan: str) -> tuple[int, float, str]:
    digits = plan_key(plan).split("/")[0]
    try:
        return (1, float(digits), plan)
    except ValueError:
        return (0, 0.0, plan)


def _plan_fr_label(plan_en: str) -> str:
    if plan_en.lower() == "talk plan":
        return "Forfait voix"
    return plan_en.replace(" GB", " Go")


# City (census metropolitan area) prices ---------------------------------------

_REPORTED_SHEET = "Prix-Prices"

_CITY_BASIS = (
    "Computed by MapleStats: mean of the reported monthly prices in the CRTC's Prix-Prices "
    "sheet. The CRTC publishes no city averages; its provincial and national averages "
    "(telecom_get_plan_prices) are its own."
)
_CITY_BASIS_FR = (
    "Calculé par MapleStats : moyenne des prix mensuels déclarés dans la feuille Prix-Prices "
    "du CRTC. Le CRTC ne publie pas de moyennes par ville; ses moyennes provinciales et "
    "nationales (telecom_get_plan_prices) sont les siennes."
)
_CITY_METHOD = (
    "Each year the CRTC asks providers for the lowest non-promotional monthly price, before "
    "taxes and fees, that meets each service basket, in 24 census metropolitan areas (CMAs) "
    "and 55 rural communities; from 2024 it surveys 13 CMAs, one per province or "
    "territory. The Prix-Prices sheet lists every reported price by provider, brand and, for "
    "mobile, network (3G, LTE, 5G). average_price is the plain mean of those rows for a year, "
    "plan and CMA, so a provider reporting several brands or networks counts several times; "
    "lowest_per_provider_average first keeps each provider's lowest price, which is how the "
    "CRTC computes its published 'lowest price' averages (in 2024 the Alberta urban average, "
    "$37.00 for 50 GB, equals Edmonton's lowest_per_provider_average). Rows without a price "
    "are left out."
)
_CITY_METHOD_FR = (
    "Chaque année, le CRTC demande aux fournisseurs le prix mensuel non promotionnel le plus "
    "bas, avant taxes et frais, qui répond à chaque ensemble de services, dans 24 régions "
    "métropolitaines de recensement (RMR) et 55 collectivités rurales; depuis 2024, il "
    "sonde 13 RMR, une par province ou territoire. La feuille Prix-Prices donne chaque "
    "prix déclaré par fournisseur, marque et, pour le mobile, réseau (3G, LTE, 5G). "
    "average_price est la moyenne simple de ces lignes pour une année, un forfait et une RMR : "
    "un fournisseur qui déclare plusieurs marques ou réseaux compte donc plusieurs fois; "
    "lowest_per_provider_average garde d'abord le prix le plus bas de chaque fournisseur, "
    "comme le CRTC pour ses moyennes publiées du « prix le plus bas » (en 2024, la moyenne "
    "urbaine de l'Alberta, 37,00 $ pour 50 Go, égale la lowest_per_provider_average "
    "d'Edmonton). Les lignes sans prix sont exclues."
)


def place_key(value: str) -> str:
    """'Montreal', 'MONTRÉAL' and 'montréal' are one place; other punctuation is a space."""
    decomposed = unicodedata.normalize("NFKD", value)
    plain = "".join(ch for ch in decomposed if not unicodedata.combining(ch)).casefold()
    # "St. John's", "St Johns" and "Saint John's" are one name.
    plain = " ".join(re.sub(r"[^a-z0-9]+", " ", re.sub(r"['’.]", "", plain)).split())
    plain = re.sub(r"^saint ", "st ", plain)
    for prefix in ("ville de ", "city of ", "region de ", "greater "):
        plain = plain.removeprefix(prefix)
    return plain.removesuffix(" city").removesuffix(" cma").removesuffix(" rmr")


def match_cmas(query: str, cmas: list[str]) -> list[str]:
    """The CMAs a name matches: the whole name or one side of a hyphen.

    'Ottawa' and 'Gatineau' find 'Ottawa-Gatineau', 'Niagara' finds
    'St Catharines - Niagara', 'St Johns' finds "St. John's".
    """
    wanted = place_key(query)
    if not wanted:
        return []
    exact = [c for c in cmas if place_key(c) == wanted]
    if exact:
        return exact
    return [c for c in cmas if wanted in {place_key(p) for p in re.split(r"-", c)}]


def _years_span(years: list[int]) -> str:
    if not years:
        return ""
    if years == list(range(years[0], years[-1] + 1)):
        return f"{years[0]}-{years[-1]}" if len(years) > 1 else str(years[0])
    return ", ".join(str(y) for y in years)


def _cma_listing(coverage: dict[str, tuple[str, list[int]]]) -> str:
    return "; ".join(
        f"{cma} ({province} {_years_span(years)})"
        for cma, (province, years) in sorted(coverage.items(), key=lambda kv: place_key(kv[0]))
    )


async def get_city_plan_prices(
    service: str = "mobile",
    *,
    city: str | list[str] | None = None,
    plan: str | None = None,
    technology: str = "all",
    from_year: int | None = None,
    to_year: int | None = None,
    include_rows: bool = False,
    limit: int = constants.ROWS_DEFAULT_LIMIT,
    lang: str = "en",
) -> CityPlanPrices:
    """Year-end mobile or Internet plan prices by census metropolitan area, averaged here."""
    lang = normalize_lang(lang)
    _check_limit(limit, lang)
    _check_years(from_year, to_year, lang)
    if service not in constants.PRICE_SHEETS:
        _raise(
            InvalidInput,
            "telecom_prices: service must be 'mobile' or 'internet'.",
            "telecom_prices : service doit être 'mobile' ou 'internet'.",
            lang,
        )
    tech = technology.strip().upper() if technology.strip().lower() != "all" else "all"
    if service == "mobile" and tech not in ("all", "3G", "LTE", "5G"):
        _raise(
            InvalidInput,
            "telecom_prices: technology must be 'all', '3G', 'LTE' or '5G'.",
            "telecom_prices : technology doit être 'all', '3G', 'LTE' ou '5G'.",
            lang,
        )
    sheets, _, resource, cached = await _workbook(
        constants.PRICES, lang, lambda n: n in (_REPORTED_SHEET, "Version"), "reported"
    )
    if _REPORTED_SHEET not in sheets:
        _raise(
            UpstreamError,
            f"telecom_prices: the price workbook has no {_REPORTED_SHEET!r} sheet.",
            f"telecom_prices : le classeur des prix n'a pas de feuille {_REPORTED_SHEET!r}.",
            lang,
        )
    try:
        reported = await run_parse(parsing.parse_reported_prices, sheets[_REPORTED_SHEET])
    except ValueError as exc:
        _raise(
            UpstreamError,
            f"telecom_prices: {exc}.",
            f"telecom_prices : la feuille Prix-Prices n'a pas la forme attendue ({exc}).",
            lang,
        )
    version, released = parsing.parse_version(sheets.get("Version", []))
    rows = [r for r in reported if r.service == service and r.region == "urban"]
    years_by_cma: dict[str, set[int]] = {}
    province_of: dict[str, str] = {}
    for r in rows:
        years_by_cma.setdefault(r.place, set()).add(r.year)
        province_of[r.place] = r.province
    coverage = {c: (province_of[c], sorted(y)) for c, y in years_by_cma.items()}
    names = sorted(coverage, key=place_key)

    queries = _split(city)
    if not queries or any(q.lower() in ("all", "tout", "toutes") for q in queries):
        wanted = names
    else:
        wanted = []
        rural = {place_key(r.place): r.place for r in reported if r.region == "rural"}
        for query in queries:
            found = match_cmas(query, names)
            if not found and place_key(query) in rural:
                _raise(
                    NotFound,
                    f"telecom_prices: {query!r} is one of the CRTC's rural communities; this "
                    "tool covers the urban census metropolitan areas. For rural prices use "
                    "telecom_get_plan_prices with region='rural'. CMAs: "
                    f"{_cma_listing(coverage)}.",
                    f"telecom_prices : {query!r} est une collectivité rurale du CRTC; cet outil "
                    "couvre les régions métropolitaines de recensement. Pour les prix ruraux, "
                    "utilisez telecom_get_plan_prices avec region='rural'. RMR : "
                    f"{_cma_listing(coverage)}.",
                    lang,
                )
            if not found:
                _raise(
                    NotFound,
                    f"telecom_prices: no census metropolitan area {query!r} in the CRTC price "
                    f"survey; CMAs and years: {_cma_listing(coverage)}.",
                    f"telecom_prices : aucune région métropolitaine de recensement {query!r} "
                    f"dans l'enquête du CRTC; RMR et années : {_cma_listing(coverage)}.",
                    lang,
                )
            wanted.extend(c for c in found if c not in wanted)

    plans_en = sorted({r.plan_en for r in rows}, key=_plan_order)
    wanted_plan = plan_key(plan) if plan else None
    if wanted_plan and wanted_plan not in {plan_key(p) for p in plans_en}:
        _raise(
            NotFound,
            f"telecom_prices: no {service} plan {plan!r}; plans are {plans_en}.",
            f"telecom_prices : aucun forfait {service} {plan!r} ; les forfaits sont {plans_en}.",
            lang,
        )
    in_scope = [
        r
        for r in rows
        if (wanted_plan is None or plan_key(r.plan_en) == wanted_plan)
        and (service != "mobile" or tech == "all" or r.technology == tech)
        and (from_year is None or r.year >= from_year)
        and (to_year is None or r.year <= to_year)
    ]
    wanted_set = set(wanted)
    matched = [r for r in in_scope if r.place in wanted_set]

    groups: dict[tuple[int, str, str], list[parsing.ReportedPrice]] = {}
    for r in matched:
        if r.value is not None:
            groups.setdefault((r.year, r.plan_en, r.place), []).append(r)
    averages: list[CityPlanPriceAverage] = []
    for (year, plan_en, place), group in groups.items():
        prices = [r.value for r in group if r.value is not None]
        lowest: dict[str, float] = {}
        for r in group:
            if r.provider and r.value is not None:
                lowest[r.provider] = min(lowest.get(r.provider, r.value), r.value)
        # 2016 rows name no provider: no per-provider figure rather than a wrong one.
        named = all(r.provider for r in group)
        averages.append(
            CityPlanPriceAverage(
                year=year,
                cma=place,
                province=province_of[place],
                plan=_plan_fr_label(plan_en) if lang == "fr" else plan_en,
                technology=(tech if service == "mobile" else None),
                average_price=round(sum(prices) / len(prices), 2),
                reported_prices=len(prices),
                providers=len(lowest) if named else None,
                lowest_per_provider_average=(
                    round(sum(lowest.values()) / len(lowest), 2) if named and lowest else None
                ),
                min_price=min(prices),
                max_price=max(prices),
            )
        )
    averages.sort(key=lambda a: (_plan_order(a.plan), place_key(a.cma), a.year))
    returned = averages[:limit]
    matched.sort(
        key=lambda r: (_plan_order(r.plan_en), place_key(r.place), r.year, r.provider or "")
    )
    provider_rows = (
        [
            CityPlanPriceRow(
                year=r.year,
                cma=r.place,
                province=r.province,
                plan=_plan_fr_label(r.plan_en) if lang == "fr" else r.plan_en,
                provider=r.provider,
                brand=r.brand,
                technology=r.technology,
                price=r.value,
                as_published=r.as_published,
            )
            for r in matched[:limit]
        ]
        if include_rows
        else []
    )

    notes = [_text(_CITY_BASIS, _CITY_BASIS_FR, lang)]
    scope_years = sorted({r.year for r in in_scope})
    for cma in wanted:
        missing = [y for y in scope_years if y not in years_by_cma[cma]]
        if not missing:
            continue
        province = province_of[cma]
        sampled = sorted(
            {r.place for r in in_scope if r.province == province and r.year in missing},
            key=place_key,
        )
        span = _years_span(missing)
        why_en = why_fr = ""
        if missing[0] >= 2024:
            why_en, why_fr = (
                "; from 2024 the CRTC surveys one CMA per province or territory",
                "; depuis 2024, le CRTC sonde une RMR par province ou territoire",
            )
        if sampled:
            why_en += f" (in {province}: {', '.join(sampled)})"
            why_fr += f" (en {province} : {', '.join(sampled)})"
        notes.append(
            _text(
                f"{cma}: no reported prices for {span}{why_en}. For {province} use the "
                f"CRTC's urban average: telecom_get_plan_prices(province='{province}', "
                "region='urban').",
                f"{cma} : aucun prix déclaré pour {span}{why_fr}. Pour {province}, utilisez "
                f"la moyenne urbaine du CRTC : telecom_get_plan_prices(province='{province}', "
                "region='urban').",
                lang,
            )
        )
    if any(a.year == 2016 for a in returned):
        notes.append(
            _text(
                "2016 rows name no provider, so providers and lowest_per_provider_average are "
                "empty for 2016.",
                "Les lignes de 2016 ne nomment pas le fournisseur : providers et "
                "lowest_per_provider_average sont vides pour 2016.",
                lang,
            )
        )
    unpriced = sum(1 for r in matched if r.value is None)
    if unpriced:
        notes.append(
            _text(
                f"{unpriced} matching rows have no price and are left out of the averages.",
                f"{unpriced} lignes retenues n'ont pas de prix et sont exclues des moyennes.",
                lang,
            )
        )

    years = sorted({r.year for r in rows})
    cut = len(averages) - len(returned)
    rows_cut = include_rows and len(matched) > len(provider_rows)
    return CityPlanPrices(
        service=service,  # type: ignore[arg-type]
        cmas_requested=wanted,
        dataset=_text(constants.PRICES.title_en, constants.PRICES.title_fr, lang),
        dataset_url=_dataset_url(constants.PRICES, lang),
        version=version,
        released=released,
        average_basis=_text(_CITY_BASIS, _CITY_BASIS_FR, lang),
        averages=returned,
        returned_count=len(returned),
        total_matched=len(averages),
        rows=provider_rows,
        rows_returned=len(provider_rows),
        rows_matched=len(matched),
        plans=[(_plan_fr_label(p) if lang == "fr" else p) for p in plans_en],
        years=years,
        cmas=[CmaCoverage(cma=c, province=coverage[c][0], years=coverage[c][1]) for c in names],
        notes=notes,
        method=_text(_CITY_METHOD, _CITY_METHOD_FR, lang),
        provenance=_provenance(
            str(resource.get("url") or ""),
            cached,
            "CityPlanPrices",
            lang,
            as_of=_as_of(released, resource),
            freshness=_text(
                "yearly: prices at December 31, published the following spring",
                "annuelle : prix au 31 décembre, publiés le printemps suivant",
                lang,
            ),
            coverage=_text(
                f"{len(rows):,} urban {service} prices in {len(names)} CMAs, {years[0]}-{years[-1]}"
                if years
                else "no rows",
                f"{len(rows):,} prix urbains ({service}) dans {len(names)} RMR, "
                f"{years[0]}-{years[-1]}"
                if years
                else "aucune ligne",
                lang,
            ),
            limits=join_limits(
                _text(
                    f"first {len(returned)} of {len(averages)} averages; narrow the filters "
                    "or raise limit",
                    f"{len(returned)} premières moyennes sur {len(averages)} ; précisez les "
                    "filtres ou augmentez limit",
                    lang,
                )
                if cut
                else None,
                _text(
                    f"first {len(provider_rows)} of {len(matched)} provider rows",
                    f"{len(provider_rows)} premières lignes de fournisseurs sur {len(matched)}",
                    lang,
                )
                if rows_cut
                else None,
                _text(
                    "City averages are computed by MapleStats from advertised list prices, "
                    "not published by the CRTC; no promotions, bundles or taxes.",
                    "Les moyennes par ville sont calculées par MapleStats à partir des prix "
                    "annoncés, et non publiées par le CRTC; sans promotions, offres groupées "
                    "ni taxes.",
                    lang,
                ),
            ),
        ),
    )


# Sector tables -----------------------------------------------------------------


def _sector(sector: str, lang: str) -> constants.Workbook:
    if sector not in constants.SECTORS:
        _raise(
            InvalidInput,
            "telecom_prices: sector must be 'mobile' or 'internet'.",
            "telecom_prices : sector doit être 'mobile' ou 'internet'.",
            lang,
        )
    return constants.SECTORS[sector]


_SECTOR_FRESHNESS = (
    "yearly: the Communications Market Report for year N comes out late in year N+1",
    (
        "annuelle : le Rapport sur le marché des communications de l'année N paraît à la "
        "fin de l'année N+1"
    ),
)


async def list_market_tables(sector: str = "mobile", *, lang: str = "en") -> MarketTableList:
    lang = normalize_lang(lang)
    workbook = _sector(sector, lang)
    sheets, package, resource, cached = await _workbook(
        workbook, lang, lambda n: n.lower().startswith(("table of contents", "table des")), "toc"
    )
    names = list(sheets[""][0])
    contents = next((rows for name, rows in sheets.items() if name), [])
    tables = parsing.parse_contents(contents, names)
    return MarketTableList(
        sector=sector,  # type: ignore[arg-type]
        dataset=_title(package, workbook, lang),
        dataset_url=_dataset_url(workbook, lang),
        tables=[
            MarketTableRef(table_id=t.table_id, label=t.label, kind=t.kind, title=t.title)  # type: ignore[arg-type]
            for t in tables
        ],
        provenance=_provenance(
            str(resource.get("url") or ""),
            cached,
            "MarketTableList",
            lang,
            as_of=_as_of(None, resource),
            freshness=_text(*_SECTOR_FRESHNESS, lang),
        ),
    )


async def get_market_table(
    table: str,
    sector: str = "mobile",
    *,
    series: str | None = None,
    from_year: int | None = None,
    to_year: int | None = None,
    lang: str = "en",
) -> MarketTable:
    lang = normalize_lang(lang)
    _check_years(from_year, to_year, lang)
    workbook = _sector(sector, lang)
    if not table.strip():
        _raise(
            InvalidInput,
            "telecom_prices: table must name a tab such as 'MB-S3' "
            "(see telecom_list_market_tables).",
            "telecom_prices : table doit nommer un onglet comme 'MB-S3' "
            "(voir telecom_list_market_tables).",
            lang,
        )
    aliases = parsing.sheet_aliases(table)
    sheets, package, resource, cached = await _workbook(
        workbook, lang, lambda n: n.upper() in aliases, "table:" + aliases[0]
    )
    found = next((name for name in sheets if name and name.upper() in aliases), None)
    if found is None:
        tabs = [n for n in sheets[""][0] if parsing.is_table_sheet(n)]
        _raise(
            NotFound,
            f"telecom_prices: no tab {table!r} in the {sector} workbook; tabs are {tabs}.",
            f"telecom_prices : aucun onglet {table!r} dans le classeur {sector} ; les onglets "
            f"sont {tabs}.",
            lang,
        )
    parsed = await run_parse(parsing.parse_table, sheets[found])
    if not parsed.observations:
        _raise(
            NotFound,
            f"telecom_prices: tab {found!r} has no table of years to read.",
            f"telecom_prices : l'onglet {found!r} ne contient aucun tableau par année.",
            lang,
        )
    words = parsing.text(series).casefold().split() if series else []
    kept = [
        o
        for o in parsed.observations
        if all(w in f"{o.block or ''} {o.series}".casefold() for w in words)
        and (from_year is None or o.year >= from_year)
        and (to_year is None or o.year <= to_year)
    ]
    mobile_only = any(o.mobile_phone_only for o in kept)
    return MarketTable(
        sector=sector,  # type: ignore[arg-type]
        table_id=found,
        label=parsed.label,
        title=parsed.title,
        dataset=_title(package, workbook, lang),
        dataset_url=_dataset_url(workbook, lang),
        observations=[MarketObservation(**o.__dict__) for o in kept],
        series=list(dict.fromkeys(o.series for o in parsed.observations)),
        notes=parsed.notes,
        derived_columns_left_out=parsed.derived_columns,
        provenance=_provenance(
            str(resource.get("url") or ""),
            cached,
            "MarketTable",
            lang,
            as_of=_as_of(None, resource),
            freshness=_text(*_SECTOR_FRESHNESS, lang),
            coverage=_text(
                f"{len(kept)} of {len(parsed.observations)} values in tab {found}",
                f"{len(kept)} valeurs sur {len(parsed.observations)} dans l'onglet {found}",
                lang,
            ),
            limits=join_limits(
                _text(
                    "Years marked mobile_phone_only (2020 on) count mobile phone services "
                    "only and do not compare with earlier years.",
                    "Les années marquées mobile_phone_only (2020 et après) ne comptent que "
                    "les services de téléphonie mobile et ne se comparent pas aux années "
                    "antérieures.",
                    lang,
                )
                if mobile_only
                else None,
                _text(
                    "Shares are stored as fractions in most tabs (0.283 is 28.3%).",
                    "Les parts sont le plus souvent des fractions (0,283 vaut 28,3 %).",
                    lang,
                )
                if any(o.value is not None and 0 < abs(o.value) < 1 for o in kept)
                else None,
            ),
        ),
    )


# Quarterly trends -----------------------------------------------------------------


_MARKETS = {"combined": "Combined", "residential": "Residential", "business": "Business"}


async def get_quarterly_trends(
    service: str = "mobile",
    *,
    metric: str | None = None,
    market: str = "all",
    measure: str | None = None,
    from_year: int | None = None,
    to_year: int | None = None,
    limit: int = constants.ROWS_DEFAULT_LIMIT,
    lang: str = "en",
) -> QuarterlyTrends:
    lang = normalize_lang(lang)
    _check_limit(limit, lang)
    _check_years(from_year, to_year, lang)
    if service not in constants.TREND_SHEETS:
        _raise(
            InvalidInput,
            "telecom_prices: service must be 'mobile', 'internet', 'wireline' or "
            "'tv_distribution'.",
            "telecom_prices : service doit être 'mobile', 'internet', 'wireline' ou "
            "'tv_distribution'.",
            lang,
        )
    if market not in (*_MARKETS, "all"):
        _raise(
            InvalidInput,
            "telecom_prices: market must be 'combined', 'residential', 'business' or 'all'.",
            "telecom_prices : market doit être 'combined', 'residential', 'business' ou 'all'.",
            lang,
        )
    sheet = constants.TREND_SHEETS[service]
    sheets, _, resource, cached = await _workbook(
        constants.TRENDS, lang, lambda n: n in (sheet, "Version"), "trends:" + service
    )
    if sheet not in sheets:
        _raise(
            UpstreamError,
            f"telecom_prices: the quarterly workbook has no {sheet!r} sheet.",
            f"telecom_prices : le classeur trimestriel n'a pas de feuille {sheet!r}.",
            lang,
        )
    try:
        rows = await run_parse(parsing.parse_trends, sheets[sheet])
    except ValueError as exc:
        _raise(
            UpstreamError,
            f"telecom_prices: {exc}.",
            f"telecom_prices : la feuille trimestrielle n'a pas la forme attendue ({exc}).",
            lang,
        )
    version, released = parsing.parse_version(sheets.get("Version", []))
    metric_words = parsing.text(metric).casefold().split() if metric else []
    measure_words = parsing.text(measure).casefold().split() if measure else []

    def has(words: list[str], *texts: str | None) -> bool:
        joined = " ".join(t or "" for t in texts).casefold()
        return all(w in joined for w in words)

    matched = [
        r
        for r in rows
        if has(metric_words, r.metric_en, r.metric_fr, r.subcategory_en, r.subcategory_fr)
        and has(measure_words, r.measure_en, r.measure_fr)
        and (market == "all" or r.market_en is None or r.market_en == _MARKETS[market])
        and (from_year is None or r.year >= from_year)
        and (to_year is None or r.year <= to_year)
    ]
    observations = [
        TrendObservation(
            period=f"{r.year}-Q{r.quarter}",
            year=r.year,
            quarter=r.quarter,
            market=r.market_fr if lang == "fr" else r.market_en,
            metric=r.metric_fr if lang == "fr" else r.metric_en,
            subcategory=r.subcategory_fr if lang == "fr" else r.subcategory_en,
            measure=r.measure_fr if lang == "fr" else r.measure_en,
            value=r.value,
        )
        for r in matched[:limit]
    ]
    metrics = sorted({(r.metric_fr if lang == "fr" else r.metric_en) for r in rows})
    quarters = sorted({(r.year, r.quarter) for r in rows})
    cut = len(matched) - len(observations)
    return QuarterlyTrends(
        service=service,  # type: ignore[arg-type]
        dataset=_text(constants.TRENDS.title_en, constants.TRENDS.title_fr, lang),
        dataset_url=_dataset_url(constants.TRENDS, lang),
        version=version,
        released=released,
        observations=observations,
        returned_count=len(observations),
        total_matched=len(matched),
        metrics=metrics,
        provenance=_provenance(
            str(resource.get("url") or ""),
            cached,
            "QuarterlyTrends",
            lang,
            as_of=_as_of(released, resource),
            freshness=_text(
                "quarterly, about five months after the quarter ends; past quarters are revised",
                "trimestrielle, environ cinq mois après la fin du trimestre; les trimestres "
                "passés sont révisés",
                lang,
            ),
            coverage=_text(
                f"{quarters[0][0]}-Q{quarters[0][1]} to {quarters[-1][0]}-Q{quarters[-1][1]}"
                if quarters
                else "no rows",
                f"{quarters[0][0]}-T{quarters[0][1]} à {quarters[-1][0]}-T{quarters[-1][1]}"
                if quarters
                else "aucune ligne",
                lang,
            ),
            limits=join_limits(
                _text(
                    f"first {len(observations)} of {len(matched)} matching values; narrow "
                    "metric, market or years, or raise limit",
                    f"{len(observations)} premières valeurs sur {len(matched)} ; précisez "
                    "metric, market ou les années, ou augmentez limit",
                    lang,
                )
                if cut
                else None,
                # Stated in the dataset's own description on open.canada.ca.
                _text(
                    "The quarterly survey (with Statistics Canada) covers the entities with "
                    "the top 90% of revenues, so totals run below the annual sector tables "
                    "(telecom_get_market_table).",
                    "L'enquête trimestrielle (avec Statistique Canada) couvre les entités qui "
                    "font 90 % des revenus : les totaux sont inférieurs à ceux des tableaux "
                    "annuels (telecom_get_market_table).",
                    lang,
                ),
            ),
        ),
    )
