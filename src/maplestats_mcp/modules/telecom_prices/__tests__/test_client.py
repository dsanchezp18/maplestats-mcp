"""Tests on rows trimmed from the CRTC Communications Market Reports workbooks (2026-10-27)."""

from __future__ import annotations

import io
import re
from datetime import datetime
from typing import Any

import pytest
from openpyxl import Workbook

from maplestats_mcp.modules.telecom_prices import client, constants, parsing
from maplestats_mcp.shared.errors import InvalidInput, NotFound, UpstreamError

_FILE = "https://applications.crtc.gc.ca/OpenData/CASP/COMMUNICATION MONITORING REPORTS/{}.xlsx"

_VERSION = [
    ("More information", "https://crtc.gc.ca/eng/publications/reports/policymonitoring/chg.htm"),
    (None, None, None),
    ("Version", "Date de publication / Release date", None),
    (11, datetime(2026, 5, 14), "Version actuelle / current version"),
    (10, datetime(2026, 2, 10), None),
]

# The real sheet mixes "Urbain / Urban" with "Urban" (2025), writes "--" for
# a price it did not compute, and lists rows newest first.
_MOBILE_PRICES = [
    (
        "Année / Year",
        "Région / Region",
        "Province ou territoire / Province or territory",
        "Forfait",
        "Plan",
        "Moyenne ($) / Average ($)",
        "Vitesse / Speed",
    ),
    (2025, "National", "Canada", "5 Go", "5GB", 30.08, "Plus bas prix / Lowest price"),
    (2025, "National", "Canada", "5 Go", "5GB", 28.95, "LTE"),
    (2025, "Urban", "Canada", "5 Go", "5GB", 29.5, "Plus bas prix / Lowest price"),
    (2025, "National", "Canada", "Plan voix", "Talk Plan", 20, "Plus bas prix / Lowest price"),
    (2021, "National", "NT", "5 Go", "5GB", "--", "5G"),
    (2020, "National", "Canada", "5 Go", "5GB", 46.90881854256854, "Plus bas prix / Lowest price"),
    (2020, "Urbain / Urban", "Canada", "5 Go", "5GB", 45.1, "Plus bas prix / Lowest price"),
    (
        2020,
        "National",
        "Territoires / Territories",
        "5 Go",
        "5GB",
        60,
        "Plus bas prix / Lowest price",
    ),
    (2020, "National", "Canada", "Forfait voix", "Talk Plan", 25, "Plus bas prix / Lowest price"),
    (2020, "National", "Canada", "10 Go", "10GB", 55, "Plus bas prix / Lowest price"),
]
_INTERNET_PRICES = [
    (
        "Année / Year",
        "Région / Region",
        "Province ou territoire / Province or territory",
        "Forfait",
        "Plan",
        "Moyenne / Average ($)",
    ),
    (2025, "National", "QC", "50/10 Mbps", "50/10 Mbps", 65.63),
    (2025, "National", "QC", "1,500/15 Mbps", "1,500/15 Mbps", 132.65),
]

# Trimmed from the real "Prix-Prices" sheet (2026-05-14): one row per provider,
# brand and network; "50GB Plan" in some years and "50GB" in others; blank
# entity in 2016; "Urban" alone in 2025; rural communities in the same
# column as the CMAs; TV rows; one row with no price.
_REPORTED = [
    (
        "Année / Year",
        "Entité / Entity",
        "Région / Region",
        "Région métropolitaine de recensement / Census metropolitan area",
        "Province ou territoire / Province or territory",
        "Forfait",
        "Plan",
        "Prix mensuel déclaré ($) / Reported monthly price ($)",
        "Marque / Brand",
        "Vitesse / Speed",
        "Chaînes / Channels",
    ),
    (2024, "Bell Mobility Inc.", "Urbain / Urban", "Toronto", "ON", "50 Go", "50GB", 34, "Virgin Mobile", "LTE", None),
    (2024, "Bell Mobility Inc.", "Urbain / Urban", "Toronto", "ON", "50 Go", "50GB", 34, "Lucky Mobile", "LTE", None),
    (2024, "Freedom Mobile Inc.", "Urbain / Urban", "Toronto", "ON", "50 Go", "50GB", 45, "Freedom Mobile", "5G", None),
    (2024, "TELUS Communications Inc.", "Urbain / Urban", "Toronto", "ON", "50 Go", "50GB", 39, "TELUS", "5G", None),
    (2024, "TELUS Communications Inc.", "Urbain / Urban", "Toronto", "ON", "50 Go", "50GB", 35, "Public Mobile", "LTE", None),
    (2024, "Bell Mobility Inc.", "Urbain / Urban", "Edmonton", "AB", "50 Go", "50GB", 34, "Virgin Mobile", "LTE", None),
    (2024, "TELUS Communications Inc.", "Urbain / Urban", "Edmonton", "AB", "50 Go", "50GB", 39, "Koodo", "LTE", None),
    (2025, "Bell Mobility Inc.", "Urban", "Edmonton", "AB", "50 Go", "50GB Plan", 39, None, "LTE", None),
    (2025, "Bell Mobility Inc.", "Urban", "Edmonton", "AB", "50 Go", "50GB Plan", None, None, "5G", None),
    (2023, "Rogers Communications Canada Inc.", "Urbain / Urban", "Calgary", "AB", "50 Go", "50GB Plan", 65, "Rogers", "5G", None),
    (2023, "Rogers Communications Canada Inc.", "Urbain / Urban", "Calgary", "AB", "50 Go", "50GB Plan", 34, "Fido", "LTE", None),
    (2023, "Bell Mobility Inc.", "Urbain / Urban", "Montréal", "QC", "5 Go", "5GB Plan", 40, "Bell", "LTE", None),
    (2016, None, "Urbain / Urban", "Ottawa-Gatineau", "ON", "50/10 Mbps", "50/10 Mbps", 70, None, None, None),
    (2016, None, "Urbain / Urban", "Ottawa-Gatineau", "ON", "50/10 Mbps", "50/10 Mbps", 75.5, None, None, None),
    (2022, "Rogers Communications Canada Inc.", "Urbain / Urban", "Ottawa-Gatineau", "ON", "50/10 Mbps", "50/10 Mbps", 80, "Rogers", None, None),
    (2025, "4141903 Canada Inc. (Storm Internet Services)", "Rural", "Ingleside", "ON", "50/10 Mbps", "50/10 Mbps", 68.95, "Storm Internet Services", None, None),
    (2020, "Bell Canada", "Urbain / Urban", "Calgary", "AB", "SRD", "DTH", None, None, None, 27),
]  # fmt: skip

_CONTENTS = [
    ("Table of Contents by infographics, figures, tables and supplementary tables",),
    ("Figures", "Tab Name", "Title"),
    ("Figure 12", "MB-F12", "National mobile ARPU ($/month), 2013-2024"),
    ("Supplementary Table 1", "MB-S1", "Retail mobile revenue and subscriber market share"),
    ("Supplementary Table 3", "MB-S3", "Average revenue per 1 GB of data ($), 2015-2024"),
    ("Table of Contents by Metric",),
    ("Supplementary Table 3", "MB-S3", "Average revenue per 1 GB of data ($), 2015-2024"),
]
_S3 = [
    ("Supplementary Table 3",),
    ("Average revenue per 1 GB of data ($), 2015-2024",),
    (None,),
    ("Type", 2015, 2019, "2020 (MP)", "2024 (MP)", "Growth (%) 2023-2024", "CAGR (%) 2020-2024"),
    ("Top 3", 30.79, 10.78, 9.46, 4.5, -20.7, -17),
    ("All providers", 27.96, 9.93, 8.73, 4.49, -20.6, -15.3),
    (None,),
    ("Source: CRTC data collection",),
    ('Note: "Top 3" includes Bell, Rogers, TELUS and their flankers.',),
]
# Two blocks in one tab, each with its own title and header row.
_S1 = [
    ("Supplementary Table 1",),
    ("Retail mobile revenue and subscriber market share by service provider, 2013-2024",),
    ("Retail mobile revenue market share by service provider",),
    ("Service provider", 2019, "2020 (MP)"),
    ("Bell", 0.295, 0.292),
    ("Other providers ", 0.099, 0.114),
    ("Retail mobile subscriber market share by service provider",),
    ("Service provider", 2019, "2020 (MP)"),
    ("Bell", 0.29, 0.285),
    ("Source: CRTC data collection",),
]
_F12 = [
    ("Figure 12",),
    ("National mobile ARPU ($/month), 2013-2024",),
    ("Type", 2013, "2024 (MP)"),
    ("ARPU", "$59.97", "$69.16"),
    ("Source: CRTC data collection",),
]
_G12 = [
    ("Graphique 12",),
    ("RMPU pour les services mobiles à l’échelle nationale ($/mois), 2013-2024",),
    ("Type", 2013, "2024 (TM)", "Croissance (%) 2023-2024"),
    ("RMPU", 59.97, 69.16, -1.5),
    ("Source : Données recueillies par le CRTC",),
]
_INFOGRAPHIC = [("Infographic 1",), ("Highlights",), ("Retail mobile revenues grew",)]

_TRENDS = [
    (
        "Date",
        "Marché d'affaires",
        "Business Market",
        "Mesure",
        "Metric",
        "Revenus / Revenues",
        "Abonnés / Subscribers",
        "Souscatégorie",
        "Subcategory",
        "Données sans fils (Go) / Wireless Data (GB)",
    ),
    ("2026 Q1", "Combinés", "Combined", "Trafic de données sans fil (Go)",
     "Wireless data traffic (GB)", None, None, None, None, 1042163326),
    ("2025 Q4", "Combinés", "Combined", "Trafic de données sans fil (Go)",
     "Wireless data traffic (GB)", None, None, None, None, 1147295899),
    ("2025 Q4", "Résidentiel", "Residential", "Téléphone mobile", "Mobile phone",
     None, 30000000, "Voix et données", "Voice and data", None),
    ("2025 Q4", "Combinés", "Combined", "Services mobile", "Mobile services",
     4123945600, None, "Données, incluant itinérance", "Data, including roaming", None),
]  # fmt: skip


def _xlsx(sheets: dict[str, list[tuple[Any, ...]]]) -> bytes:
    book = Workbook()
    book.remove(book.active)  # type: ignore[arg-type]
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(list(row))
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _package(workbook: constants.Workbook, files: dict[str, str]) -> dict[str, Any]:
    resources = [
        {"id": rid, "format": "XLSX", "language": [lang], "url": _FILE.format(name)}
        for rid, (lang, name) in {
            workbook.resource_en: ("en", files["en"]),
            workbook.resource_fr: ("fr", files.get("fr", files["en"])),
        }.items()
    ]
    return {
        "success": True,
        "result": {
            "id": workbook.dataset_id,
            "title_translated": {"en": workbook.title_en, "fr": workbook.title_fr},
            "resources": resources,
        },
    }


def _mock(httpx_mock, workbook: constants.Workbook, files: dict[str, str], bodies: dict) -> None:
    httpx_mock.add_response(
        url=re.compile(re.escape(constants.CKAN_PACKAGE_SHOW) + r"\?id=" + workbook.dataset_id),
        json=_package(workbook, files),
        is_reusable=True,
    )
    for name, body in bodies.items():
        # A test reads one language's workbook, so the other is optional.
        httpx_mock.add_response(
            url=_FILE.format(name).replace(" ", "%20"),
            content=body,
            is_reusable=True,
            is_optional=True,
        )


def _prices(httpx_mock) -> None:
    body = _xlsx(
        {
            "Version": _VERSION,
            "Prix-Prices": _REPORTED,
            "Moyennes-Averages Internet": _INTERNET_PRICES,
            "Moyennes-Averages Mobile": _MOBILE_PRICES,
        }
    )
    _mock(httpx_mock, constants.PRICES, {"en": "prices"}, {"prices": body})


# Parsers ----------------------------------------------------------------------


def test_number_reads_dollar_percent_and_symbols():
    assert parsing.number("$59.97") == (59.97, "$59.97")
    assert parsing.number(" 28.4 ") == (28.4, "28.4")
    assert parsing.number("30.6%") == (30.6, "30.6%")
    assert parsing.number("1,042") == (1042.0, "1,042")
    assert parsing.number("--") == (None, "--")
    assert parsing.number("More than one quarter") == (None, "More than one quarter")
    assert parsing.number(0.283) == (0.283, None)


def test_prices_unify_labels():
    rows = parsing.parse_prices(_MOBILE_PRICES, "mobile")
    assert {r.region for r in rows} == {"national", "urban"}
    assert {r.geography for r in rows} == {"Canada", "NT", "Territories"}
    talk = [r for r in rows if r.plan_en == "Talk Plan"]
    assert {r.plan_fr for r in talk} == {"Forfait voix"}
    assert rows[0].plan_en == "5 GB" and rows[0].technology == "lowest"
    missing = next(r for r in rows if r.geography == "NT")
    assert missing.value is None and missing.as_published == "--"


def test_reported_prices_sheet():
    rows = parsing.parse_reported_prices(_REPORTED)
    assert {r.service for r in rows} == {"mobile", "internet", "tv"}
    assert {r.plan_en for r in rows if r.service == "mobile"} == {"50 GB", "5 GB"}
    assert {r.region for r in rows} == {"urban", "rural"}
    old = next(r for r in rows if r.year == 2016)
    assert old.provider is None and old.technology is None and old.value == 70
    assert next(r for r in rows if r.value is None and r.service == "mobile").technology == "5G"
    with pytest.raises(ValueError):
        parsing.parse_reported_prices([("Année / Year", "Plan")])


def test_cma_names_match_without_accents_case_or_hyphen_sides():
    cmas = ["Montréal", "Ottawa-Gatineau", "Québec", "St. John's", "St Catharines - Niagara"]
    assert client.match_cmas("MONTREAL", cmas) == ["Montréal"]
    assert client.match_cmas("quebec city", cmas) == ["Québec"]
    assert client.match_cmas("Ottawa", cmas) == ["Ottawa-Gatineau"]
    assert client.match_cmas("Gatineau", cmas) == ["Ottawa-Gatineau"]
    assert client.match_cmas("Saint Johns", cmas) == ["St. John's"]
    assert client.match_cmas("niagara", cmas) == ["St Catharines - Niagara"]
    assert client.match_cmas("Toronto", cmas) == []


def test_version_sheet():
    assert parsing.parse_version(_VERSION)[0] == "11"
    assert parsing.parse_version(_VERSION)[1].isoformat() == "2026-05-14"  # type: ignore[union-attr]
    csv_like = [("Version", "Release date", None), ("17.0", "8/21/2026", "current version")]
    assert parsing.parse_version(csv_like)[1].isoformat() == "2026-08-21"  # type: ignore[union-attr]


def test_table_drops_derived_columns_and_marks_mobile_phone_only():
    table = parsing.parse_table(_S3)
    assert table.label == "Supplementary Table 3"
    assert table.derived_columns == ["Growth (%) 2023-2024", "CAGR (%) 2020-2024"]
    top = [
        (o.year, o.value, o.mobile_phone_only) for o in table.observations if o.series == "Top 3"
    ]
    assert top == [
        (2015, 30.79, False),
        (2019, 10.78, False),
        (2020, 9.46, True),
        (2024, 4.5, True),
    ]
    assert table.notes[0] == "Source: CRTC data collection" and len(table.notes) == 2


def test_table_blocks_and_text_values():
    table = parsing.parse_table(_S1)
    blocks = {o.block for o in table.observations}
    assert blocks == {
        "Retail mobile revenue market share by service provider",
        "Retail mobile subscriber market share by service provider",
    }
    assert len(table.observations) == 6
    arpu = parsing.parse_table(_F12)
    assert [(o.year, o.value, o.as_published) for o in arpu.observations] == [
        (2013, 59.97, "$59.97"),
        (2024, 69.16, "$69.16"),
    ]


def test_qualifier_row_under_the_years():
    rows = [
        ("Supplementary table 3",),
        ("Internet access revenue overages summary, 2016-2024",),
        (None, 2016, 2016),
        (None, "Residential", "Business"),
        ("Overage fees ($M)", 100.4, 2.5),
        ("Source: CRTC data collection",),
    ]
    table = parsing.parse_table(rows)
    assert [(o.qualifier, o.value) for o in table.observations] == [
        ("Residential", 100.4),
        ("Business", 2.5),
    ]


def test_contents_and_aliases():
    refs = parsing.parse_contents(_CONTENTS, ["MB-F12", "MB-S1", "MB-S3"])
    assert [r.table_id for r in refs] == ["MB-F12", "MB-S1", "MB-S3"]
    assert refs[0].kind == "figure" and refs[2].kind == "supplementary"
    assert parsing.sheet_aliases("mb-f12") == ["MB-F12", "MB-G12"]
    assert parsing.sheet_aliases("N-F2 & F13") == ["N-F2 & F13", "N-G2 & G13"]


def test_trends_sort_and_split_measures():
    rows = parsing.parse_trends(_TRENDS)
    traffic = [(r.year, r.quarter, r.value) for r in rows if r.measure_en == "Wireless Data (GB)"]
    assert traffic == [(2025, 4, 1147295899.0), (2026, 1, 1042163326.0)]
    subscribers = next(r for r in rows if r.measure_en == "Subscribers")
    assert (
        subscribers.subcategory_fr == "Voix et données" and subscribers.market_en == "Residential"
    )


def test_plan_key():
    assert client.plan_key("5GB") == client.plan_key("5 Go") == client.plan_key("5") == "5"
    assert client.plan_key("Talk Plan") == client.plan_key("forfait voix") == "talk"
    assert client.plan_key("1,500/15 Mbps") == "1500/15"


# Client -----------------------------------------------------------------------


async def test_plan_prices_default_is_lowest_national_canada(httpx_mock):
    _prices(httpx_mock)
    result = await client.get_plan_prices("mobile", plan="5 GB")
    assert [(p.year, p.average_price) for p in result.points] == [
        (2020, 46.90881854256854),
        (2025, 30.08),
    ]
    assert result.version == "11" and result.released.isoformat() == "2026-05-14"  # type: ignore[union-attr]
    assert result.plans == ["Talk Plan", "5 GB", "10 GB"]
    assert result.provenance.source == "crtc-cmr"
    assert "Open Government Licence" in (result.provenance.licence or "")
    assert result.provenance.as_of is not None and result.provenance.as_of.year == 2026


async def test_plan_prices_french_regions_and_technologies(httpx_mock):
    _prices(httpx_mock)
    result = await client.get_plan_prices(
        "mobile", plan="voix", region="all", technology="all", lang="fr"
    )
    assert {p.plan for p in result.points} == {"Forfait voix"}
    assert "Forfait voix" in result.plans
    urban = await client.get_plan_prices("mobile", plan="5", region="urban")
    assert [(p.year, p.average_price) for p in urban.points] == [(2020, 45.1), (2025, 29.5)]
    lte = await client.get_plan_prices("mobile", plan="5", technology="LTE")
    assert [p.average_price for p in lte.points] == [28.95]


async def test_internet_prices_by_province(httpx_mock):
    _prices(httpx_mock)
    result = await client.get_plan_prices("internet", province="Quebec")
    assert [p.plan for p in result.points] == ["50/10 Mbps", "1,500/15 Mbps"]
    assert all(p.technology is None for p in result.points)


async def test_plan_prices_errors(httpx_mock):
    with pytest.raises(InvalidInput):
        await client.get_plan_prices(province="Atlantis")
    with pytest.raises(InvalidInput):
        await client.get_plan_prices(from_year=2025, to_year=2020)
    with pytest.raises(InvalidInput):
        await client.get_plan_prices(limit=0)
    _prices(httpx_mock)
    with pytest.raises(NotFound, match="plans are"):
        await client.get_plan_prices(plan="7 GB")


async def test_limit_is_reported(httpx_mock):
    _prices(httpx_mock)
    result = await client.get_plan_prices(region="all", technology="all", province="all", limit=2)
    assert result.returned_count == 2 and result.total_matched > 2
    assert "first 2 of" in (result.provenance.limits or "")


async def test_plan_prices_for_several_provinces(httpx_mock):
    _prices(httpx_mock)
    listed = await client.get_plan_prices("mobile", plan="5", province=["Canada", "Territories"])
    assert {(p.geography, p.year) for p in listed.points} == {
        ("Canada", 2020),
        ("Canada", 2025),
        ("Territories", 2020),
    }
    text = await client.get_plan_prices("mobile", plan="5", province="canada, NT", technology="all")
    assert {p.geography for p in text.points} == {"Canada", "NT"}
    every = await client.get_plan_prices("mobile", plan="5", province=["NT", "all"])
    assert {p.geography for p in every.points} == {"Canada", "Territories"}
    with pytest.raises(InvalidInput, match="Atlantis"):
        await client.get_plan_prices(province="ON, Atlantis")
    with pytest.raises(InvalidInput):
        await client.get_plan_prices(province=" , ")


async def test_city_prices_are_computed_means_with_counts(httpx_mock):
    _prices(httpx_mock)
    result = await client.get_city_plan_prices("mobile", city="toronto, EDMONTON", plan="50 GB")
    assert result.cmas_requested == ["Toronto", "Edmonton"]
    by = {(a.cma, a.year): a for a in result.averages}
    toronto = by[("Toronto", 2024)]
    # Mean of the five rows; the per-provider mean keeps Bell 34, Freedom 45, TELUS 35.
    assert toronto.average_price == 37.4 and toronto.reported_prices == 5
    assert toronto.providers == 3 and toronto.lowest_per_provider_average == 38.0
    assert (toronto.min_price, toronto.max_price, toronto.technology) == (34, 45, "all")
    # 2025 has one priced row and one without a price.
    assert by[("Edmonton", 2025)].reported_prices == 1
    assert any("1 matching rows have no price" in n for n in result.notes)
    assert result.notes[0].startswith("Computed by MapleStats: mean of the reported monthly")
    assert "Computed by MapleStats" in result.average_basis
    assert [c.cma for c in result.cmas] == ["Calgary", "Edmonton", "Montréal", "Toronto"]
    assert result.provenance.schema_name == "telecom_prices.CityPlanPrices"
    assert "not published by the CRTC" in (result.provenance.limits or "")
    lte = await client.get_city_plan_prices("mobile", city=["Toronto"], technology="LTE")
    assert [(a.average_price, a.technology) for a in lte.averages] == [(34.33, "LTE")]


async def test_city_without_rows_for_a_year_points_to_the_provincial_average(httpx_mock):
    _prices(httpx_mock)
    result = await client.get_city_plan_prices("mobile", city="Calgary", plan="50GB")
    assert [a.year for a in result.averages] == [2023]
    missing = [n for n in result.notes if n.startswith("Calgary:")]
    assert len(missing) == 1
    assert "2024-2025" in missing[0] and "(in AB: Edmonton)" in missing[0]
    assert "telecom_get_plan_prices(province='AB', region='urban')" in missing[0]
    french = await client.get_city_plan_prices("mobile", city="Calgary", lang="fr")
    assert any(n.startswith("Calgary\u00a0: aucun prix déclaré") for n in french.notes)
    assert french.notes[0].startswith("Calculé par MapleStats")


async def test_city_internet_2016_has_no_provider_counts(httpx_mock):
    _prices(httpx_mock)
    result = await client.get_city_plan_prices(
        "internet", city="Ottawa", include_rows=True, lang="fr"
    )
    first = result.averages[0]
    assert (first.year, first.average_price, first.providers) == (2016, 72.75, None)
    assert first.lowest_per_provider_average is None and first.technology is None
    assert result.averages[1].providers == 1
    assert result.rows_returned == 3 and result.rows[0].provider is None
    assert any("2016" in n for n in result.notes)


async def test_city_rows_and_averages_are_limited(httpx_mock):
    _prices(httpx_mock)
    result = await client.get_city_plan_prices("mobile", include_rows=True, limit=2)
    assert result.returned_count == 2 and result.total_matched > 2
    assert result.rows_returned == 2 and result.rows_matched > 2
    assert "first 2 of" in (result.provenance.limits or "")
    plain = await client.get_city_plan_prices("mobile", city="all")
    assert plain.rows == [] and plain.rows_matched > 0


async def test_city_errors(httpx_mock):
    with pytest.raises(InvalidInput):
        await client.get_city_plan_prices(technology="4G")
    with pytest.raises(InvalidInput):
        await client.get_city_plan_prices(limit=0)
    _prices(httpx_mock)
    with pytest.raises(NotFound, match=r"Calgary \(AB 2023\)"):
        await client.get_city_plan_prices(city="Atlantis")
    with pytest.raises(NotFound, match="rural communities"):
        await client.get_city_plan_prices("internet", city="Ingleside")
    with pytest.raises(NotFound, match="plans are"):
        await client.get_city_plan_prices(city="Toronto", plan="7 GB")


def _sector(httpx_mock) -> None:
    english = _xlsx(
        {
            "Table of Contents by I,F,T,S": _CONTENTS,
            "MB-I1": _INFOGRAPHIC,
            "MB-F12": _F12,
            "MB-S1": _S1,
            "MB-S3": _S3,
        }
    )
    french = _xlsx({"Table des matières - I,G,T,S": _CONTENTS, "MB-G12": _G12})
    _mock(
        httpx_mock,
        constants.SECTORS["mobile"],
        {"en": "mobile-en", "fr": "mobile-fr"},
        {"mobile-en": english, "mobile-fr": french},
    )


async def test_market_tables(httpx_mock):
    _sector(httpx_mock)
    listing = await client.list_market_tables("mobile")
    assert [t.table_id for t in listing.tables] == ["MB-F12", "MB-S1", "MB-S3"]
    table = await client.get_market_table("mb-s3", series="all providers", from_year=2020)
    assert [(o.year, o.value) for o in table.observations] == [(2020, 8.73), (2024, 4.49)]
    assert table.series == ["Top 3", "All providers"]
    assert "mobile_phone_only" in (table.provenance.limits or "")
    shares = await client.get_market_table("MB-S1", series="bell")
    assert len(shares.observations) == 4
    assert "fractions" in (shares.provenance.limits or "")


async def test_french_workbook_uses_g_tabs(httpx_mock):
    _sector(httpx_mock)
    table = await client.get_market_table("MB-F12", lang="fr")
    assert table.table_id == "MB-G12"
    assert table.derived_columns_left_out == ["Croissance (%) 2023-2024"]
    assert [o.value for o in table.observations] == [59.97, 69.16]


async def test_market_table_errors(httpx_mock):
    _sector(httpx_mock)
    with pytest.raises(NotFound, match="tabs are"):
        await client.get_market_table("MB-S99")
    with pytest.raises(NotFound, match="no table of years"):
        await client.get_market_table("MB-I1")
    with pytest.raises(InvalidInput):
        await client.get_market_table(" ")
    with pytest.raises(InvalidInput):
        await client.get_market_table("MB-S3", sector="tv")


async def test_quarterly_trends(httpx_mock):
    body = _xlsx({"Version": _VERSION, "Mobile": _TRENDS})
    _mock(httpx_mock, constants.TRENDS, {"en": "trends"}, {"trends": body})
    result = await client.get_quarterly_trends("mobile", metric="wireless data traffic")
    assert [(o.period, o.value) for o in result.observations] == [
        ("2025-Q4", 1147295899.0),
        ("2026-Q1", 1042163326.0),
    ]
    assert result.provenance.coverage == "2025-Q4 to 2026-Q1"
    residential = await client.get_quarterly_trends(
        "mobile", market="residential", measure="abonnés", lang="fr"
    )
    assert [(o.metric, o.subcategory, o.value) for o in residential.observations] == [
        ("Téléphone mobile", "Voix et données", 30000000.0)
    ]
    with pytest.raises(UpstreamError, match="no 'Internet' sheet"):
        await client.get_quarterly_trends("internet")


async def test_resource_found_again_by_format_when_its_id_changes(httpx_mock):
    package = _package(constants.PRICES, {"en": "prices"})
    for resource in package["result"]["resources"]:
        resource["id"] = "replaced"
    picked = client.pick_resource(package["result"], constants.PRICES.resource_en, "en")
    assert picked["url"].endswith("prices.xlsx")
    with pytest.raises(NotFound):
        client.pick_resource({"resources": [{"id": "x", "format": "PDF"}]}, "y", "en")


async def test_a_web_page_instead_of_the_workbook_is_an_upstream_error(httpx_mock):
    _mock(httpx_mock, constants.PRICES, {"en": "prices"}, {"prices": b"<html>maintenance</html>"})
    with pytest.raises(UpstreamError, match="not an Excel workbook"):
        await client.get_plan_prices()


async def test_dataset_gone_is_not_found(httpx_mock):
    httpx_mock.add_response(
        url=re.compile(re.escape(constants.CKAN_PACKAGE_SHOW) + r".*"), status_code=404
    )
    with pytest.raises(NotFound):
        await client.get_plan_prices()
