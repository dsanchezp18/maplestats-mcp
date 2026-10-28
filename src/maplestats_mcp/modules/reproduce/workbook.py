"""reproduce_workbook: a MapleStats result as a formatted Excel workbook.

The rows are the ones the tool returned (the tool runs once here, through
the server, so its cache and rate limits apply), or rows the caller passes
back. They are cleaned as the R scripts clean data (snake_case names as
janitor gives them, trimmed text, empty text as missing, numbers and ISO
dates stored as text converted, code columns kept as text as
cleaning.CODE_COLUMNS keeps them; the response's notes name the columns
converted and the result's fields left out of the rows) and written following the house Excel
conventions (R Code Conventions section 11, applied with openpyxl since
this server is Python): a data sheet named after its content holding one
Excel table object with a frozen header, explicit number formats and
fitted column widths; a Source sheet with the call, provenance, licence,
what cleaning changed and the attribution in the form the terms ask for
(Statistics Canada's "Adapted from" notice once cleaning changed anything,
its "Source:" notice otherwise; a note of the changes for the Bank of Canada
and Creative Commons Attribution); and, where the rows are a time series
or a short list of values, a native (editable) Excel line or bar chart on
its own sheet, drawn from a Chart data sheet.

Delivery. A hosted server cannot write to the user's disk, so the
workbook comes back as base64 with its file name. That text is a third
larger than the file and MCP clients pass it through the model's context
(FastMCP sends the result as structured content and as JSON text), so rows
are capped (max_rows, at most MAX_ROWS) and so is the file (MAX_BYTES).
On a local stdio server the workbook is written to MAPLE_EXPORT_DIR
(default: the user's Downloads folder) and only its path comes back.
For data past these caps, reproduce_code with language="excel" gives a
Power Query that loads the whole table inside Excel and refreshes it.
"""

from __future__ import annotations

import base64
import io
import json
import math
import os
import re
import unicodedata
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any, Literal

from maplestats_mcp.modules.reproduce.cleaning import CODE_COLUMNS
from maplestats_mcp.modules.reproduce.schemas import ExcelWorkbook
from maplestats_mcp.shared.envelope import STATCAN_LICENCE, STATCAN_LICENCE_FR, make_provenance
from maplestats_mcp.shared.errors import InvalidInput, NotFound
from maplestats_mcp.shared.executor import check_deadline, run_parse
from maplestats_mcp.shared.i18n import NBSP, french_spacing, french_text, normalize_lang
from maplestats_mcp.shared.licences import BOC_TERMS, LICENCES_FR

MAX_ROWS = 20_000
DEFAULT_ROWS = 2_000
MAX_COLUMNS = 256
MAX_CELLS = 200_000
MAX_TEXT_CHARS = 8_000_000
# The workbook file; base64 adds a third, and the client sees it twice.
MAX_BYTES = 2_000_000
MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# Excel stores at most 32,767 characters in a cell.
_MAX_CELL = 32_767
_NOT_DATA = ("reproduce_code", "reproduce_workbook", "plan_query", "search_tools", "call_tool")
# One palette for every chart (R Code Conventions section 10: declared once).
PALETTE = ["0D3692", "E60F2D", "2A9D8F", "E9C46A", "6D597A", "F4A261", "264653", "8AB17D"]

Delivery = Literal["auto", "base64", "file"]

LABELS: dict[str, dict[str, str]] = {
    "en": {
        "data": "Data",
        "source": "Source",
        "chart_data": "Chart data",
        "chart": "Chart",
        "field": "Field",
        "value": "Value",
        "title": "Title",
        "produced_by": "Produced by",
        "tool": "Tool",
        "arguments": "Arguments (JSON)",
        "source_name": "Source",
        "source_url": "Source URL",
        "queried_at": "Queried at (UTC)",
        "as_of": "Data as of",
        "coverage": "Coverage",
        "limits": "Limits",
        "licence": "Licence",
        "attribution": "Attribution",
        "rows": "Rows",
        "cleaning": "Cleaning",
        "reproduce": "Reproduce",
        "created": "Workbook created (UTC)",
        "rows_text": "{written} of {available} rows",
        "rows_capped": " (capped by max_rows)",
        "caller_rows": "rows passed to reproduce_workbook by the caller",
        "no_licence": (
            "Not stated in the tool's provenance: check the publisher's terms at the source URL."
        ),
        "ogl_attribution": (
            "Contains information licensed under the Open Government Licence - {region}."
        ),
        "cite": "Cite the publisher and the source URL above.",
        "statcan_source": (
            "Source: Statistics Canada, {product}, {date}. Reproduced and distributed on an "
            '"as is" basis with the permission of Statistics Canada.'
        ),
        "statcan_adapted": (
            "Adapted from Statistics Canada, {product}, {date}. This does not constitute an "
            "endorsement by Statistics Canada of this product."
        ),
        "table": "Table",
        "changes_flag": "Changes were made to the data: see Changes made.",
        "changes": "Changes made",
        "changes_list": "{changes}.",
        "no_changes": "None: column names and values as the tool returned them.",
        "renamed": "column names put in snake_case ({count} of {total}: {examples})",
        "trimmed": "cells trimmed of spaces: {count}",
        "emptied": "cells left blank (empty text or a non-finite number): {count}",
        "converted": "text converted to numbers or dates in {columns}",
        "long_ints": "whole numbers over 15 digits stored as text in {columns}",
        "cleaning_text": (
            "Names in snake_case (as janitor::clean_names in R), text trimmed, empty text as "
            "missing, numbers and ISO dates stored as text converted; codes with leading "
            "zeros, and columns named like a code (coordinate, vector or product id, DGUID, "
            "postal code, NAICS, NOC, SGC), kept as text."
        ),
        "cleaned_note": (
            "Cleaned as the R scripts clean data: snake_case column names, text trimmed, empty "
            "text as missing."
        ),
        "converted_note": "Converted from text to numbers or dates: {columns}.",
        "codes_note": "Kept as text because they hold codes: {columns}.",
        "left_out": "Fields of the result",
        "left_out_note": (
            "Not in the data rows (other parts of the tool's result): {fields}. Their values "
            "are on the Source sheet; lists are named only."
        ),
        "reproduce_text": (
            "Call reproduce_code with the same tool and arguments for R, Python, Stata and "
            "Julia scripts, or language='excel' for a Power Query that refreshes the full data."
        ),
        "default_title": "MapleStats data",
        "chart_text": "{kind} chart of {columns} by {x}",
        "line": "line",
        "bar": "bar",
        "max_rows": "max_rows must be between 1 and {limit}, got {got}.",
        "no_call": "Pass tool_name (and its arguments), or the rows to write.",
        "not_data": "{tool} does not return data rows to write to a workbook.",
        "too_wide": (
            "The workbook exceeds {columns} columns, {cells} cells or 8 million text "
            "characters. Select fewer columns or rows."
        ),
        "delivery": "delivery must be 'auto', 'base64' or 'file'.",
        "hosted": "This is a hosted server: it cannot write to your disk. Use delivery='base64'.",
        "rows_shape": "rows must be a list of objects, one per row.",
        "no_rows": "{what} returned no rows to write to a workbook.",
        "the_rows": "The rows",
        "too_big": (
            "The workbook is {size} bytes, over the {limit}-byte limit for a tool response. "
            "Lower max_rows, or call reproduce_code with language='excel' for a Power Query "
            "that loads the full data inside Excel."
        ),
        "failed": "{tool} failed with these arguments: {text}",
        "unstructured": "{tool} did not return a structured result.",
        "first_rows": "Wrote the first {written} of {available} rows (max_rows).",
        "saved": "Saved to {path}.",
        "decode": (
            "Decode workbook_base64 and save it as {file} (a hosted server cannot write to your "
            "disk). Large results pass through the client's context: keep max_rows small, or "
            "use reproduce_code language='excel' for the full data."
        ),
        "limits_text": "at most {rows} rows (max_rows, default {default}) and {size} bytes per workbook",
    },
    "fr": {
        "data": "Données",
        "source": "Source",
        "chart_data": "Données du graphique",
        "chart": "Graphique",
        "field": "Champ",
        "value": "Valeur",
        "title": "Titre",
        "produced_by": "Produit par",
        "tool": "Outil",
        "arguments": "Arguments (JSON)",
        "source_name": "Source",
        "source_url": "URL de la source",
        "queried_at": "Interrogé le (UTC)",
        "as_of": "Données en date du",
        "coverage": "Couverture",
        "limits": "Limites",
        "licence": "Licence",
        "attribution": "Mention de la source",
        "rows": "Lignes",
        "cleaning": "Nettoyage",
        "reproduce": "Reproduire",
        "created": "Classeur créé le (UTC)",
        "rows_text": "{written} lignes sur {available}",
        "rows_capped": " (limite max_rows)",
        "caller_rows": "lignes transmises à reproduce_workbook par l'appelant",
        "no_licence": (
            "Non précisée dans la provenance de l'outil : consultez les conditions de "
            "l'éditeur à l'URL de la source."
        ),
        "ogl_attribution": (
            "Contient des renseignements visés par la Licence du gouvernement ouvert – {region}."
        ),
        "cite": "Citez l'éditeur et l'URL de la source ci-dessus.",
        "statcan_source": (
            "Source : Statistique Canada, {product}, {date}. Reproduit et diffusé « tel quel » "
            "avec la permission de Statistique Canada."
        ),
        "statcan_adapted": (
            "Adapté de Statistique Canada, {product}, {date}. Cela ne constitue pas une "
            "approbation de ce produit par Statistique Canada."
        ),
        "table": "Tableau",
        "changes_flag": "Les données ont été modifiées : voir Modifications apportées.",
        "changes": "Modifications apportées",
        "changes_list": "{changes}.",
        "no_changes": "Aucune : noms de colonnes et valeurs tels que l'outil les a renvoyés.",
        "renamed": "noms de colonnes mis en snake_case ({count} sur {total} : {examples})",
        "trimmed": "cellules dont les espaces superflus ont été retirés : {count}",
        "emptied": "cellules laissées vides (texte vide ou nombre non fini) : {count}",
        "converted": "texte converti en nombres ou en dates dans {columns}",
        "long_ints": "nombres entiers de plus de 15 chiffres stockés comme texte dans {columns}",
        "cleaning_text": (
            "Noms en snake_case (comme janitor::clean_names en R), espaces superflus retirés "
            "du texte, texte vide traité comme valeur manquante, nombres et dates ISO stockés "
            "sous forme de texte, convertis en valeurs ; codes à zéros initiaux, et colonnes "
            "dont le nom désigne un code (coordonnée, identifiant de vecteur ou de produit, "
            "DGUID, code postal, SCIAN, CNP, CGT), conservés comme texte."
        ),
        "cleaned_note": (
            "Nettoyé comme dans les scripts R : noms de colonnes en snake_case, espaces "
            "superflus retirés du texte, texte vide traité comme valeur manquante."
        ),
        "converted_note": "Colonnes converties de texte en nombres ou en dates : {columns}.",
        "codes_note": "Colonnes conservées comme texte parce qu'elles contiennent des codes : {columns}.",
        "left_out": "Champs du résultat",
        "left_out_note": (
            "Absent des lignes de données (autres parties du résultat de l'outil) : {fields}. "
            "Leurs valeurs figurent dans la feuille Source ; les listes y sont seulement nommées."
        ),
        "reproduce_text": (
            "Appelez reproduce_code avec le même outil et les mêmes arguments pour obtenir des "
            "scripts R, Python, Stata et Julia, ou avec language='excel' pour une requête "
            "Power Query qui actualise toutes les données."
        ),
        "default_title": "Données MapleStats",
        "chart_text": "graphique {kind} de {columns} selon {x}",
        "line": "linéaire",
        "bar": "à barres",
        "max_rows": "max_rows doit être compris entre 1 et {limit}; valeur reçue : {got}.",
        "no_call": "Indiquez tool_name (et ses arguments), ou les lignes à écrire.",
        "not_data": "{tool} ne renvoie pas de lignes de données à écrire dans un classeur.",
        "too_wide": (
            "Le classeur dépasse {columns} colonnes, {cells} cellules ou 8 millions de "
            "caractères. Choisissez moins de colonnes ou de lignes."
        ),
        "delivery": "delivery doit valoir 'auto', 'base64' ou 'file'.",
        "hosted": (
            "Ce serveur est hébergé : il ne peut pas écrire sur votre disque. Utilisez "
            "delivery='base64'."
        ),
        "rows_shape": "rows doit être une liste d'objets, un par ligne.",
        "no_rows": "{what} n'a renvoyé aucune ligne à écrire dans un classeur.",
        "the_rows": "La liste rows",
        "too_big": (
            "Le classeur fait {size} octets, au-delà de la limite de {limit} octets d'une "
            "réponse d'outil. Réduisez max_rows, ou appelez reproduce_code avec "
            "language='excel' pour une requête Power Query qui charge toutes les données dans "
            "Excel."
        ),
        "failed": "{tool} a échoué avec ces arguments : {text}",
        "unstructured": "{tool} n'a pas renvoyé de résultat structuré.",
        "first_rows": "Les {written} premières lignes sur {available} ont été écrites (max_rows).",
        "saved": "Enregistré dans {path}.",
        "decode": (
            "Décodez workbook_base64 et enregistrez-le sous le nom {file} (un serveur hébergé "
            "ne peut pas écrire sur votre disque). Les gros résultats passent par le contexte "
            "du client : gardez max_rows petit, ou utilisez reproduce_code avec "
            "language='excel' pour toutes les données."
        ),
        "limits_text": (
            "au maximum {rows} lignes (max_rows, {default} par défaut) et {size} octets par classeur"
        ),
    },
}


# The federal, provincial and territorial governments that issue an Open Government
# Licence, in English and in French. A licence from anyone else (a city, say) gets the
# generic citation line rather than a jurisdiction we would have to guess.
_OGL_REGIONS = (
    ("Canada", "Canada"),
    ("Alberta", "Alberta"),
    ("British Columbia", "Colombie-Britannique"),
    ("Ontario", "Ontario"),
    ("Yukon", "Yukon"),
    ("Newfoundland and Labrador", "Terre-Neuve-et-Labrador"),
    ("Northwest Territories", "Territoires du Nord-Ouest"),
    ("Nova Scotia", "Nouvelle-Écosse"),
    ("New Brunswick", "Nouveau-Brunswick"),
    ("Saskatchewan", "Saskatchewan"),
    ("Manitoba", "Manitoba"),
)

# "Open Government Licence - <jurisdiction>" and its French form, hyphen or en dash.
_OGL_NAME = re.compile(
    r"(?:open government licen[cs]e|licence du gouvernement ouvert)\s*[-–—]\s*(?P<rest>[^.;:(]*)",
    re.IGNORECASE,
)


def _licence_region(licence: str | None, lang: str) -> str | None:
    """The jurisdiction the licence text names, in the workbook's language."""
    for named in _OGL_NAME.finditer(licence or ""):
        rest = named.group("rest").strip().lower()
        for english, french in _OGL_REGIONS:
            if rest.startswith((english.lower(), french.lower())):
                return french if normalize_lang(lang) == "fr" else english
    return None


def _labels(lang: str) -> dict[str, str]:
    return LABELS["fr" if normalize_lang(lang) == "fr" else "en"]


def _say(lang: str, key: str, **values: object) -> str:
    """A label filled in; French gets no-break spacing and 1 234-style numbers."""
    french = normalize_lang(lang) == "fr"
    filled = {
        k: (f"{v:,}".replace(",", NBSP) if french else f"{v:,}") if isinstance(v, int) else v
        for k, v in values.items()
    }
    text = _labels(lang)[key].format(**filled)
    return french_text(text) if french else text


# Rows ------------------------------------------------------------------------


async def _call(tool: str, arguments: dict[str, Any], lang: str = "en") -> Any:
    """The tool's structured result, called through the server like any client."""
    from fastmcp import Client
    from fastmcp.exceptions import DisabledError, NotFoundError
    from mcp.types import TextContent

    from maplestats_mcp.server import mcp

    # A French workbook needs the source text in French too: ask for it when the tool
    # takes a lang argument and the caller did not choose one.
    if normalize_lang(lang) == "fr" and "lang" not in arguments:
        try:
            target = await mcp.get_tool(tool)
        except (NotFoundError, DisabledError):  # an unknown tool is reported by the call below
            target = None
        if target is not None and "lang" in target.parameters.get("properties", {}):
            arguments = {**arguments, "lang": "fr"}

    async with Client(mcp) as client:
        result = await client.call_tool(
            "call_tool", {"name": tool, "arguments": arguments}, raise_on_error=False
        )
    text = next((block.text for block in result.content if isinstance(block, TextContent)), "")
    if result.is_error:
        raise InvalidInput(_say(lang, "failed", tool=tool, text=text[:300]))
    try:
        return json.loads(text)
    except ValueError as exc:
        raise NotFound(_say(lang, "unstructured", tool=tool)) from exc


def _record_lists(payload: Any) -> list[list[dict[str, Any]]]:
    """Every list of objects in the payload, three levels deep, provenance aside."""
    found: list[list[dict[str, Any]]] = []
    frontier: list[Any] = [payload]
    for _ in range(4):
        next_frontier: list[Any] = []
        for node in frontier:
            if isinstance(node, list) and node and all(isinstance(v, dict) for v in node):
                found.append(node)
            elif isinstance(node, dict):
                next_frontier += [v for k, v in node.items() if k != "provenance"]
        frontier = next_frontier
    return found


def _flatten(record: dict[str, Any], prefix: str = "") -> dict[str, Any]:
    """Nested objects become dotted names; lists become text."""
    check_deadline()
    flat: dict[str, Any] = {}
    for key, value in record.items():
        if key == "provenance":
            continue  # each item's provenance goes to the Source sheet, not the rows
        name = f"{prefix}{key}"
        if isinstance(value, dict):
            flat.update(_flatten(value, f"{name}."))
        elif isinstance(value, list):
            scalar = all(not isinstance(v, dict | list) for v in value)
            flat[name] = (
                "; ".join("" if v is None else str(v) for v in value)
                if scalar
                else json.dumps(value, ensure_ascii=False, default=str)
            )
        else:
            flat[name] = value
    return flat


def payload_provenance(payload: Any) -> dict[str, Any]:
    """The result's provenance, or its first item's for a list result (WDS)."""
    if isinstance(payload, dict) and isinstance(payload.get("result"), list):
        payload = payload["result"]
    if isinstance(payload, list) and payload and isinstance(payload[0], dict):
        payload = payload[0]
    found = payload.get("provenance") if isinstance(payload, dict) else None
    return found if isinstance(found, dict) else {}


def _left_out(node: Any, best: list[dict[str, Any]], prefix: str = "") -> dict[str, Any]:
    """The result's fields outside the rows, by dotted name.

    Scalars keep their value (a total, a truncated flag, a dataset id) so
    the Source sheet can show them; a list or object outside the path to
    the rows is named only, as "[list]" or "[object]".
    """
    left: dict[str, Any] = {}
    if not isinstance(node, dict):
        return left
    for key, value in node.items():
        name = f"{prefix}{key}"
        if key == "provenance" or value is best:
            continue
        if isinstance(value, dict) and _holds(value, best):
            left.update(_left_out(value, best, f"{name}."))
        elif isinstance(value, dict):
            left[name] = "[object]"
        elif isinstance(value, list):
            left[name] = "[list]"
        else:
            left[name] = value
    return left


def _holds(node: Any, best: list[dict[str, Any]]) -> bool:
    if node is best:
        return True
    if isinstance(node, dict):
        return any(_holds(v, best) for k, v in node.items() if k != "provenance")
    return False


def rows_from_payload(payload: Any) -> list[dict[str, Any]]:
    """The table a tool result holds: its largest list of objects.

    When each object holds its own list of objects (a series and its
    observations: sdmx_get_data, wds_get_data_from_vectors), every inner row
    is one row, carrying the outer object's other fields.
    """
    return rows_and_left_out(payload)[0]


def rows_and_left_out(payload: Any) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """`rows_from_payload`, plus the result's fields the rows leave out.

    A result often carries a total, a truncated flag or a second list beside
    its rows; the workbook says which (see `_left_out`) instead of dropping
    them unseen.
    """
    if isinstance(payload, dict) and isinstance(payload.get("result"), list):
        payload = payload["result"]  # FastMCP wraps a list result under "result"
    lists = _record_lists(payload)
    if not lists:
        if isinstance(payload, dict):
            single = {k: v for k, v in payload.items() if k != "provenance"}
            return ([_flatten(single)] if single else []), {}
        return [], {}
    best = max(lists, key=len)
    left = _left_out(payload, best)
    child_keys = [
        key
        for key, value in best[0].items()
        if isinstance(value, list) and value and all(isinstance(v, dict) for v in value)
    ]
    if not child_keys:
        return [_flatten(row) for row in best], left
    key = max(child_keys, key=lambda k: sum(len(row.get(k) or []) for row in best))
    rows = []
    for parent in best:
        outer = _flatten({k: v for k, v in parent.items() if k != key})
        rows += [{**outer, **_flatten(child)} for child in parent.get(key) or []]
    return rows, left


# Cleaning, as the R scripts clean ---------------------------------------------


def clean_name(name: str) -> str:
    plain = unicodedata.normalize("NFKD", str(name)).encode("ascii", "ignore").decode()
    split = re.sub(r"([a-z0-9])([A-Z])", r"\1_\2", plain)
    return re.sub(r"[^0-9a-z]+", "_", split.lower()).strip("_") or "x"


def clean_names(names: list[str]) -> list[str]:
    """janitor's clean_names: snake_case, and names that clean alike numbered."""
    cleaned = [clean_name(n) for n in names]
    taken: set[str] = set()
    out = []
    for name in cleaned:
        candidate, suffix = name, 2
        while candidate in taken:
            candidate = f"{name}_{suffix}"
            suffix += 1
        taken.add(candidate)
        out.append(candidate)
    return out


_NUMBER = re.compile(r"[-+]?(\d+(\.\d*)?|\.\d+)([eE][-+]?\d+)?")
_DATE = re.compile(r"\d{4}-\d{2}-\d{2}")
_DATETIME = re.compile(r"\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}(:\d{2}(\.\d+)?)?(Z|[+-]\d{2}:?\d{2})?")


def _convert(text: str) -> Any:
    """A number or date written as text, else the text (type_convert in R)."""
    if _NUMBER.fullmatch(text):
        # A code such as "01" or "0042" keeps its leading zeros.
        unsigned = text.lstrip("-+")
        if len(unsigned) > 1 and unsigned[0] == "0" and unsigned[1] != ".":
            return text
        # Excel preserves at most 15 significant digits; long codes must stay text.
        if re.fullmatch(r"[-+]?\d+", text):
            return text if len(unsigned) > 15 else int(text)
        number = float(text)
        return number if math.isfinite(number) else text
    if _DATE.fullmatch(text):
        try:
            return date.fromisoformat(text)
        except ValueError:
            return text
    if _DATETIME.fullmatch(text):
        try:
            parsed = datetime.fromisoformat(text)
        except ValueError:
            return text
        # Excel has no time zones: UTC, written without the offset.
        return parsed.astimezone(UTC).replace(tzinfo=None) if parsed.tzinfo else parsed
    return text


# Same rule as the scripts (cleaning.CODE_COLUMNS): StatCan coordinates "2.1"
# and "2.10" are different members and collide as numbers, and a NAICS or
# NOC code is not a quantity. Matched on the cleaned name, and on the last
# part of a flattened one ("series.coordinate").
_CODE_COLUMN = re.compile(CODE_COLUMNS)


def _is_code_column(name: str) -> bool:
    return any(_CODE_COLUMN.fullmatch(clean_name(part)) for part in {name, name.rsplit(".", 1)[-1]})


def clean_rows(rows: list[dict[str, Any]], lang: str = "en") -> tuple[list[str], list[list[Any]]]:
    """Column names and cleaned values; a column converts only if all its values do."""
    names, table, _, _ = clean_rows_report(rows, lang)
    return names, table


def clean_rows_report(
    rows: list[dict[str, Any]], lang: str = "en"
) -> tuple[list[str], list[list[Any]], list[str], list[str]]:
    """`clean_rows`, plus the cleaned names of the columns converted from text
    and of the code columns kept as text although their values read as numbers."""
    columns: list[str] = []
    seen: set[str] = set()
    text_chars = 0
    for row in rows:
        check_deadline()
        for key, value in row.items():
            if key not in seen:
                seen.add(key)
                columns.append(key)
            text_chars += len(value) if isinstance(value, str) else 0
            if (
                len(columns) > MAX_COLUMNS
                or len(columns) * len(rows) > MAX_CELLS
                or text_chars > MAX_TEXT_CHARS
            ):
                raise InvalidInput(_say(lang, "too_wide", columns=MAX_COLUMNS, cells=MAX_CELLS))
    table: list[list[Any]] = []
    long_int_columns: set[int] = set()
    for row in rows:
        check_deadline()
        values = []
        for index, column in enumerate(columns):
            value = row.get(column)
            if isinstance(value, str):
                value = value.strip() or None
            elif isinstance(value, float) and not math.isfinite(value):
                value = None
            elif (
                isinstance(value, int) and not isinstance(value, bool) and len(str(abs(value))) > 15
            ):
                value = str(value)
                long_int_columns.add(index)
            values.append(value)
        table.append(values)
    names = clean_names(columns)
    converted_columns: list[str] = []
    code_columns: list[str] = []
    for index in range(len(columns)):
        check_deadline()
        texts = [r[index] for r in table if isinstance(r[index], str)]
        if not texts:
            continue
        converted = [_convert(t) for t in texts]
        kinds = {type(v) for v in converted}
        if not (kinds <= {int, float} or kinds == {date} or kinds == {datetime}):
            continue  # still text, or mixed kinds: leave the column as text
        if kinds <= {int, float} and _is_code_column(columns[index]):
            code_columns.append(names[index])
            continue
        converted_columns.append(names[index])
        lookup = dict(zip(texts, converted, strict=True))
        for r in table:
            if isinstance(r[index], str):
                r[index] = lookup[r[index]]
    # A code too long for Excel's 15 digits makes its whole column text, so it
    # never mixes numbers and text.
    for index in long_int_columns:
        for r in table:
            if isinstance(r[index], int) and not isinstance(r[index], bool):
                r[index] = str(r[index])
    return names, table, converted_columns, code_columns


# Workbook ------------------------------------------------------------------------


def _sheet_name(text: str, taken: set[str]) -> str:
    name = re.sub(r"[\[\]:*?/\\]", " ", text).strip().strip("'")[:31] or "Data"
    base, n = name, 2
    while name.lower() in {t.lower() for t in taken}:
        suffix = f" ({n})"
        name, n = base[: 31 - len(suffix)] + suffix, n + 1
    taken.add(name)
    return name


def _column_kind(name: str, values: list[Any]) -> str:
    present = [v for v in values if v is not None]
    if not present:
        return "empty"
    if all(isinstance(v, bool) for v in present):
        return "text"
    if all(isinstance(v, datetime) for v in present):
        return "datetime"
    if all(isinstance(v, date) and not isinstance(v, datetime) for v in present):
        return "date"
    if all(isinstance(v, int | float) and not isinstance(v, bool) for v in present):
        integers = all(isinstance(v, int) or float(v).is_integer() for v in present)
        identifier = re.search(r"(^|_)(id|code|vector|year|annee|pid|key|dguid)($|_)", name)
        if integers and (identifier or name in ("d", "yr")):
            return "identifier"
        return "integer" if integers else "decimal"
    return "text"


def _decimals(values: list[Any]) -> int:
    places = 0
    for value in values[:5000]:
        if isinstance(value, float) and math.isfinite(value):
            text = f"{value:.6f}".rstrip("0")
            places = max(places, len(text.split(".")[1]) if "." in text else 0)
    return min(places, 6)


def number_format(kind: str, values: list[Any]) -> str:
    """Explicit formats (section 11): never Excel's auto-detection."""
    if kind == "integer":
        return "#,##0"
    if kind == "decimal":
        places = max(_decimals(values), 1)
        return "#,##0." + "0" * places
    if kind == "date":
        return "yyyy-mm-dd"
    if kind == "datetime":
        return "yyyy-mm-dd hh:mm:ss"
    if kind == "identifier":
        return "0"
    return "@"


def _cell_value(value: Any) -> Any:
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
        return value[: _MAX_CELL - 1] + "…" if len(value) > _MAX_CELL else value
    if isinstance(value, bool):
        return str(value).upper()
    return value


def _write_table(ws: Any, header: list[str], rows: list[list[Any]], table_name: str) -> list[str]:
    """One Excel table object on the sheet: header, rows, formats, widths, frozen header."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        # Chart headers are series values from the data, so a name starting with
        # "=" must stay text: a formula header breaks the table's column names.
        cell.data_type = "s"
    for row in rows:
        check_deadline()
        ws.append([_cell_value(v) for v in row])
    kinds = []
    for index, name in enumerate(header, start=1):
        check_deadline()
        values = [row[index - 1] for row in rows]
        kind = _column_kind(name, values)
        kinds.append(kind)
        fmt = number_format(kind, values)
        letter = get_column_letter(index)
        for cell in ws[letter][1:]:
            cell.number_format = fmt
            # A text cell starting with "=" stays text, never a formula.
            if isinstance(cell.value, str):
                cell.data_type = "s"
        sample = [name, *("" if v is None else str(v) for v in values[:500])]
        width = max(len(text) for text in sample) + 2
        if kind == "date":
            width = max(width, 12)
        ws.column_dimensions[letter].width = min(max(width, 8), 60)
    ref = f"A1:{get_column_letter(len(header))}{len(rows) + 1}"
    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "A2"
    return kinds


_VALUE_NAMES = ("value", "obs_value", "valeur", "value_normalized", "count", "total", "amount")
_SERIES_NAMES = (
    "series", "series_name", "vector_id", "vector", "geo", "geography", "name", "label",
    "indicator", "region", "province",
)  # fmt: skip
_PERIOD_NAMES = re.compile(
    r"(^|_)(date|ref_date|ref_per|time_period|period|periode|d|year|annee|month|mois)$"
)
_PERIOD_TEXT = re.compile(r"\d{4}(-\d{2}(-\d{2})?|-?Q[1-4])?")


def _chart_plan(
    header: list[str], rows: list[list[Any]], kinds: list[str]
) -> tuple[str, int, list[tuple[str, int]], int | None] | None:
    """(chart type, x column, value columns by name, series column) or None."""
    if len(rows) < 2:
        return None
    numeric = [i for i, k in enumerate(kinds) if k in ("integer", "decimal")]
    if not numeric:
        return None

    def is_period(i: int) -> bool:
        return kinds[i] in ("date", "datetime") or all(
            v is None or _PERIOD_TEXT.fullmatch(str(v)) for v in (r[i] for r in rows)
        )

    named = [i for i, name in enumerate(header) if _PERIOD_NAMES.search(name) and is_period(i)]
    dated = [i for i, kind in enumerate(kinds) if kind in ("date", "datetime")]
    period = (named or dated or [None])[0]
    value = next((i for name in _VALUE_NAMES for i in numeric if header[i] == name), None)
    if period is not None:
        series = None
        if value is not None:
            series = next(
                (
                    i
                    for name in _SERIES_NAMES
                    for i, column in enumerate(header)
                    if column == name
                    and i not in (period, value)
                    and 2 <= len({r[i] for r in rows}) <= len(PALETTE)
                ),
                None,
            )
            values = [value]
        else:
            values = [i for i in numeric if kinds[i] == "decimal" and i != period][: len(PALETTE)]
            values = values or [i for i in numeric if i != period][:1]
        # A line needs one value per period in each series; otherwise the rows
        # are not a time series (several places per date, say).
        keys = [(r[period], r[series] if series is not None else None) for r in rows]
        if values and len(set(keys)) == len(keys):
            return "line", period, [(header[i], i) for i in values], series
        return None
    labels = [i for i, k in enumerate(kinds) if k == "text"]
    if labels and len(rows) <= 30:
        target = value if value is not None else numeric[0]
        return "bar", labels[0], [(header[target], target)], None
    return None


def _add_chart(
    wb: Any, plan: tuple[Any, ...], header, rows, labels: dict[str, str], title: str, taken
) -> str:
    from openpyxl.chart import BarChart, LineChart, Reference

    chart_type, x, values, series = plan
    if series is not None:
        names = sorted({r[series] for r in rows if r[series] is not None}, key=str)
        wide: dict[Any, dict[Any, Any]] = {}
        for r in rows:
            if r[x] is not None:
                wide.setdefault(r[x], {})[r[series]] = r[values[0][1]]
        columns = [header[x], *(str(n) for n in names)]
        table = [[key, *(wide[key].get(n) for n in names)] for key in sorted(wide, key=str)]
    else:
        columns = [header[x], *(name for name, _ in values)]
        table = [[r[x], *(r[i] for _, i in values)] for r in rows if r[x] is not None]
        if chart_type == "line":
            table.sort(key=lambda r: str(r[0]))
    data_ws = wb.create_sheet(_sheet_name(labels["chart_data"], taken))
    _write_table(data_ws, columns, table, "chart_data")
    chart = LineChart() if chart_type == "line" else BarChart()
    chart.title = title[:250]
    chart.y_axis.title = values[0][0] if len(values) == 1 and series is None else labels["value"]
    chart.x_axis.title = header[x]
    chart.height, chart.width = 12, 24
    data = Reference(data_ws, min_col=2, max_col=len(columns), min_row=1, max_row=len(table) + 1)
    categories = Reference(data_ws, min_col=1, min_row=2, max_row=len(table) + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    if chart.legend is not None:
        chart.legend.position = "b"
    for index, line in enumerate(chart.series):
        colour = PALETTE[index % len(PALETTE)]
        if chart_type == "line":
            line.graphicalProperties.line.solidFill = colour
            line.graphicalProperties.line.width = 22000
            line.smooth = False
        else:
            line.graphicalProperties.solidFill = colour
    sheet = wb.create_chartsheet(_sheet_name(labels["chart"], taken))
    sheet.add_chart(chart)
    return labels["chart_text"].format(
        kind=labels[chart_type], columns=", ".join(columns[1:4]), x=header[x]
    )


def _licence(provenance: dict[str, Any]) -> str | None:
    for key, value in provenance.items():
        if "licen" in key.lower() and value:
            return value if isinstance(value, str) else json.dumps(value, ensure_ascii=False)
    return None


def _changes(
    rows: list[dict[str, Any]], header: list[str], converted: list[str], lang: str
) -> list[str]:
    """What cleaning changed from the rows as the tool returned them, for the Source sheet.

    Some publishers' terms depend on it: Statistics Canada's Open Licence asks
    for "Adapted from" rather than "Source:" once the data are changed, and the
    Bank of Canada and Creative Commons Attribution ask that changes be indicated.
    """
    columns: list[str] = []
    seen: set[str] = set()
    trimmed = emptied = 0
    long_ints: set[str] = set()
    for row in rows:
        check_deadline()
        for key, value in row.items():
            if key not in seen:
                seen.add(key)
                columns.append(key)
            if isinstance(value, str):
                stripped = value.strip()
                trimmed += bool(stripped) and stripped != value
                emptied += not stripped
            elif isinstance(value, float) and not math.isfinite(value):
                emptied += 1
            elif (
                isinstance(value, int) and not isinstance(value, bool) and len(str(abs(value))) > 15
            ):
                long_ints.add(key)
    renamed = [f"{old} → {new}" for old, new in zip(columns, header, strict=True) if old != new]
    changes = []
    if renamed:
        examples = ", ".join(renamed[:5]) + (", …" if len(renamed) > 5 else "")
        changes.append(
            _say(lang, "renamed", count=len(renamed), total=len(columns), examples=examples)
        )
    if trimmed:
        changes.append(_say(lang, "trimmed", count=trimmed))
    if emptied:
        changes.append(_say(lang, "emptied", count=emptied))
    if converted:
        changes.append(_say(lang, "converted", columns=", ".join(converted)))
    if long_ints:
        named = [new for old, new in zip(columns, header, strict=True) if old in long_ints]
        changes.append(_say(lang, "long_ints", columns=", ".join(named)))
    return changes


def _plain(text: str | None) -> str:
    """French typography's no-break spaces as plain ones, for comparing licence texts."""
    return (text or "").replace(NBSP, " ").replace("\u202f", " ").strip()


# Tool results carry the French texts with French spacing (make_provenance).
_STATCAN = {_plain(STATCAN_LICENCE), _plain(french_spacing(STATCAN_LICENCE_FR))}
_BOC = {_plain(BOC_TERMS), _plain(french_spacing(LICENCES_FR[BOC_TERMS]))}
# CC BY and its NC/SA/ND variants ask that changes be indicated; CC0 does not.
_CC_BY = re.compile(r"\bCC[ -]BY\b|Creative Commons Attribution", re.IGNORECASE)


def _product(title: str, tool: str, arguments: dict[str, Any] | None, lang: str) -> str | None:
    """The product a StatCan citation names: its table number, and the title the caller gave."""
    labels = _labels(lang)
    named = title if title and title not in {tool, labels["default_title"]} else None
    pid = next(
        (
            str((arguments or {}).get(key))
            for key in ("product_id", "productId", "pid", "table_id")
            if re.fullmatch(r"\d{8}(\d{2})?", str((arguments or {}).get(key, "")))
        ),
        None,
    )
    table = (
        f"{labels['table']} {pid[:2]}-{pid[2:4]}-{pid[4:8]}-{pid[8:10] or '01'}" if pid else None
    )
    return " ".join(part for part in (table, named) if part) or None


def _attribution(
    licence: str | None,
    provenance: dict[str, Any],
    changed: bool,
    product: str | None,
    lang: str,
) -> str:
    """The credit line the publisher's terms ask for, in the form that fits the changes."""
    labels = _labels(lang)
    if _plain(licence) in _STATCAN:
        when = str(provenance.get("as_of") or provenance.get("queried_at") or "")[:10]
        key = "statcan_adapted" if changed else "statcan_source"
        return _say(
            lang,
            key,
            product=product or provenance.get("url") or "",
            date=when or datetime.now(UTC).strftime("%Y-%m-%d"),
        )
    open_licence = re.search(r"open government licen|gouvernement ouvert", (licence or "").lower())
    region = _licence_region(licence, lang) if open_licence else None
    credit = labels["ogl_attribution"].format(region=region) if region else labels["cite"]
    flag = changed and (
        _plain(licence) in _BOC
        or provenance.get("source") in {"boc", "valet"}
        or bool(_CC_BY.search(licence or ""))
    )
    return f"{credit} {_say(lang, 'changes_flag')}" if flag else credit


def build_workbook(
    rows: list[dict[str, Any]],
    *,
    title: str,
    tool: str,
    arguments: dict[str, Any] | None,
    provenance: dict[str, Any],
    available: int,
    lang: str = "en",
    left_out: dict[str, Any] | None = None,
) -> tuple[bytes, list[str], str | None, list[str]]:
    """The .xlsx bytes, its sheet names, a description of its chart and notes
    on how the rows were cleaned."""
    from openpyxl import Workbook

    labels = _labels(lang)
    header, table, converted, codes = clean_rows_report(rows, lang)
    cleaning = [_say(lang, "cleaned_note")]
    if converted:
        cleaning.append(_say(lang, "converted_note", columns=", ".join(converted)))
    if codes:
        cleaning.append(_say(lang, "codes_note", columns=", ".join(codes)))
    wb = Workbook()
    taken: set[str] = set()
    data_ws = wb.worksheets[0]
    data_ws.title = _sheet_name(title or labels["data"], taken)
    kinds = _write_table(data_ws, header, table, "maplestats_data")

    plan = _chart_plan(header, table, kinds)
    chart = _add_chart(wb, plan, header, table, labels, title, taken) if plan else None

    licence = _licence(provenance)
    changes = _changes(rows, header, converted, lang)
    product = _product(title, tool, arguments, lang)
    attribution = _attribution(licence, provenance, bool(changes), product, lang)
    capped = labels["rows_capped"] if len(table) < available else ""
    notes = [
        (labels["title"], title),
        (labels["produced_by"], "MapleStats MCP, reproduce_workbook"),
        (labels["tool"], tool or labels["caller_rows"]),
        (labels["arguments"], json.dumps(arguments or {}, ensure_ascii=False, default=str)),
        (labels["source_name"], provenance.get("source")),
        (labels["source_url"], provenance.get("url")),
        (labels["queried_at"], provenance.get("queried_at")),
        (labels["as_of"], provenance.get("as_of")),
        (labels["coverage"], provenance.get("coverage")),
        (labels["limits"], provenance.get("limits")),
        (labels["licence"], licence or labels["no_licence"]),
        (labels["attribution"], attribution),
        (
            labels["rows"],
            _say(lang, "rows_text", written=len(table), available=available) + capped,
        ),
        (labels["cleaning"], labels["cleaning_text"]),
        (
            labels["changes"],
            _say(lang, "changes_list", changes="; ".join(changes))
            if changes
            else _say(lang, "no_changes"),
        ),
        (
            labels["left_out"],
            json.dumps(left_out, ensure_ascii=False, default=str)[:_MAX_CELL] if left_out else None,
        ),
        (labels["reproduce"], labels["reproduce_text"]),
        (labels["created"], datetime.now(UTC).strftime("%Y-%m-%d %H:%M:%S")),
    ]
    source_ws = wb.create_sheet(_sheet_name(labels["source"], taken))
    _write_table(
        source_ws,
        [labels["field"], labels["value"]],
        [[field, "" if value is None else str(value)] for field, value in notes],
        "source_notes",
    )
    source_ws.column_dimensions["B"].width = 100
    wb.properties.title = title
    wb.properties.creator = "MapleStats MCP"
    buffer = io.BytesIO()
    check_deadline()
    wb.save(buffer)
    check_deadline()
    return buffer.getvalue(), wb.sheetnames, chart, cleaning


def _export_dir() -> Path:
    configured = os.environ.get("MAPLE_EXPORT_DIR", "").strip()
    if configured:
        return Path(configured).expanduser()
    downloads = Path.home() / "Downloads"
    return downloads if downloads.is_dir() else Path.home()


def _file_name(title: str) -> str:
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", unicodedata.normalize("NFKD", title))
    stem = stem.encode("ascii", "ignore").decode().strip("._")[:60] or "maplestats"
    return f"{stem}_{datetime.now(UTC).strftime('%Y%m%d_%H%M%S')}.xlsx"


async def export(
    tool_name: str | None,
    arguments: dict[str, Any] | None,
    rows: list[dict[str, Any]] | None,
    title: str | None,
    lang: str,
    max_rows: int,
    delivery: Delivery,
) -> ExcelWorkbook:
    from maplestats_mcp import config

    if not 1 <= max_rows <= MAX_ROWS:
        raise InvalidInput(_say(lang, "max_rows", limit=MAX_ROWS, got=str(max_rows)))
    if rows is None and not tool_name:
        raise InvalidInput(_say(lang, "no_call"))
    if tool_name in _NOT_DATA:
        raise InvalidInput(_say(lang, "not_data", tool=tool_name))
    if delivery not in ("auto", "base64", "file"):
        raise InvalidInput(_say(lang, "delivery"))
    local = config.get_transport() == "stdio"
    if delivery == "file" and not local:
        raise InvalidInput(_say(lang, "hosted"))
    provenance: dict[str, Any] = {}
    left_out: dict[str, Any] = {}
    if rows is None:
        payload = await _call(str(tool_name), arguments or {}, lang)
        provenance = payload_provenance(payload)
        records, left_out = await run_parse(rows_and_left_out, payload)
        available = len(records)
    else:
        if not all(isinstance(row, dict) for row in rows):
            raise InvalidInput(_say(lang, "rows_shape"))
        available = len(rows)
        records = await run_parse(lambda: [_flatten(row) for row in rows[:max_rows]])
    if not records:
        what = tool_name or _labels(lang)["the_rows"]
        raise NotFound(_say(lang, "no_rows", what=what))
    title = (title or tool_name or _labels(lang)["default_title"]).strip()[:120]
    content, sheets, chart, cleaning = await run_parse(
        build_workbook,
        records[:max_rows],
        title=title,
        tool=tool_name or "",
        arguments=arguments,
        provenance=provenance,
        available=available,
        lang=lang,
        left_out=left_out,
    )
    if len(content) > MAX_BYTES:
        raise InvalidInput(_say(lang, "too_big", size=len(content), limit=MAX_BYTES))
    file_name = _file_name(title)
    notes = []
    if available > max_rows:
        notes.append(_say(lang, "first_rows", written=max_rows, available=available))
    notes += cleaning
    if left_out:
        notes.append(_say(lang, "left_out_note", fields=", ".join(left_out)))
    saved_path = None
    encoded = None
    if delivery == "file" or (delivery == "auto" and local):
        folder = _export_dir()
        folder.mkdir(parents=True, exist_ok=True)
        target = folder / file_name
        target.write_bytes(content)
        saved_path = str(target)
        notes.append(_say(lang, "saved", path=saved_path))
    else:
        encoded = base64.b64encode(content).decode("ascii")
        notes.append(_say(lang, "decode", file=file_name))
    return ExcelWorkbook(
        file_name=file_name,
        media_type=MEDIA_TYPE,
        size_bytes=len(content),
        sheets=sheets,
        rows_written=min(available, max_rows),
        rows_available=available,
        truncated=available > max_rows,
        chart=chart,
        saved_path=saved_path,
        workbook_base64=encoded,
        notes=notes,
        provenance=make_provenance(
            source="maplestats-workbook",
            url=str(provenance.get("url") or "about:blank"),
            cached=False,
            schema_name="reproduce.ExcelWorkbook",
            limits=_say(lang, "limits_text", rows=MAX_ROWS, default=DEFAULT_ROWS, size=MAX_BYTES),
            lang=lang,
        ),
    )
