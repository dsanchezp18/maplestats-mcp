"""Newfoundland and Labrador Statistics Agency: file listing and Excel reading.

The agency has no API. Its topic pages list the Excel files with a
description, and each file is read sheet by sheet (`.xlsx` with openpyxl,
legacy `.xls` with xlrd). Sheets keep the agency's layout, so rows come
back as published text with a guessed header row.
"""

from __future__ import annotations

import io
from datetime import date, datetime
from typing import Any, Literal
from urllib.parse import urljoin, urlparse

import httpx
from bs4 import BeautifulSoup, Tag

from maplestats_mcp.modules.nl_stats import constants
from maplestats_mcp.modules.nl_stats.schemas import (
    FileData,
    FileEntry,
    FileList,
    SheetInfo,
    TopicInfo,
)
from maplestats_mcp.shared.cache import cached_fetch
from maplestats_mcp.shared.envelope import make_provenance
from maplestats_mcp.shared.errors import InvalidInput, NotFound, UpstreamError, UpstreamUnavailable
from maplestats_mcp.shared.executor import run_parse
from maplestats_mcp.shared.http import get_raw
from maplestats_mcp.shared.rate_limiter import get_limiter

_LIMITER = get_limiter(
    constants.RATE_LIMIT_SOURCE,
    rate=constants.RATE_LIMIT_PER_SECOND,
    capacity=constants.RATE_LIMIT_CAPACITY,
)

Lang = Literal["en", "fr"]

_EXCEL_SUFFIXES = (".xlsx", ".xls")


def topic_title(topic: str, lang: Lang) -> str:
    en, fr = constants.TOPICS[topic]
    return en if lang == "en" else fr


def _own_text(li: Tag) -> str:
    """The text of a list item before its first link or nested list."""
    parts: list[str] = []
    for child in li.contents:
        if isinstance(child, Tag) and child.name in ("a", "ul"):
            break
        parts.append(child.get_text() if isinstance(child, Tag) else str(child))
    return " ".join("".join(parts).split()).strip(" :|")


def _title_for(link: Tag) -> str:
    li = link.find_parent("li")
    if li is None:
        return " ".join(link.get_text().split())
    title = _own_text(li)
    parent_li = li.find_parent("li")
    if parent_li is not None:
        parent = _own_text(parent_li)
        title = f"{parent}: {title}" if parent and title else parent or title
    return title


def parse_topic_page(html: str, topic: str) -> list[FileEntry]:
    soup = BeautifulSoup(html, "html.parser")
    container = soup.find(id="ContentPlaceHolder1_links")
    if container is None:
        return []
    files: list[FileEntry] = []
    seen: set[tuple[str, str]] = set()
    for link in container.find_all("a", href=True):
        href = str(link["href"])
        if not href.lower().endswith(_EXCEL_SUFFIXES):
            continue
        url = urljoin(constants.SITE + "/Statistics/", href)
        title = _title_for(link)
        if (url, title) in seen:
            continue
        seen.add((url, title))
        heading = link.find_previous("h4")
        files.append(
            FileEntry(
                topic=topic,
                section=" ".join(heading.get_text().split()) if heading else None,
                title=title,
                url=url,
                format=url.lower().rsplit(".", 1)[1],
            )
        )
    return files


async def _topic_files(topic: str) -> tuple[list[FileEntry], bool]:
    url = constants.TOPIC_PAGE_URL.format(topic=topic)

    async def fetch() -> list[FileEntry]:
        await _LIMITER.acquire()
        try:
            response = await get_raw(url, timeout=60.0)
        except httpx.HTTPStatusError as exc:
            raise UpstreamError(
                f"nl_stats: {url} returned HTTP {exc.response.status_code}."
            ) from exc
        except httpx.HTTPError as exc:
            raise UpstreamUnavailable(f"nl_stats: {url} did not respond in time.") from exc
        return parse_topic_page(response.text, topic)

    return await cached_fetch(f"nl_stats:topic:{topic}", constants.CACHE_TTL_PAGE_SECONDS, fetch)


async def list_files(
    topic: str | None = None,
    query: str | None = None,
    limit: int = constants.FILES_LIMIT_MAX,
    lang: Lang = "en",
) -> FileList:
    if topic is not None and topic not in constants.TOPICS:
        raise InvalidInput(f"nl_stats: topic must be one of {list(constants.TOPICS)}.")
    if not 1 <= limit <= constants.FILES_LIMIT_MAX:
        raise InvalidInput(f"nl_stats: limit must be 1 to {constants.FILES_LIMIT_MAX}.")
    chosen = [topic] if topic else list(constants.TOPICS)

    files: list[FileEntry] = []
    all_cached = True
    for name in chosen:
        entries, cached = await _topic_files(name)
        files.extend(entries)
        all_cached &= cached

    if query:
        words = query.casefold().split()
        files = [
            f
            for f in files
            if all(w in f"{f.title} {f.section or ''} {f.url}".casefold() for w in words)
        ]
    total = len(files)
    return FileList(
        topics=[
            TopicInfo(
                topic=name,
                title=topic_title(name, lang),
                page_url=constants.TOPIC_PAGE_URL.format(topic=name),
            )
            for name in constants.TOPICS
        ],
        files=files[:limit],
        total_files=total,
        truncated=total > limit,
        provenance=make_provenance(
            source=constants.PROVENANCE_SOURCE,
            url=constants.SITE + "/Statistics/Statistics.aspx",
            cached=all_cached,
            schema_name="nl_stats.FileList",
            freshness="The agency updates monthly, quarterly and annually by table.",
            coverage="Excel files only (each topic page also links PDFs); census, "
            "environment and justice topics publish no Excel files.",
            limits=f"Showing {limit} of {total} files." if total > limit else None,
        ),
    )


def check_file_url(url: str) -> None:
    parsed = urlparse(url)
    if (
        parsed.scheme != "https"
        or parsed.hostname != constants.DOMAIN
        or not parsed.path.startswith(constants.FILE_PATH_PREFIX)
        or not parsed.path.lower().endswith(_EXCEL_SUFFIXES)
    ):
        raise InvalidInput(
            f"nl_stats: url must be an https .xlsx or .xls link under "
            f"{constants.SITE}{constants.FILE_PATH_PREFIX} (see nl_stats_list_files)."
        )


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() and abs(value) < 1e15 else repr(value)
    if isinstance(value, datetime):
        return value.isoformat(sep=" ") if value.time() else value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return " ".join(str(value).split())


def _trim(rows: list[list[str]]) -> list[list[str]]:
    rows = rows[: constants.MAX_ROWS_PER_SHEET]
    while rows and not any(rows[-1]):
        rows.pop()
    width = max((max((i for i, c in enumerate(r) if c), default=-1) + 1 for r in rows), default=0)
    return [r[:width] + [""] * (width - len(r[:width])) for r in rows]


def _parse_xlsx(body: bytes) -> dict[str, list[list[str]]]:
    # openpyxl takes ~0.7 s to import, so it loads on first use.
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    try:
        return {
            sheet.title: _trim(
                [[_text(c) for c in row] for row in sheet.iter_rows(values_only=True)]
            )
            for sheet in workbook.worksheets
        }
    finally:
        workbook.close()


def _parse_xls(body: bytes) -> dict[str, list[list[str]]]:
    import xlrd

    book = xlrd.open_workbook(file_contents=body)
    sheets: dict[str, list[list[str]]] = {}
    for sheet in book.sheets():
        rows: list[list[str]] = []
        for r in range(min(sheet.nrows, constants.MAX_ROWS_PER_SHEET)):
            row: list[str] = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    serial = float(cell.value)
                    row.append(_text(xlrd.xldate.xldate_as_datetime(serial, book.datemode)))
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    row.append("")
                else:
                    row.append(_text(cell.value))
            rows.append(row)
        sheets[sheet.name] = _trim(rows)
    return sheets


async def _workbook(url: str) -> tuple[dict[str, list[list[str]]], bool]:
    async def fetch() -> dict[str, list[list[str]]]:
        await _LIMITER.acquire()
        try:
            response = await get_raw(url, timeout=120.0)
        except httpx.HTTPStatusError as exc:
            if exc.response.status_code == 404:
                raise NotFound(f"nl_stats: no file at {url}.") from exc
            raise UpstreamError(
                f"nl_stats: {url} returned HTTP {exc.response.status_code}."
            ) from exc
        except httpx.HTTPError as exc:
            raise UpstreamUnavailable(f"nl_stats: {url} did not respond in time.") from exc
        body = response.content
        if len(body) > constants.MAX_FILE_BYTES:
            raise UpstreamError(f"nl_stats: {url} is larger than this tool reads.")
        if body.lstrip()[:5].lower() in (b"<!doc", b"<html"):
            raise NotFound(f"nl_stats: {url} returned a web page, not an Excel file.")
        parse = _parse_xlsx if url.lower().endswith(".xlsx") else _parse_xls
        try:
            return await run_parse(parse, body)
        except Exception as exc:  # openpyxl and xlrd raise several unrelated types
            raise UpstreamError(f"nl_stats: could not read {url}: {exc}") from exc

    return await cached_fetch(f"nl_stats:file:{url}", constants.CACHE_TTL_FILE_SECONDS, fetch)


def guess_header(rows: list[list[str]]) -> int | None:
    """0-based index of the first row with at least three filled cells."""
    for index, row in enumerate(rows[:30]):
        if sum(1 for cell in row if cell) >= 3:
            return index
    return None


async def read_file(
    url: str,
    sheet: str | None = None,
    contains: str | None = None,
    limit: int = constants.ROWS_LIMIT_DEFAULT,
    offset: int = 0,
    lang: Lang = "en",
) -> FileData:
    check_file_url(url)
    if not 1 <= limit <= constants.ROWS_LIMIT_MAX:
        raise InvalidInput(f"nl_stats: limit must be 1 to {constants.ROWS_LIMIT_MAX}.")
    if offset < 0:
        raise InvalidInput("nl_stats: offset must be 0 or more.")

    sheets, cached = await _workbook(url)
    if not sheets:
        raise UpstreamError(f"nl_stats: {url} has no sheets.")
    if sheet is None:
        chosen = next(iter(sheets))
    else:
        matches = [name for name in sheets if name.casefold() == sheet.strip().casefold()]
        if not matches:
            raise InvalidInput(f"nl_stats: no sheet {sheet!r}; sheets are {list(sheets)}.")
        chosen = matches[0]

    rows = sheets[chosen]
    header_index = guess_header(rows)
    header = rows[header_index] if header_index is not None else []
    body = rows[header_index + 1 :] if header_index is not None else rows
    body = [r for r in body if any(r)]
    if contains:
        wanted = contains.casefold()
        body = [r for r in body if any(wanted in cell.casefold() for cell in r)]
    total = len(body)
    return FileData(
        url=url,
        format=url.lower().rsplit(".", 1)[1],
        sheets=[
            SheetInfo(name=name, rows=len(r), columns=len(r[0]) if r else 0)
            for name, r in sheets.items()
        ],
        sheet=chosen,
        header_row=None if header_index is None else header_index + 1,
        header=header,
        rows=body[offset : offset + limit],
        total_rows=total,
        offset=offset,
        truncated=offset + limit < total,
        provenance=make_provenance(
            source=constants.PROVENANCE_SOURCE,
            url=url,
            cached=cached,
            schema_name="nl_stats.FileData",
            freshness="As published by the NL Statistics Agency.",
            coverage=(
                f"Sheet {chosen!r} of {len(sheets)}."
                if lang == "en"
                else f"Feuille {chosen!r} sur {len(sheets)}."
            ),
            limits=(
                f"Sheets are read up to {constants.MAX_ROWS_PER_SHEET} rows; showing rows "
                f"{offset + 1} to {offset + min(limit, total - offset)} of {total}."
                if offset + limit < total
                else None
            ),
        ),
    )
