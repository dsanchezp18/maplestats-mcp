"""Tests for modules/cihi/client.py, shaped on live 2026-09-23 pages and files."""

from __future__ import annotations

import io
import re

import pytest
from openpyxl import Workbook

from maple_data_mcp.modules.cihi import client, constants
from maple_data_mcp.shared import cache as cache_module
from maple_data_mcp.shared.errors import InvalidInput, NotFound


@pytest.fixture(autouse=True)
def _clear_cache():
    cache_module._caches.clear()
    yield


_SLUG = "30-day-stroke-in-hospital-mortality"
_EN_PAGE_URL = f"{constants.BASE_URL}{constants.INDICATOR_PATH}{_SLUG}"
_FR_PAGE_URL = f"{constants.BASE_URL}/fr/indicateurs/mortalite-avc"
_EN_FILE = (
    f"{constants.BASE_URL}/sites/default/files/document/data-file/955-{_SLUG}-data-table-en.xlsx"
)
_FR_FILE = _EN_FILE.replace("-en.xlsx", "-fr.xlsx")


def _page(file_url: str, name: str) -> str:
    return f"""<html><head><link hreflang="fr" href="{_FR_PAGE_URL}"></head><body><main>
<h1>{name}</h1>
<div class="view-metadata-summary"><div class="col-12"><div>Risk-adjusted rate of deaths.</div></div>
<ul class="indicator-meta-summary"><li>Data updated: <strong>September 2026</strong></li>
<li>Update frequency: <strong>Every year</strong></li></ul>
<div class="indicator-topics"><ul><li>Acute care </li><li>Health outcomes </li></ul></div></div>
<a href="{file_url}">download the data file</a></main></body></html>"""


def _workbook(header: list[str], rows: list[list[object]]) -> bytes:
    book = Workbook()
    book.create_sheet("Instructions", 0)
    book.remove(book.worksheets[1])
    sheet = book.create_sheet("Table 1")
    sheet.append(["Table 1 30-Day Stroke In-Hospital Mortality"])
    sheet.append([*header, None])
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


_HEADER = ["Reporting level", "Place or organization", "Time frame", "Risk-adjusted rate"]
_ROWS = [
    ["Province/territory", "Alberta", "2024–2025", 12.1],
    ["Province/territory", "Alberta", "2025–2026", 10.8],
    ["Facility", "University of Alberta Hospital (Alta.)", "2025–2026", 15.9],
    ["Province/territory", "Ontario", "2025–2026", 11.0],
]


async def test_search_crawls_library_until_no_new_indicators(httpx_mock):
    page = '<main><a href="/en/indicators/30-day-stroke-in-hospital-mortality">30-Day Stroke In-Hospital Mortality</a><a href="/en/indicators/hospitalized-strokes">Hospitalized Strokes</a></main>'
    httpx_mock.add_response(url=f"{constants.LIBRARY_URL}?page=0", text=page)
    httpx_mock.add_response(url=f"{constants.LIBRARY_URL}?page=1", text=page)
    result = await client.search_indicators("stroke mortality")
    assert [r.slug for r in result.indicators] == [_SLUG]


async def test_get_indicator_reads_facts_and_data_file(httpx_mock):
    httpx_mock.add_response(url=_EN_PAGE_URL, text=_page(_EN_FILE, "30-Day Stroke"))
    detail = await client.get_indicator(_SLUG)
    assert detail.description == "Risk-adjusted rate of deaths."
    assert detail.facts["Data updated"] == "September 2026"
    assert detail.facts["Topics"] == "Acute care, Health outcomes"
    assert detail.data_file_url == _EN_FILE


async def test_get_indicator_data_filters_place_and_columns(httpx_mock):
    httpx_mock.add_response(url=_EN_PAGE_URL, text=_page(_EN_FILE, "30-Day Stroke"))
    httpx_mock.add_response(url=_EN_FILE, content=_workbook(_HEADER, _ROWS))
    data = await client.get_indicator_data(
        _SLUG, place="alberta", filters={"reporting level": "Province/territory"}, limit=1
    )
    assert data.tables == ["Table 1"]
    assert data.columns == _HEADER
    assert data.matching_rows == 2
    assert data.rows == [
        {
            "Reporting level": "Province/territory",
            "Place or organization": "Alberta",
            "Time frame": "2025–2026",
            "Risk-adjusted rate": "10.8",
        }
    ]


async def test_french_data_follows_hreflang(httpx_mock):
    httpx_mock.add_response(url=_EN_PAGE_URL, text=_page(_EN_FILE, "30-Day Stroke"))
    httpx_mock.add_response(url=_FR_PAGE_URL, text=_page(_FR_FILE, "Mortalité AVC"))
    french_header = ["Niveau de déclaration", "Lieu ou organisme", "Période", "Taux"]
    httpx_mock.add_response(url=_FR_FILE, content=_workbook(french_header, _ROWS))
    data = await client.get_indicator_data(_SLUG, place="Ontario", lang="fr")
    assert data.rows[0]["Lieu ou organisme"] == "Ontario"


async def test_input_validation(httpx_mock):
    with pytest.raises(InvalidInput):
        await client.get_indicator("../../etc")
    with pytest.raises(InvalidInput):
        await client.get_indicator_data(_SLUG, limit=0)
    httpx_mock.add_response(url=_EN_PAGE_URL, text=_page(_EN_FILE, "30-Day Stroke"))
    httpx_mock.add_response(url=_EN_FILE, content=_workbook(_HEADER, _ROWS))
    with pytest.raises(InvalidInput, match="Unknown column"):
        await client.get_indicator_data(_SLUG, filters={"Bogus": "x"})
    with pytest.raises(InvalidInput, match="Unknown table"):
        await client.get_indicator_data(_SLUG, table="Table 9")


async def test_missing_data_file_is_not_found(httpx_mock):
    httpx_mock.add_response(
        url=re.compile(re.escape(_EN_PAGE_URL)), text="<main><h1>No file</h1></main>"
    )
    with pytest.raises(NotFound, match="no downloadable data table"):
        await client.get_indicator_data(_SLUG)
