"""Client for CIHI's Indicator Library pages and XLSX data tables.

Pages are parsed with BeautifulSoup; data tables with openpyxl in
read-only mode. Indicators are addressed by their English page slug;
French pages and files are reached through each page's hreflang link.
"""

from __future__ import annotations

import io
import re
from typing import Any
from urllib.parse import urlparse

import httpx
from bs4 import BeautifulSoup, Tag
from openpyxl import load_workbook

from maple_data_mcp.modules.cihi import constants
from maple_data_mcp.modules.cihi.schemas import (
    IndicatorData,
    IndicatorDetail,
    IndicatorRef,
    IndicatorSearchResult,
)
from maple_data_mcp.shared.cache import cached_fetch
from maple_data_mcp.shared.envelope import make_provenance
from maple_data_mcp.shared.errors import InvalidInput, NotFound, UpstreamError, UpstreamUnavailable
from maple_data_mcp.shared.http import get_raw
from maple_data_mcp.shared.rate_limiter import get_limiter

_LIMITER = get_limiter(
    constants.RATE_LIMIT_SOURCE,
    rate=constants.RATE_LIMIT_PER_SECOND,
    capacity=constants.RATE_LIMIT_CAPACITY,
)
_SLUG = re.compile(r"^[a-z0-9-]{3,200}$")
_DATA_FILE = re.compile(r"data-table-(en|fr)\.xlsx$")


async def _get(url: str) -> httpx.Response:
    await _LIMITER.acquire()
    try:
        return await get_raw(url, timeout=120.0)
    except httpx.HTTPStatusError as exc:
        if exc.response.status_code == 404:
            raise NotFound(f"cihi: nothing published at {url}.") from exc
        raise UpstreamError(f"cihi: {url} returned HTTP {exc.response.status_code}.") from exc
    except httpx.HTTPError as exc:
        raise UpstreamUnavailable(f"cihi: {url} did not respond in time.") from exc


async def _page(url: str) -> tuple[str, bool]:
    async def fetch() -> str:
        return (await _get(url)).text

    return await cached_fetch(f"cihi:page:{url}", constants.CACHE_TTL_PAGE_SECONDS, fetch)


def _clean(text: str) -> str:
    return " ".join(text.split())


def _parse_library(html: str) -> list[IndicatorRef]:
    soup = BeautifulSoup(html, "html.parser")
    main = soup.find("main") or soup
    refs: list[IndicatorRef] = []
    for link in main.find_all("a", href=True):
        href = str(link["href"])
        if href.startswith(constants.INDICATOR_PATH):
            slug = href[len(constants.INDICATOR_PATH) :].strip("/")
            refs.append(
                IndicatorRef(slug=slug, name=_clean(link.get_text()), url=constants.BASE_URL + href)
            )
    return refs


async def _library() -> tuple[list[IndicatorRef], bool]:
    async def fetch() -> list[IndicatorRef]:
        seen: dict[str, IndicatorRef] = {}
        for page in range(constants.LIBRARY_MAX_PAGES):
            response = await _get(f"{constants.LIBRARY_URL}?page={page}")
            refs = _parse_library(response.text)
            new = [r for r in refs if r.slug not in seen]
            if not new:
                break
            seen.update({r.slug: r for r in new})
        return list(seen.values())

    return await cached_fetch("cihi:library", constants.CACHE_TTL_LIBRARY_SECONDS, fetch)


async def search_indicators(query: str = "") -> IndicatorSearchResult:
    refs, cached = await _library()
    words = query.lower().split()
    matches = [r for r in refs if all(w in r.name.lower() for w in words)]
    return IndicatorSearchResult(
        indicators=matches,
        total_matches=len(matches),
        provenance=make_provenance(
            source=constants.RATE_LIMIT_SOURCE,
            url=constants.LIBRARY_URL,
            cached=cached,
            schema_name="cihi.IndicatorSearchResult",
            freshness="indicator list cached 7 days",
        ),
    )


def _slug(indicator: str) -> str:
    value = indicator.strip().rstrip("/").lower()
    page_prefix = (constants.BASE_URL + constants.INDICATOR_PATH).lower()
    value = value.removeprefix(page_prefix)
    if not _SLUG.match(value):
        raise InvalidInput(
            f"indicator must be a slug from cihi_search_indicators, got {indicator!r}."
        )
    return value


async def _page_url(slug: str, lang: str) -> str:
    english = f"{constants.BASE_URL}{constants.INDICATOR_PATH}{slug}"
    if lang != "fr":
        return english
    html, _ = await _page(english)
    soup = BeautifulSoup(html, "html.parser")
    alternate = soup.find("link", hreflang="fr")
    if not isinstance(alternate, Tag) or not alternate.get("href"):
        raise NotFound(f"cihi: no French page for {slug!r}.")
    return str(alternate["href"])


def _data_file(soup: BeautifulSoup) -> str | None:
    for link in soup.find_all("a", href=True):
        href = str(link["href"])
        if _DATA_FILE.search(href):
            return href if href.startswith("http") else constants.BASE_URL + href
    return None


async def get_indicator(indicator: str, lang: str = "en") -> IndicatorDetail:
    slug = _slug(indicator)
    url = await _page_url(slug, lang)
    html, cached = await _page(url)
    soup = BeautifulSoup(html, "html.parser")
    heading = soup.find("h1")
    summary = soup.find(class_="view-metadata-summary")
    description = None
    facts: dict[str, str] = {}
    if isinstance(summary, Tag):
        first = summary.find("div", class_="col-12")
        description = _clean(first.get_text()) if isinstance(first, Tag) else None
        for item in summary.select("ul.indicator-meta-summary li"):
            label, _, value = _clean(item.get_text()).partition(":")
            if value:
                facts[label.strip()] = value.strip()
        topics = [_clean(t.get_text()) for t in summary.select(".indicator-topics li")]
        if topics:
            facts["Thèmes" if lang == "fr" else "Topics"] = ", ".join(topics)
    return IndicatorDetail(
        slug=slug,
        name=_clean(heading.get_text()) if isinstance(heading, Tag) else slug,
        description=description,
        facts=facts,
        page_url=url,
        data_file_url=_data_file(soup),
        provenance=make_provenance(
            source=constants.RATE_LIMIT_SOURCE,
            url=url,
            cached=cached,
            schema_name="cihi.IndicatorDetail",
        ),
    )


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_workbook(body: bytes) -> dict[str, tuple[str, list[str], list[list[str]]]]:
    """Each Table sheet as (title, header, rows)."""
    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    tables: dict[str, tuple[str, list[str], list[list[str]]]] = {}
    for sheet in workbook.worksheets:
        if not sheet.title.lower().startswith(("table", "tableau")):
            continue
        rows = sheet.iter_rows(values_only=True)
        title_row = next(rows, ())
        header_row = next(rows, ())
        header = [_cell(v) for v in header_row]
        while header and not header[-1]:
            header.pop()
        width = len(header)
        data = [cells for row in rows if any(cells := [_cell(v) for v in row[:width]])]
        tables[sheet.title] = (_cell(title_row[0] if title_row else ""), header, data)
    workbook.close()
    return tables


async def _tables(url: str) -> tuple[dict[str, tuple[str, list[str], list[list[str]]]], bool]:
    if urlparse(url).hostname not in constants.ALLOWED_HOSTS:
        raise UpstreamError(f"cihi: unexpected data file host in {url}.")

    async def fetch() -> dict[str, tuple[str, list[str], list[list[str]]]]:
        response = await _get(url)
        if len(response.content) > constants.MAX_FILE_BYTES:
            raise UpstreamError(f"cihi: {url} is larger than this tool reads.")
        try:
            return _parse_workbook(response.content)
        except Exception as exc:  # openpyxl raises several unrelated types
            raise UpstreamError(f"cihi: {url} is not a readable XLSX file.") from exc

    return await cached_fetch(f"cihi:data:{url}", constants.CACHE_TTL_DATA_SECONDS, fetch)


async def get_indicator_data(
    indicator: str,
    *,
    place: str | None = None,
    filters: dict[str, str] | None = None,
    table: str | None = None,
    limit: int = constants.ROWS_DEFAULT,
    lang: str = "en",
) -> IndicatorData:
    """Rows of an indicator's data table; the latest rows come last in the file."""
    if limit < 1 or limit > constants.ROWS_MAX:
        raise InvalidInput(f"limit must be between 1 and {constants.ROWS_MAX}, got {limit}.")
    detail = await get_indicator(indicator, lang)
    if not detail.data_file_url:
        raise NotFound(f"cihi: {detail.name!r} has no downloadable data table.")
    tables, tables_cached = await _tables(detail.data_file_url)
    if not tables:
        raise UpstreamError(f"cihi: {detail.data_file_url} has no data table sheets.")
    sheet = table or next(iter(tables))
    if sheet not in tables:
        raise InvalidInput(f"Unknown table {sheet!r}; tables are {list(tables)}.")
    title, header, data = tables[sheet]
    by_lower = {h.lower(): i for i, h in enumerate(header)}

    def index(name: str) -> int:
        position = by_lower.get(name.strip().lower())
        if position is None:
            raise InvalidInput(f"Unknown column {name!r}; columns are {header}.")
        return position

    wanted = [(index(k), v.strip().lower()) for k, v in (filters or {}).items()]
    place_index = next(
        (
            i
            for name, i in by_lower.items()
            if name in ("place or organization", "lieu ou organisme", "lieu ou organisation")
        ),
        None,
    )
    needle = (place or "").strip().lower()
    if needle and place_index is None:
        raise InvalidInput(f"This table has no place column; columns are {header}.")

    def keep(row: list[str]) -> bool:
        cells = row + [""] * (len(header) - len(row))
        if needle and place_index is not None and needle not in cells[place_index].lower():
            return False
        return all(cells[i].lower() == v for i, v in wanted)

    matching = [r for r in data if keep(r)]
    kept = matching[-limit:]
    return IndicatorData(
        slug=detail.slug,
        table=title or sheet,
        tables=list(tables),
        columns=header,
        rows=[dict(zip(header, r + [""] * (len(header) - len(r)), strict=True)) for r in kept],
        total_rows=len(data),
        matching_rows=len(matching),
        returned_count=len(kept),
        data_file_url=detail.data_file_url,
        provenance=make_provenance(
            source=constants.RATE_LIMIT_SOURCE,
            url=detail.data_file_url,
            cached=detail.provenance.cached and tables_cached,
            schema_name="cihi.IndicatorData",
            coverage=f"last {len(kept)} of {len(matching)} matching rows",
        ),
    )
